from flask import Flask, request, jsonify, redirect, url_for, render_template, render_template_string, session
import gspread
from google.oauth2.service_account import Credentials
from datetime import datetime, timedelta
import base64
import hmac
import json
import os
import re
import logging
import csv
import io
import secrets
from pathlib import Path
from zoneinfo import ZoneInfo
from werkzeug.security import check_password_hash, generate_password_hash
from werkzeug.utils import secure_filename

# =========================
# 로깅 설정
# =========================
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
)
logger = logging.getLogger(__name__)

app = Flask(__name__)

# =========================
# 설정
# =========================
MENU_SHEET_NAME = "food_menu"
STUDENT_SHEET_NAME = "rfid"
LOGIN_SHEET_NAME = "id"
LOG_SHEET_NAME = "logs"

# 운영 배포에서는 시트 ID를 환경변수로 넣어 정확한 문서에 연결한다.
MENU_SHEET_ID = os.getenv("MENU_SHEET_ID", "")
STUDENT_SHEET_ID = os.getenv("STUDENT_SHEET_ID", "")
LOGIN_SHEET_ID = os.getenv("LOGIN_SHEET_ID", "")
LOG_SHEET_ID = os.getenv("LOG_SHEET_ID", "")

APP_ENV = os.getenv("APP_ENV", os.getenv("FLASK_ENV", "local")).strip().lower()
IS_PRODUCTION = APP_ENV in {"prod", "production"} or bool(
    os.getenv("RENDER") or os.getenv("RAILWAY_ENVIRONMENT") or os.getenv("FLY_APP_NAME")
)

if IS_PRODUCTION:
    missing_sheet_ids = [
        name
        for name, value in {
            "MENU_SHEET_ID": MENU_SHEET_ID,
            "STUDENT_SHEET_ID": STUDENT_SHEET_ID,
            "LOGIN_SHEET_ID": LOGIN_SHEET_ID,
            "LOG_SHEET_ID": LOG_SHEET_ID,
        }.items()
        if not value.strip()
    ]
    if missing_sheet_ids:
        raise RuntimeError("운영 배포에는 시트 ID 환경변수가 필요해: " + ", ".join(missing_sheet_ids))


def get_secret_setting(name, local_default="", min_length=16):
    value = os.getenv(name, "").strip()
    if not value:
        if IS_PRODUCTION:
            raise RuntimeError(f"{name} 환경변수는 운영 배포에서 반드시 설정해야 해.")
        logger.warning("%s 환경변수가 없어 로컬 개발용 기본값을 사용해.", name)
        return local_default
    if IS_PRODUCTION and len(value) < min_length:
        raise RuntimeError(f"{name} 값이 너무 짧아. 운영 배포에서는 최소 {min_length}자 이상으로 설정해줘.")
    return value


SERVICE_ACCOUNT_JSON = os.getenv("SERVICE_ACCOUNT_JSON", "").strip()
SERVICE_ACCOUNT_JSON_B64 = os.getenv("SERVICE_ACCOUNT_JSON_B64", "").strip()
SERVICE_ACCOUNT_JSON_FILE = os.getenv("SERVICE_ACCOUNT_JSON_FILE", "config/service_account.json").strip()
TIME_FORMAT = "%Y-%m-%d %H:%M:%S"
DEFAULT_ADMIN_NAME = os.getenv("DEFAULT_ADMIN_NAME", "관리자")
RFID_DASHBOARD_URL = os.getenv("RFID_DASHBOARD_URL", "http://allergy-monitoring.duckdns.org:5000")
PUBLIC_ADMIN_URL = os.getenv("PUBLIC_ADMIN_URL", "http://allergy-admin.duckdns.org:5001")
EMAIL_DISPATCH_TOKEN = os.getenv("EMAIL_DISPATCH_TOKEN", "").strip()
NOTIFICATION_TIMEZONE = os.getenv("NOTIFICATION_TIMEZONE", "Asia/Seoul").strip() or "Asia/Seoul"
KIOSK_SCAN_API_TOKEN = get_secret_setting("KIOSK_SCAN_API_TOKEN", "local-dev-only-kiosk-token", min_length=24)
DEFAULT_STUDENT_PASSWORD = get_secret_setting("DEFAULT_STUDENT_PASSWORD", "1234", min_length=8)
ALLERGY_DICT_PATH = Path(os.getenv("ALLERGY_DICT_PATH", "data/allergy_dict.json"))
MEAL_UPLOAD_DIR = Path(os.getenv("MEAL_UPLOAD_DIR", "uploads"))
MEAL_UPLOAD_DIR.mkdir(parents=True, exist_ok=True)

app.secret_key = get_secret_setting("FLASK_SECRET_KEY", secrets.token_urlsafe(32), min_length=32)
app.config.update(
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE=os.getenv("SESSION_COOKIE_SAMESITE", "Lax"),
    SESSION_COOKIE_SECURE=os.getenv("SESSION_COOKIE_SECURE", "1" if IS_PRODUCTION else "0").strip() == "1",
    PERMANENT_SESSION_LIFETIME=timedelta(minutes=int(os.getenv("SESSION_TIMEOUT_MINUTES", "60"))),
)

# 최근 태그된 RFID를 학생 추가/수정 화면에서 바로 불러오기 위한 임시 저장소
LATEST_SCANNED_RFID_UID = ""
LATEST_SCANNED_RFID_AT = ""
AUTO_RESET_SECONDS = int(os.getenv("AUTO_RESET_SECONDS", "8"))
LAST_SCAN_STATE = {
    "view": "waiting",
    "title": "학생증을 태그해 주세요",
    "subtitle": "RFID 리더기에 학생증을 가까이 대면\n알러지 위험도를 바로 확인할 수 있어요.",
    "status_text": "대기중",
    "name": "",
    "student_number": "",
    "menu_name": "",
    "student_allergy_state": "",
    "student_allergy_names": [],
    "warning_names": [],
    "unsafe_menus": [],
    "matched_allergies": [],
    "risk_explanation": "",
    "alternative_recommendation": "",
    "needs_staff_review": False,
    "reason": "",
    "uid": "",
    "updated_at": "",
}

def reset_last_scan_state():
    global LAST_SCAN_STATE
    LAST_SCAN_STATE = {
        "view": "waiting",
        "title": "학생증을 태그해 주세요",
        "subtitle": "RFID 리더기에 학생증을 가까이 대면\n알러지 위험도를 바로 확인할 수 있어요.",
        "status_text": "대기중",
        "name": "",
        "student_number": "",
        "menu_name": "",
        "student_allergy_state": "",
        "student_allergy_names": [],
        "warning_names": [],
        "unsafe_menus": [],
        "matched_allergies": [],
        "risk_explanation": "",
        "alternative_recommendation": "",
        "needs_staff_review": False,
        "reason": "",
        "uid": "",
        "updated_at": now_time_str(),
    }

def update_last_scan_state_from_result(result):
    global LAST_SCAN_STATE

    scan = (result or {}).get("scan", {}) or {}
    status = safe_str(scan.get("status")).strip()
    uid = safe_str(scan.get("uid")).strip()
    name = safe_str(scan.get("name")).strip()
    student_number = safe_str(scan.get("student_number")).strip()
    menu_name = safe_str(scan.get("menu_name")).strip()
    reason = safe_str(scan.get("reason")).strip()
    warning_names = normalize_text_list(scan.get("hit_names"))
    student_allergy_names = normalize_text_list(scan.get("student_allergy_names"))
    unsafe_menus = normalize_text_list(scan.get("unsafe_menus"))
    matched_allergies = normalize_text_list(scan.get("matched_allergies"))
    risk_explanation = safe_str(scan.get("risk_explanation")).strip()
    alternative_recommendation = safe_str(scan.get("alternative_recommendation")).strip()
    needs_staff_review = bool(scan.get("needs_staff_review"))

    student_allergy_state = "있음" if student_allergy_names else "없음"

    if status == "미등록":
        LAST_SCAN_STATE = {
            "view": "notfound",
            "title": "등록되지 않은 학생입니다",
            "subtitle": "태그는 감지됐지만 학생 정보가 시스템에 등록되어 있지 않아요.\n관리자 페이지에서 RFID와 학생 정보를 먼저 등록해 주세요.",
            "status_text": "미등록",
            "name": "",
            "student_number": "",
            "menu_name": "",
            "student_allergy_state": "",
            "student_allergy_names": [],
            "warning_names": [],
            "unsafe_menus": [],
            "matched_allergies": [],
            "risk_explanation": "",
            "alternative_recommendation": "",
            "needs_staff_review": False,
            "reason": reason,
            "uid": uid,
            "updated_at": now_time_str(),
        }
        return

    if status == "메뉴없음":
        LAST_SCAN_STATE = {
            "view": "notfound",
            "title": "오늘 급식 메뉴가 등록되지 않았어요",
            "subtitle": "오늘 날짜 메뉴를 관리자 페이지에서 먼저 등록해 주세요.",
            "status_text": "메뉴없음",
            "name": name,
            "student_number": student_number,
            "menu_name": "",
            "student_allergy_state": student_allergy_state,
            "student_allergy_names": student_allergy_names,
            "warning_names": [],
            "unsafe_menus": [],
            "matched_allergies": [],
            "risk_explanation": "",
            "alternative_recommendation": "",
            "needs_staff_review": False,
            "reason": reason,
            "uid": uid,
            "updated_at": now_time_str(),
        }
        return

    if status in ("경고", "OK"):
        LAST_SCAN_STATE = {
            "view": "warning" if status == "경고" else "ok",
            "title": "학생 정보 감지됨",
            "subtitle": "",
            "status_text": "주의" if status == "경고" else "안전",
            "name": name,
            "student_number": student_number,
            "menu_name": menu_name,
            "student_allergy_state": student_allergy_state,
            "student_allergy_names": student_allergy_names,
            "warning_names": warning_names if status == "경고" else [],
            "unsafe_menus": unsafe_menus if status == "경고" else [],
            "matched_allergies": matched_allergies if status == "경고" else [],
            "risk_explanation": risk_explanation if status == "경고" else "",
            "alternative_recommendation": alternative_recommendation if status == "경고" else "",
            "needs_staff_review": needs_staff_review if status == "경고" else False,
            "reason": reason,
            "uid": uid,
            "updated_at": now_time_str(),
        }
        return

    reset_last_scan_state()

def get_display_state():
    state = dict(LAST_SCAN_STATE)
    updated_at = safe_str(state.get("updated_at")).strip()
    if not updated_at:
        return state

    try:
        updated_dt = datetime.strptime(updated_at, TIME_FORMAT)
        elapsed = (datetime.now() - updated_dt).total_seconds()
        if elapsed >= max(AUTO_RESET_SECONDS, 1) and state.get("view") in {"warning", "ok", "notfound"}:
            reset_last_scan_state()
            return dict(LAST_SCAN_STATE)
    except Exception as e:
        logger.warning(f"last_scan 자동 초기화 시간 계산 실패: {e}")

    return state


ALLERGY_MAP = {
    1: "계란",
    2: "우유",
    3: "메밀",
    4: "땅콩",
    5: "콩",
    6: "밀가루",
    7: "고등어",
    8: "게",
    9: "새우",
    10: "돼지고기",
    11: "복숭아",
    12: "토마토",
    13: "보존제(아황산)",
    14: "호두",
    15: "닭고기",
    16: "소고기",
    17: "오징어",
    18: "조개류(굴, 전복, 홍합 포함)",
    19: "잣",
}

def safe_str(x) -> str:
    return "" if x is None else str(x)

# =========================
# 구글시트 연결
# =========================
SCOPES = [
    "https://www.googleapis.com/auth/spreadsheets",
    "https://www.googleapis.com/auth/drive",
]

def _load_service_account_credentials():
    if SERVICE_ACCOUNT_JSON_B64:
        decoded = base64.b64decode(SERVICE_ACCOUNT_JSON_B64).decode("utf-8")
        return Credentials.from_service_account_info(json.loads(decoded), scopes=SCOPES)

    if SERVICE_ACCOUNT_JSON:
        value = SERVICE_ACCOUNT_JSON
        if value.startswith("base64:"):
            decoded = base64.b64decode(value[len("base64:"):]).decode("utf-8")
            return Credentials.from_service_account_info(json.loads(decoded), scopes=SCOPES)
        if value.lstrip().startswith("{"):
            return Credentials.from_service_account_info(json.loads(value), scopes=SCOPES)
        return Credentials.from_service_account_file(value, scopes=SCOPES)

    return Credentials.from_service_account_file(SERVICE_ACCOUNT_JSON_FILE, scopes=SCOPES)

def _connect_sheets(retries=3, delay=3.0):
    import time
    last_exc = None
    for attempt in range(1, retries + 1):
        try:
            creds = _load_service_account_credentials()
            gc = gspread.authorize(creds)

            def _open_sheet(sheet_id, title):
                try:
                    if safe_str(sheet_id).strip():
                        return gc.open_by_key(sheet_id).sheet1
                except Exception as e:
                    logger.warning(f"시트 ID 연결 실패 ({title}): {e}")
                return gc.open(title).sheet1

            m_ws   = _open_sheet(MENU_SHEET_ID, MENU_SHEET_NAME)
            s_ws   = _open_sheet(STUDENT_SHEET_ID, STUDENT_SHEET_NAME)
            i_ws   = _open_sheet(LOGIN_SHEET_ID, LOGIN_SHEET_NAME)
            lo_ws  = _open_sheet(LOG_SHEET_ID, LOG_SHEET_NAME)
            logger.info("Google Sheets 연결 성공")
            return m_ws, s_ws, i_ws, lo_ws
        except FileNotFoundError:
            raise RuntimeError(
                f"서비스 계정 파일을 찾을 수 없어: {SERVICE_ACCOUNT_JSON or SERVICE_ACCOUNT_JSON_FILE}"
            )
        except Exception as e:
            last_exc = e
            logger.warning(f"Google Sheets 연결 실패 (시도 {attempt}/{retries}): {e}")
            if attempt < retries:
                time.sleep(delay)
    raise RuntimeError(f"Google Sheets 연결에 실패했어 ({retries}번 시도): {last_exc}")

menu_ws, student_ws, id_ws, logs_ws = _connect_sheets()

# =========================
# 유틸
# =========================
def now_time_str() -> str:
    return datetime.now().strftime(TIME_FORMAT)

def today_sheet_str() -> str:
    return datetime.now().strftime("%Y%m%d")

NGROK_URL_FILE = os.getenv("NGROK_URL_FILE", "runtime/ngrok_url.txt")

def load_public_base_url() -> str:
    env_url = safe_str(os.getenv("PUBLIC_BASE_URL", "")).strip()
    if env_url:
        return env_url.rstrip("/")

    try:
        url_file = Path(NGROK_URL_FILE)
        if url_file.exists():
            file_url = safe_str(url_file.read_text(encoding="utf-8")).lstrip("\ufeff").strip()
            if file_url:
                return file_url.rstrip("/")
    except Exception as e:
        logger.warning(f"공개 주소 파일 읽기 실패: {e}")

    fallback = safe_str(PUBLIC_ADMIN_URL).strip()
    return fallback.rstrip("/") if fallback else ""

'''

def normalize_key(key) -> str:
    s = safe_str(key)
    for bad in ("\ufeff", "\u200b", "\xa0"):
        s = s.replace(bad, "")
    return " ".join(s.strip().lower().split())

def clean_record_keys(record: dict) -> dict:
    cleaned = {}
    for k, v in (record or {}).items():
        cleaned[normalize_key(k)] = v
    return cleaned

def pick_first_value(record: dict, *keys, default=""):
    record = clean_record_keys(record)
    for key in keys:
        nk = normalize_key(key)
        if nk in record and safe_str(record.get(nk)).strip() != "":
            return safe_str(record.get(nk)).strip()
    return default


def get_sheet_records_raw(ws):
    """get_all_records보다 헤더/빈값 문제에 덜 민감한 raw 파서"""
    values = ws.get_all_values()
    if not values:
        return []
    headers = [normalize_key(h) for h in values[0]]
    records = []
    for row in values[1:]:
        if len(row) < len(headers):
            row = row + [""] * (len(headers) - len(row))
        rec = {}
        for i, header in enumerate(headers):
            if not header:
                continue
            rec[header] = row[i] if i < len(row) else ""
            rec[chr(97+i)] = row[i] if i < len(row) else ""  # a,b,c... positional fallback
        records.append(rec)
    return records

# =========================
# 간단한 TTL 캐시 (Sheets API 과호출 방지)
# =========================
import threading as _threading
import time as _time

class _TTLCache:
    def __init__(self, ttl: float = 8.0):
        self._ttl = ttl
        self._store: dict = {}
        self._lock = _threading.Lock()

    def get(self, key):
        with self._lock:
            entry = self._store.get(key)
            if entry and (_time.monotonic() - entry["ts"]) < self._ttl:
                return entry["value"], True
        return None, False

    def set(self, key, value):
        with self._lock:
            self._store[key] = {"value": value, "ts": _time.monotonic()}

    def invalidate(self, key):
        with self._lock:
            self._store.pop(key, None)

    def invalidate_prefix(self, prefix: str):
        with self._lock:
            for k in list(self._store.keys()):
                if k.startswith(prefix):
                    del self._store[k]

_cache = _TTLCache(ttl=8.0)

def _sheet_cache_key(ws):
    title = getattr(ws, 'title', None) or 'sheet'
    return f"records::{title}"

def _cached_get_all_records(ws):
    key = _sheet_cache_key(ws)
    value, hit = _cache.get(key)
    if hit:
        return value
    records = ws.get_all_records()
    _cache.set(key, records)
    return records

def _invalidate_sheet_cache(ws):
    _cache.invalidate(_sheet_cache_key(ws))


def parse_codes(s):
    if s is None:
        return set()

    if isinstance(s, (list, tuple, set)):
        codes = set()
        for item in s:
            codes |= parse_codes(item)
        return codes

    s = str(s).strip()
    if not s:
        return set()

    codes = set()
    for token in re.findall("[0-9]+", s):
        if not token.isdigit():
            continue
        n = int(token)
        if 1 <= n <= 19:
            codes.add(n)

    return codes

def codes_to_names(codes):
    return [ALLERGY_MAP.get(code, f"알수없음({code})") for code in codes]

def code_string_to_names(s, empty_text="없음"):
    codes = sorted(list(parse_codes(s)))
    if not codes:
        return empty_text
    return ", ".join(codes_to_names(codes))

def parse_timestamp(value):
    s = safe_str(value).strip()
    if not s:
        return None

    for fmt in (TIME_FORMAT, "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d %H:%M", "%Y/%m/%d %H:%M:%S"):
        try:
            return datetime.strptime(s, fmt)
        except Exception:
            pass
    return None


def build_menu_conflicts(student_codes, menus):
    eaten_menu_names = []
    blocked_menu_names = []
    blocked_codes = set()

    for menu in menus:
        menu_name = safe_str(menu.get("menu_name")).strip()
        menu_codes = parse_codes(menu.get("allergy_codes"))
        matched = student_codes & menu_codes

        if matched:
            if menu_name:
                blocked_menu_names.append(menu_name)
            blocked_codes |= matched
        else:
            if menu_name:
                eaten_menu_names.append(menu_name)

    return eaten_menu_names, blocked_menu_names, sorted(blocked_codes)


def merge_menu_allergies(menus):
    merged_codes = set()
    menu_names = []

    for menu in menus:
        menu_name = safe_str(menu.get("menu_name")).strip()
        allergy_codes = parse_codes(menu.get("allergy_codes"))
        if menu_name:
            menu_names.append(menu_name)
        merged_codes |= allergy_codes

    return sorted(list(merged_codes)), menu_names


def get_menu_hit_details(student_codes, menus):
    details = []
    for menu in menus:
        menu_name = safe_str(menu.get("menu_name")).strip()
        menu_codes = parse_codes(menu.get("allergy_codes"))
        hit = sorted(list(student_codes & menu_codes))
        if hit:
            details.append({
                "menu_name": menu_name,
                "hit_codes": hit,
                "hit_names": codes_to_names(hit),
            })
    return details

def normalize_trash(v):
    return "1" if str(v).strip() == "1" else "0"

def find_header_col(ws, header_name):
    headers = ws.row_values(1)
    for idx, h in enumerate(headers, start=1):
        if str(h).strip() == header_name:
            return idx
    return None

def update_cell_by_header(ws, row_index, header_name, value):
    col = find_header_col(ws, header_name)
    if col is None:
        raise ValueError(f"{header_name} 헤더를 찾지 못했어")
    ws.update_cell(row_index, col, value)
    _invalidate_sheet_cache(ws)


def find_header_col_any(ws, header_names):
    for header_name in header_names:
        col = find_header_col(ws, header_name)
        if col is not None:
            return col, header_name
    return None, None


def update_cell_by_headers(ws, row_index, header_names, value):
    col, found = find_header_col_any(ws, header_names)
    if col is None:
        raise ValueError(f"헤더를 찾지 못했어: {', '.join(header_names)}")
    ws.update_cell(row_index, col, value)
    _invalidate_sheet_cache(ws)
    return found


def get_record_value(record, *keys, default=""):
    for key in keys:
        if key in record:
            return record.get(key)
    return default


def student_number_to_parts(student_number):
    s = ''.join(ch for ch in safe_str(student_number) if ch.isdigit())
    if len(s) != 5:
        return {
            "student_number": safe_str(student_number).strip(),
            "grade": None,
            "class_no": None,
            "number": None,
            "class_key": "unknown",
            "class_label": "미분류",
            "student_label": safe_str(student_number).strip() or "-",
        }

    grade = int(s[0])
    class_no = int(s[1:3])
    number = int(s[3:5])
    return {
        "student_number": s,
        "grade": grade,
        "class_no": class_no,
        "number": number,
        "class_key": f"{grade}-{class_no}",
        "class_label": f"{grade}학년 {class_no}반",
        "student_label": f"{number}번",
    }


def build_student_number(grade, class_no, number):
    try:
        grade = int(str(grade).strip())
        class_no = int(str(class_no).strip())
        number = int(str(number).strip())
    except Exception:
        return ""

    if grade <= 0 or class_no <= 0 or number <= 0:
        return ""
    if grade >= 10 or class_no >= 100 or number >= 100:
        return ""
    return f"{grade}{class_no:02d}{number:02d}"


def enrich_student_like_row(row):
    parts = student_number_to_parts(row.get("student_number", ""))
    new_row = dict(row)
    new_row.update(parts)
    return new_row


def sort_class_key(class_key):
    if class_key == "unknown":
        return (999, 999)
    try:
        grade, class_no = class_key.split("-")
        return (int(grade), int(class_no))
    except Exception:
        return (999, 999)


def get_available_class_keys(student_rows=None, lunch_log_rows=None):
    keys = {"1-1"}
    for row in student_rows or []:
        if row.get("class_key") and row.get("class_key") != "unknown":
            keys.add(row["class_key"])
    for row in lunch_log_rows or []:
        if row.get("class_key") and row.get("class_key") != "unknown":
            keys.add(row["class_key"])
    return sorted(keys, key=sort_class_key)


def build_class_groups(rows):
    grouped = {}
    for row in rows:
        key = row.get("class_key", "unknown")
        if key not in grouped:
            grouped[key] = {
                "class_key": key,
                "class_label": row.get("class_label", "미분류"),
                "rows": [],
            }
        grouped[key]["rows"].append(row)

    groups = []
    for key in sorted(grouped.keys(), key=sort_class_key):
        item = grouped[key]
        item["rows"].sort(key=lambda x: ((x.get("number") or 999), x.get("student_number", ""), x.get("name", "")))
        groups.append(item)
    return groups


def append_row_by_headers(ws, row_dict, default_value=""):
    headers = [str(x).strip() for x in ws.row_values(1)]
    row = []
    for header in headers:
        row.append(row_dict.get(header, default_value))
    ws.append_row(row)
    _invalidate_sheet_cache(ws)


def append_rows_by_headers(ws, row_dicts, default_value=""):
    rows_to_add = list(row_dicts or [])
    if not rows_to_add:
        return
    headers = [str(x).strip() for x in ws.row_values(1)]
    rows = [[row_dict.get(header, default_value) for header in headers] for row_dict in rows_to_add]
    if hasattr(ws, "append_rows"):
        ws.append_rows(rows)
    else:
        for row in rows:
            ws.append_row(row)
    _invalidate_sheet_cache(ws)


def soft_delete_rows(ws, row_indexes):
    for row_index in row_indexes:
        update_cell_by_header(ws, row_index, "trash", "1")
    _invalidate_sheet_cache(ws)

def restore_rows(ws, row_indexes):
    for row_index in row_indexes:
        update_cell_by_header(ws, row_index, "trash", "0")
    _invalidate_sheet_cache(ws)

def hard_delete_rows(ws, row_indexes):
    for row_index in sorted(row_indexes, reverse=True):
        ws.delete_rows(row_index)
    _invalidate_sheet_cache(ws)

'''

from core.admin_domain import (
    build_class_groups,
    build_lunch_log_stats,
    build_menu_conflicts,
    build_not_eaten_students,
    build_student_number,
    code_string_to_names,
    codes_to_names,
    enrich_student_like_row,
    get_available_class_keys,
    get_menu_hit_details,
    merge_menu_allergies,
    parse_codes,
    parse_timestamp,
    set_allergy_map,
    sort_class_key,
    student_number_to_parts,
)
from core.admin_sheet_utils import (
    _cached_get_all_records,
    _cached_get_sheet_records_raw,
    _invalidate_sheet_cache,
    append_row_by_headers,
    append_rows_by_headers,
    clean_record_keys,
    find_header_col,
    find_header_col_any,
    get_record_value,
    get_sheet_records_raw,
    normalize_key,
    normalize_trash,
    pick_first_value,
    restore_rows,
    hard_delete_rows,
    soft_delete_rows,
    update_cell_by_header,
    update_cell_by_headers,
)
from services.ai_service import (
    analyze_meal_ocr_text,
    generate_allergy_safety_plan,
    generate_daily_ai_brief,
    generate_menu_allergy_review,
    generate_risk_assistance,
    load_allergy_dict,
)
from services.email_service import (
    build_daily_email,
    email_delivery_configured,
    is_valid_email,
    send_email,
)
from services.ocr_service import OCRServiceError, extract_text
from services.sheets_service import merge_duplicate_menu_rows, parse_preview_form, save_ai_meal_analysis

try:
    AI_ALLERGY_DICT = load_allergy_dict(ALLERGY_DICT_PATH)
except Exception as e:
    logger.warning(f"AI 알러지 사전 로드 실패: {e}")
    AI_ALLERGY_DICT = {
        "allergies": {str(code): {"name": name, "keywords": []} for code, name in ALLERGY_MAP.items()},
        "menu_hints": [],
    }

set_allergy_map(ALLERGY_MAP)


def normalize_text_list(value):
    if isinstance(value, (list, tuple, set)):
        return [safe_str(item).strip() for item in value if safe_str(item).strip()]

    text = safe_str(value).strip()
    if not text:
        return []

    return [item.strip() for item in re.split(r"[,\n]+", text) if item.strip()]



def value_is_enabled(value):
    return safe_str(value).strip().lower() in {"1", "true", "yes", "on", "enabled"}


def ensure_sheet_header(ws, header_name):
    headers = [safe_str(value).strip() for value in ws.row_values(1)]
    wanted = safe_str(header_name).strip().lower()
    for index, header in enumerate(headers, start=1):
        if header.lower() == wanted:
            return index

    column = len(headers) + 1
    ws.update_cell(1, column, header_name)
    _invalidate_sheet_cache(ws)
    return column


def update_cell_by_header_create(ws, row_index, header_name, value):
    column = ensure_sheet_header(ws, header_name)
    ws.update_cell(row_index, column, value)
    _invalidate_sheet_cache(ws)


def compact_date_value(value):
    digits = ''.join(ch for ch in safe_str(value) if ch.isdigit())
    return digits[:8] if len(digits) >= 8 else digits


def date_value_to_input(value):
    digits = compact_date_value(value)
    if len(digits) == 8:
        return f"{digits[:4]}-{digits[4:6]}-{digits[6:8]}"
    return safe_str(value).strip()


def is_allowed_meal_upload(filename):
    suffix = Path(safe_str(filename)).suffix.lower()
    return suffix in {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".xlsx", ".csv", ".tsv"}


def is_spreadsheet_meal_upload(filename):
    suffix = Path(safe_str(filename)).suffix.lower()
    return suffix in {".xlsx", ".csv", ".tsv"}


def extract_meal_spreadsheet_text(path):
    suffix = Path(path).suffix.lower()
    rows = []

    if suffix in {".csv", ".tsv"}:
        delimiter = "\t" if suffix == ".tsv" else ","
        raw = Path(path).read_bytes()
        text = None
        for encoding in ("utf-8-sig", "cp949", "euc-kr"):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            text = raw.decode("utf-8", errors="replace")
        reader = csv.reader(io.StringIO(text), delimiter=delimiter)
        for row in reader:
            values = [safe_str(cell).strip() for cell in row]
            if any(values):
                rows.append(values)

    elif suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
        except Exception as exc:
            raise ValueError("엑셀(.xlsx) 파일을 읽으려면 openpyxl 설치가 필요해. requirements.txt 설치를 다시 해줘.") from exc

        wb = load_workbook(path, read_only=True, data_only=True)
        for ws in wb.worksheets:
            rows.append([f"[시트] {ws.title}"])
            for row in ws.iter_rows(values_only=True):
                values = [safe_str(cell).strip() for cell in row]
                if any(values):
                    rows.append(values)
        wb.close()
    else:
        raise ValueError("지원하지 않는 급식표 파일 형식이야.")

    lines = []
    for row in rows[:500]:
        lines.append(" | ".join(cell for cell in row if cell))

    if not lines:
        raise ValueError("스프레드시트에서 읽을 수 있는 내용이 없어.")

    return "\n".join(lines)


def _sheet_cell_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return safe_str(value).strip()


def _student_sheet_key(value):
    text = _sheet_cell_text(value).lower()
    return re.sub(r"[\s_\-./()\[\]:]+", "", text)


def _normalize_student_sheet_rows(raw_rows):
    rows = []
    for row in raw_rows:
        values = [_sheet_cell_text(cell) for cell in row]
        if any(values):
            rows.append(values)
    return rows


def _student_sheet_column_map(header):
    aliases = {
        "rfid_id": {"rfidid", "rfid", "uid", "카드번호", "카드id", "태그id"},
        "name": {"name", "studentname", "이름", "성명", "학생명"},
        "allergy_codes": {"allergycodes", "allergy", "allergies", "알러지", "알레르기", "알러지번호", "알레르기번호"},
        "student_number": {"studentnumber", "studentno", "학번", "학생번호"},
        "grade": {"grade", "학년"},
        "class_no": {"class", "classno", "반"},
        "student_seq": {"number", "no", "seq", "번호", "번"},
        "trash": {"trash", "삭제", "휴지통"},
    }
    mapped = {}
    for index, name in enumerate(header):
        key = _student_sheet_key(name)
        for target, names in aliases.items():
            if key in names and target not in mapped:
                mapped[target] = index
                break
    return mapped


def _student_sheet_value(row, columns, key):
    index = columns.get(key)
    if index is None or index >= len(row):
        return ""
    return _sheet_cell_text(row[index])


def _normalize_student_number_from_sheet(row, columns):
    direct = re.sub(r"\D", "", _student_sheet_value(row, columns, "student_number"))
    if direct:
        return direct

    grade = re.sub(r"\D", "", _student_sheet_value(row, columns, "grade"))
    class_no = re.sub(r"\D", "", _student_sheet_value(row, columns, "class_no"))
    student_seq = re.sub(r"\D", "", _student_sheet_value(row, columns, "student_seq"))
    if grade and class_no and student_seq:
        return f"{grade[-1]}{class_no.zfill(2)[-2:]}{student_seq.zfill(2)[-2:]}"
    return ""


def parse_student_spreadsheet_file(path):
    suffix = Path(path).suffix.lower()
    sheet_sets = []

    if suffix in {".csv", ".tsv"}:
        delimiter = "\t" if suffix == ".tsv" else ","
        raw = Path(path).read_bytes()
        text = None
        for encoding in ("utf-8-sig", "cp949", "euc-kr"):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            text = raw.decode("utf-8", errors="replace")
        sheet_sets.append(("csv", list(csv.reader(io.StringIO(text), delimiter=delimiter))))
    elif suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
        except Exception as exc:
            raise ValueError("엑셀(.xlsx) 파일을 읽으려면 openpyxl 설치가 필요해. requirements.txt 설치를 다시 해줘.") from exc

        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            for ws in wb.worksheets:
                sheet_sets.append((ws.title, list(ws.iter_rows(values_only=True))))
        finally:
            wb.close()
    else:
        raise ValueError("지원하지 않는 학생 시트 파일 형식이야.")

    parsed_rows = []
    warnings = []
    for sheet_name, raw_rows in sheet_sets:
        rows = _normalize_student_sheet_rows(raw_rows)
        if not rows:
            continue

        header_index = None
        columns = {}
        for index, row in enumerate(rows[:20]):
            current = _student_sheet_column_map(row)
            if "name" in current and ("student_number" in current or {"grade", "class_no", "student_seq"} <= set(current)):
                header_index = index
                columns = current
                break

        if header_index is None:
            warnings.append(f"{sheet_name}: 학생 명단 헤더를 찾지 못해서 건너뜀")
            continue

        for row in rows[header_index + 1:]:
            name = _student_sheet_value(row, columns, "name")
            student_number = _normalize_student_number_from_sheet(row, columns)
            rfid_id = _student_sheet_value(row, columns, "rfid_id")
            allergy_codes = sorted(parse_codes(_student_sheet_value(row, columns, "allergy_codes")))
            trash = _student_sheet_value(row, columns, "trash")
            if not name and not student_number and not rfid_id:
                continue
            if trash == "1":
                continue
            parsed_rows.append({
                "rfid_id": rfid_id,
                "name": name,
                "student_number": student_number,
                "allergy_codes": allergy_codes,
            })

    if not parsed_rows:
        raise ValueError("학생 시트에서 추가할 학생을 찾지 못했어.")

    return {"rows": parsed_rows, "warnings": warnings}


def build_meal_preview_context(analysis):
    rows = []
    for item in merge_duplicate_menu_rows((analysis or {}).get("rows", [])):
        allergies = sorted(parse_codes((item or {}).get("allergy_codes", [])))
        try:
            confidence = round(float((item or {}).get("confidence", 0) or 0), 2)
        except Exception:
            confidence = 0.0

        rows.append({
            "date": compact_date_value((item or {}).get("date", "")),
            "menu_name": safe_str((item or {}).get("menu_name")).strip(),
            "allergy_codes": allergies,
            "allergy_codes_text": ",".join(str(code) for code in allergies),
            "confidence": confidence,
            "needs_review": bool((item or {}).get("needs_review")),
            "review_reason": safe_str((item or {}).get("review_reason")).strip(),
            "created_by": safe_str((item or {}).get("created_by")).strip(),
            "created_at": safe_str((item or {}).get("created_at")).strip(),
            "trash": "1" if safe_str((item or {}).get("trash", "0")).strip() == "1" else "0",
        })

    warnings = []
    for item in (analysis or {}).get("warnings", []):
        menu = safe_str((item or {}).get("menu_name") or (item or {}).get("menu")).strip()
        message = safe_str((item or {}).get("message")).strip()
        if menu and message:
            warnings.append({
                "date": compact_date_value((item or {}).get("date", "")),
                "menu_name": menu,
                "message": message,
            })

    warning_count = sum(1 for item in rows if item.get("needs_review"))
    return {
        "rows": rows,
        "warnings": warnings,
        "warning_count": warning_count,
    }


def build_meal_preview_context_from_form(form):
    rows = []
    try:
        row_count = max(int(form.get("row_count", "0") or 0), 0)
    except Exception:
        row_count = 0

    raw_rows = []
    for index in range(row_count):
        menu_name = safe_str(form.get(f"row_menu_name_{index}")).strip()
        if not menu_name:
            continue
        allergy_codes = sorted(parse_codes(form.get(f"row_allergy_codes_{index}", "")))
        try:
            confidence = round(float(form.get(f"row_confidence_{index}", "0") or 0), 2)
        except Exception:
            confidence = 0.0

        raw_rows.append({
            "date": compact_date_value(form.get(f"row_date_{index}", "")),
            "menu_name": menu_name,
            "allergy_codes": allergy_codes,
            "confidence": confidence,
            "needs_review": form.get(f"row_needs_review_{index}") == "on",
            "review_reason": safe_str(form.get(f"row_review_reason_{index}")).strip(),
            "created_by": safe_str(form.get(f"row_created_by_{index}")).strip(),
            "created_at": safe_str(form.get(f"row_created_at_{index}")).strip(),
            "trash": "1" if safe_str(form.get(f"row_trash_{index}", "0")).strip() == "1" else "0",
        })

    for item in merge_duplicate_menu_rows(raw_rows):
        allergy_codes = sorted(parse_codes(item.get("allergy_codes", [])))
        rows.append({
            "date": compact_date_value(item.get("date", "")),
            "menu_name": safe_str(item.get("menu_name")).strip(),
            "allergy_codes": allergy_codes,
            "allergy_codes_text": ",".join(str(code) for code in allergy_codes),
            "confidence": item.get("confidence", 0),
            "needs_review": bool(item.get("needs_review")),
            "review_reason": safe_str(item.get("review_reason")).strip(),
            "created_by": safe_str(item.get("created_by")).strip(),
            "created_at": safe_str(item.get("created_at")).strip(),
            "trash": "1" if safe_str(item.get("trash", "0")).strip() == "1" else "0",
        })

    warnings_raw = safe_str(form.get("warnings_json", "[]")).strip() or "[]"
    try:
        warnings = json.loads(warnings_raw)
    except Exception:
        warnings = []

    preview = {
        "rows": rows,
        "warnings": warnings,
        "warning_count": sum(1 for item in rows if item.get("needs_review")),
    }
    return preview


def build_ai_risk_context(date_value):
    date_value = compact_date_value(date_value) or today_sheet_str()
    menus = get_menu_rows_by_date(date_value, include_trash=False)
    menu_names = [safe_str(row.get("menu_name")).strip() for row in menus if safe_str(row.get("menu_name")).strip()]
    risk_items = []

    for student in get_all_student_rows(include_trash=False):
        student_codes = parse_codes(student.get("allergy_codes"))
        if not student_codes:
            continue

        details = get_menu_hit_details(student_codes, menus)
        if not details:
            continue

        unsafe_menus = [item["menu_name"] for item in details if safe_str(item.get("menu_name")).strip()]
        hit_codes = sorted({code for item in details for code in item.get("hit_codes", [])})
        risk_items.append({
            "student_name": safe_str(student.get("name")).strip(),
            "student_number": safe_str(student.get("student_number")).strip(),
            "student_allergies": codes_to_names(sorted(student_codes)),
            "matched_allergies": codes_to_names(hit_codes),
            "unsafe_menus": unsafe_menus,
            "class_label": safe_str(student.get("class_label")).strip(),
            "student_label": safe_str(student.get("student_label")).strip(),
        })

    return {
        "date": date_value,
        "menus": menus,
        "menu_names": menu_names,
        "risk_items": risk_items,
    }


def build_student_ai_context(student, date_value):
    date_value = compact_date_value(date_value) or today_sheet_str()
    menus = get_menu_rows_by_date(date_value, include_trash=False)
    _, menu_names = merge_menu_allergies(menus)
    student_codes = parse_codes((student or {}).get("allergy_codes"))
    details = get_menu_hit_details(student_codes, menus)
    unsafe_menus = [item["menu_name"] for item in details if safe_str(item.get("menu_name")).strip()]
    hit_codes = sorted({code for item in details for code in item.get("hit_codes", [])})

    return {
        "date": date_value,
        "student": student,
        "menus": menus,
        "menu_names": menu_names,
        "student_allergy_names": codes_to_names(sorted(student_codes)),
        "matched_allergies": codes_to_names(hit_codes),
        "unsafe_menus": unsafe_menus,
    }



def build_student_email_message(student, date_value):
    context = build_student_ai_context(student, date_value)
    return build_daily_email(
        student_name=safe_str((student or {}).get("name")).strip() or "학생",
        date_value=context["date"],
        menu_names=context["menu_names"],
        allergy_names=context["student_allergy_names"],
        unsafe_menus=context["unsafe_menus"],
        matched_names=context["matched_allergies"],
    )


def notification_dispatch_authorized():
    if not EMAIL_DISPATCH_TOKEN:
        return False
    authorization = safe_str(request.headers.get("Authorization")).strip()
    supplied = authorization[7:].strip() if authorization.lower().startswith("bearer ") else ""
    if not supplied:
        supplied = safe_str(request.headers.get("X-Notification-Token")).strip()
    return bool(supplied) and hmac.compare_digest(supplied, EMAIL_DISPATCH_TOKEN)

# =========================
# 로그인 / 권한
# =========================
PASSWORD_HASH_PREFIXES = ("scrypt:", "pbkdf2:", "argon2:", "bcrypt:")


def is_password_hash(value):
    value = safe_str(value).strip()
    return value.startswith(PASSWORD_HASH_PREFIXES)


def hash_password(password):
    return generate_password_hash(safe_str(password).strip())


def verify_password(stored_password, candidate_password):
    stored_password = safe_str(stored_password).strip()
    candidate_password = safe_str(candidate_password).strip()
    if not stored_password or not candidate_password:
        return False
    if is_password_hash(stored_password):
        try:
            return check_password_hash(stored_password, candidate_password)
        except Exception:
            logger.warning("비밀번호 해시 검증 실패")
            return False
    return hmac.compare_digest(stored_password, candidate_password)


def get_all_login_rows(include_trash=False):
    records = id_ws.get_all_records()
    rows = []

    for idx, r in enumerate(records, start=2):
        r = clean_record_keys(r)
        trash = normalize_trash(r.get("trash", 0))

        if not include_trash and trash == "1":
            continue

        rows.append({
            "row_index": idx,
            "id": safe_str(r.get("id")).strip(),
            "name": safe_str(r.get("name")).strip(),
            "pw": safe_str(r.get("pw")).strip(),
            "role": safe_str(r.get("role")).strip().lower(),
            "email": safe_str(r.get("email")).strip(),
            "email_notifications": value_is_enabled(r.get("email_notifications")),
            "notification_time": safe_str(r.get("notification_time")).strip() or "07:30",
            "trash": trash,
        })

    return rows

def is_logged_in():
    return session.get("logged_in") is True

def is_admin():
    return session.get("role") == "a"

def is_student():
    return session.get("role") == "s"

def require_admin():
    if not is_logged_in():
        return False, redirect(url_for("login_page"))
    if not is_admin():
        return False, "관리자만 접근할 수 있어"
    return True, None

def require_admin_api():
    if not is_logged_in():
        return False, (jsonify({"ok": False, "error": "로그인이 필요해"}), 401)
    if not is_admin():
        return False, (jsonify({"ok": False, "error": "관리자만 사용할 수 있어"}), 403)
    return True, None

def require_student():
    if not is_logged_in():
        return False, redirect(url_for("login_page"))
    if not is_student():
        return False, "학생만 접근할 수 있어"
    return True, None

# =========================
# 데이터 조회
# =========================
def get_all_menu_rows(include_trash=False):
    # food_menu는 get_all_records가 헤더/빈값을 이상하게 읽는 경우가 있어서 raw 파서 사용
    records = get_sheet_records_raw(menu_ws)
    rows = []

    for idx, r in enumerate(records, start=2):
        r = clean_record_keys(r)
        trash = normalize_trash(pick_first_value(r, "trash", default=0))

        if not include_trash and trash == "1":
            continue

        date_value = pick_first_value(r, "date", "menu_date", "급식 날짜", "날짜", "일자")
        menu_name_value = pick_first_value(r, "menu_name", "menu", "메뉴명", "식단표", "meal_name", "음식명")
        allergy_value = pick_first_value(r, "allergy_codes", "allergy_code", "알러지코드", "알레르기코드")
        created_by_value = pick_first_value(r, "created_by", "writer", "추가자", default="-")
        created_at_value = pick_first_value(r, "created_at", "timestamp", "등록시각", default="-")

        date_value = safe_str(date_value).strip()
        menu_name_value = safe_str(menu_name_value).strip()
        allergy_value = safe_str(allergy_value).strip()

        # 헤더 인식이 꼬인 경우를 대비해서 food_menu 기본 열 순서(A=date, B=menu_name, C=allergy_codes)를 마지막으로 한 번 더 본다.
        if not date_value:
            date_value = safe_str(r.get('date') or r.get('a') or '').strip()
        if not menu_name_value:
            menu_name_value = safe_str(r.get('menu name') or r.get('menu_name') or r.get('b') or '').strip()
        if not allergy_value:
            allergy_value = safe_str(r.get('allergy codes') or r.get('allergy_codes') or r.get('c') or '').strip()

        if not date_value and not menu_name_value and not allergy_value:
            continue

        rows.append({
            "row_index": idx,
            "date": date_value or "-",
            "menu_name": menu_name_value or "-",
            "allergy_codes": allergy_value,
            "allergy_names": code_string_to_names(allergy_value),
            "created_by": safe_str(created_by_value).strip() or "-",
            "created_at": safe_str(created_at_value).strip() or "-",
            "trash": trash,
        })

    rows.sort(key=lambda x: (x["date"], x["menu_name"], -x["row_index"]), reverse=True)
    return rows

def get_menu_groups(include_trash=False):
    rows = get_all_menu_rows(include_trash=include_trash)
    grouped = {}

    for row in rows:
        date = row["date"]
        if date not in grouped:
            grouped[date] = {
                "date": date,
                "row_indexes": [],
                "allergy_names_set": set(),
                "created_by": row["created_by"],
                "created_at": row["created_at"],
                "trash": row["trash"],
                "menus": [],
            }

        grouped[date]["row_indexes"].append(row["row_index"])
        grouped[date]["menus"].append(row)

        for code in parse_codes(row["allergy_codes"]):
            grouped[date]["allergy_names_set"].add(ALLERGY_MAP.get(code, f"알수없음({code})"))

        if row["created_at"] > grouped[date]["created_at"]:
            grouped[date]["created_at"] = row["created_at"]
            grouped[date]["created_by"] = row["created_by"]

    result = []
    for date, item in grouped.items():
        allergy_names_text = ", ".join(sorted(item["allergy_names_set"])) if item["allergy_names_set"] else "없음"
        if not item["menus"]:
            allergy_names_text = "-"
        item["menus"].sort(key=lambda row: row.get("row_index", 0))
        result.append({
            "date": item["date"],
            "row_indexes": sorted(item["row_indexes"]),
            "menu_count": len(item["menus"]),
            "menus": item["menus"],
            "allergy_names": allergy_names_text,
            "created_by": item["created_by"],
            "created_at": item["created_at"],
            "trash": item["trash"],
        })

    today = today_sheet_str()

    def date_sort_key(item):
        date = safe_str(item.get("date")).strip()
        digits = compact_date_value(date)
        if len(digits) != 8:
            return (3, date)
        if digits >= today:
            return (0, digits)
        return (1, -int(digits))

    result.sort(key=date_sort_key)
    return result

def get_all_student_rows(include_trash=False):
    # allergy_codes like "1,2,3,12" are mangled by get_all_records() into 12312,
    # so student rows must be read from raw values.
    records = _cached_get_sheet_records_raw(student_ws)
    rows = []

    for idx, raw in enumerate(records, start=2):
        r = clean_record_keys(raw)
        trash = normalize_trash(pick_first_value(r, "trash", default=0))

        if not include_trash and trash == "1":
            continue

        rfid_id = safe_str(pick_first_value(r, "rfid_id", "uid")).strip()
        name = safe_str(pick_first_value(r, "name", "student_name")).strip()
        student_number = safe_str(pick_first_value(r, "student_number", "student_no", "학번")).strip()
        allergy_value = safe_str(pick_first_value(r, "allergy_codes", "allergy_code", "알러지코드", "알레르기코드")).strip()
        created_by = safe_str(pick_first_value(r, "created_by", "writer", "추가자")).strip()
        created_at = safe_str(pick_first_value(r, "created_at", "timestamp", "등록시각")).strip()

        if not allergy_value:
            allergy_value = safe_str(r.get("c")).strip()

        row = {
            "row_index": idx,
            "rfid_id": rfid_id,
            "name": name,
            "student_number": student_number,
            "allergy_codes": allergy_value,
            "allergy_names": code_string_to_names(allergy_value, empty_text="없음"),
            "created_by": created_by,
            "created_at": created_at,
            "trash": trash,
        }
        rows.append(enrich_student_like_row(row))

    rows.sort(key=lambda x: ((x.get("grade") or 999), (x.get("class_no") or 999), (x.get("number") or 999), x["name"]))
    return rows

def get_all_lunch_log_rows(include_trash=False):
    """logs 시트에서 스캔 로그를 읽어 최신순으로 반환.
    지원 헤더 예시: scanned_at / timestamp, uid / rfid_id
    학생 정보(학번, 반 등)는 rfid 시트와 조인해서 보완.
    """
    records = get_sheet_records_raw(logs_ws)
    student_map = {row["rfid_id"]: row for row in get_all_student_rows(include_trash=True)}
    rows = []

    for idx, raw in enumerate(records, start=2):
        r = clean_record_keys(raw)

        uid = pick_first_value(r, "uid", "rfid_id")
        name = pick_first_value(r, "name", "student_name")
        result = pick_first_value(r, "result", default="-")
        hit_codes = pick_first_value(r, "hit_codes", "hit_code", "allergy_hit_codes")
        menu_date = pick_first_value(r, "menu_date", "date")
        scanned_at = pick_first_value(r, "scanned_at", "timestamp", "created_at", default="-")

        student = student_map.get(uid, {})
        student_number = safe_str(student.get("student_number", "")).strip()
        parsed_dt = parse_timestamp(scanned_at)

        if not name:
            name = safe_str(student.get("name", "")).strip()

        resolved_hit_codes = safe_str(hit_codes).strip()
        if not resolved_hit_codes and safe_str(result).strip() == "경고":
            try:
                student_codes = parse_codes(student.get("allergy_codes"))
                menu_rows = get_menu_rows_by_date(menu_date, include_trash=True) if menu_date else []
                _, _, blocked_codes = build_menu_conflicts(student_codes, menu_rows)
                if blocked_codes:
                    resolved_hit_codes = ",".join(str(x) for x in blocked_codes)
            except Exception:
                pass

        normalized_result = safe_str(result).strip()
        if normalized_result != "경고":
            hit_name_text = "-"
        else:
            candidates = [resolved_hit_codes, hit_codes]
            hit_name_text = "-"
            for candidate in candidates:
                names = code_string_to_names(candidate, empty_text="-")
                if names != "-":
                    hit_name_text = names
                    break
            if hit_name_text == "-":
                try:
                    student_codes = parse_codes(student.get("allergy_codes"))
                    if student_codes:
                        hit_name_text = ", ".join(codes_to_names(sorted(student_codes)))
                    else:
                        hit_name_text = "확인 필요"
                except Exception:
                    hit_name_text = "확인 필요"

        scan_date = parsed_dt.strftime("%Y%m%d") if parsed_dt else ""
        row = {
            "row_index": idx,
            "date": menu_date or "-",
            "scan_date": scan_date,
            "rfid_id": uid,
            "student_number": student_number,
            "name": name,
            "result": result,
            "hit_codes": resolved_hit_codes,
            "hit_names": hit_name_text,
            "scanned_at": scanned_at,
            "parsed_dt": parsed_dt,
        }
        rows.append(enrich_student_like_row(row))

    rows.sort(
        key=lambda x: (x.get("parsed_dt") or datetime.min, x.get("row_index", 0)),
        reverse=True,
    )
    return rows

def build_lunch_log_stats(rows, all_students=None):
    """최신 스캔 기준으로 오늘 이용 현황을 집계한다.
    같은 학생이 여러 번 찍어도 가장 최근 기록 1개만 반영한다.
    오늘 총 인원과 이용률은 등록된 학생 기준으로만 계산한다.
    """
    latest_by_person = {}
    for row in rows:
        student_key = (
            safe_str(row.get("rfid_id")).strip()
            or safe_str(row.get("student_number")).strip()
            or safe_str(row.get("name")).strip()
            or f"row-{row.get('row_index', '')}"
        )
        if student_key not in latest_by_person:
            latest_by_person[student_key] = row

    latest_rows = list(latest_by_person.values())
    registered_students = [
        row for row in (all_students or [])
        if safe_str(row.get("trash", "0")) != "1"
    ]
    registered_student_numbers = {
        safe_str(row.get("student_number")).strip()
        for row in registered_students
        if safe_str(row.get("student_number")).strip()
    }
    registered_rfid_ids = {
        safe_str(row.get("rfid_id")).strip()
        for row in registered_students
        if safe_str(row.get("rfid_id")).strip()
    }

    if registered_students:
        latest_rows = [
            row for row in latest_rows
            if (
                safe_str(row.get("student_number")).strip() in registered_student_numbers
                or safe_str(row.get("rfid_id")).strip() in registered_rfid_ids
            )
        ]

    total_count = len(latest_rows)
    warning_count = sum(1 for row in latest_rows if safe_str(row.get("result")) == "경고")
    ok_count = sum(1 for row in latest_rows if safe_str(row.get("result")) == "OK")
    other_count = total_count - warning_count - ok_count
    total_students = len(registered_students)
    utilization_rate = round((total_count / total_students) * 100, 1) if total_students else 0.0
    return {
        "total_count": total_count,
        "warning_count": warning_count,
        "ok_count": ok_count,
        "other_count": other_count,
        "total_students": total_students,
        "utilization_rate": utilization_rate,
        "scan_count": len(rows),
    }


def build_not_eaten_students(all_students, date_rows, selected_class="all"):
    """해당 날짜에 실제로 급식을 먹은 학생(OK/경고)을 제외한 학생 목록을 만든다."""
    eligible_students = []
    for student in all_students or []:
        if safe_str(student.get("trash", "0")) == "1":
            continue
        if selected_class != "all" and safe_str(student.get("class_key")) != selected_class:
            continue
        eligible_students.append(student)

    eaten_student_numbers = set()
    latest_by_person = {}
    for row in date_rows or []:
        student_key = (
            safe_str(row.get("student_number")).strip()
            or safe_str(row.get("rfid_id")).strip()
            or safe_str(row.get("name")).strip()
        )
        if not student_key:
            continue
        if student_key not in latest_by_person:
            latest_by_person[student_key] = row

    for row in latest_by_person.values():
        result = safe_str(row.get("result")).strip()
        student_no = safe_str(row.get("student_number")).strip()
        if result in {"OK", "경고"} and student_no:
            eaten_student_numbers.add(student_no)

    not_eaten_rows = []
    eaten_rows = []
    for student in eligible_students:
        student_no = safe_str(student.get("student_number")).strip()
        if student_no and student_no in eaten_student_numbers:
            eaten_rows.append(student)
        else:
            not_eaten_rows.append(student)

    sort_key = lambda row: (
        int(safe_str(row.get("grade", "0")) or 0),
        int(safe_str(row.get("class_no", "0")) or 0),
        int(safe_str(row.get("number", "0")) or 0),
        safe_str(row.get("student_number")),
    )
    not_eaten_rows.sort(key=sort_key)
    eaten_rows.sort(key=sort_key)

    return {
        "total_students": len(eligible_students),
        "eaten_count": len(eaten_rows),
        "not_eaten_count": len(not_eaten_rows),
        "not_eaten_rows": not_eaten_rows,
        "eaten_rows": eaten_rows,
    }

def get_student_by_student_number(student_number, include_trash=False):
    for row in get_all_student_rows(include_trash=include_trash):
        if row["student_number"] == student_number:
            return row
    return None

def get_login_by_id(login_id, include_trash=False):
    for row in get_all_login_rows(include_trash=include_trash):
        if row["id"] == login_id:
            return row
    return None

def get_student_by_rfid_id(rfid_id, include_trash=False):
    for row in get_all_student_rows(include_trash=include_trash):
        if row["rfid_id"] == safe_str(rfid_id).strip():
            return row
    return None

def get_menu_rows_by_date(date_str, include_trash=False):
    return [row for row in get_all_menu_rows(include_trash=include_trash) if row["date"] == safe_str(date_str).strip()]

def write_scan_log(uid, name, result, hit_codes_list, menu_date):
    row_dict = {
        "scanned_at": now_time_str(),
        "timestamp": now_time_str(),
        "uid": safe_str(uid),
        "rfid_id": safe_str(uid),
        "name": safe_str(name),
        "student_name": safe_str(name),
        "result": safe_str(result),
        "hit_codes": ",".join(str(x) for x in hit_codes_list) if hit_codes_list else "",
        "menu_date": safe_str(menu_date),
        "date": safe_str(menu_date),
    }
    append_row_by_headers(logs_ws, row_dict)

def process_scan(uid):
    uid = safe_str(uid).strip()
    today = today_sheet_str()

    if not uid:
        return {
            "ok": False,
            "error": "uid missing",
            "scan": {
                "uid": None,
                "name": None,
                "status": "오류",
                "reason": "uid가 비어있음",
                "menu_date": today,
                "menu_name": None,
                "student_allergy_names": "",
                "hit_codes": "",
                "hit_names": "",
                "unsafe_menus": [],
                "matched_allergies": [],
                "risk_explanation": "",
                "alternative_recommendation": "",
                "needs_staff_review": False,
            },
        }

    student = get_student_by_rfid_id(uid, include_trash=False)
    if not student:
        write_scan_log(uid, None, "미등록", [], today)
        return {
            "ok": True,
            "scan": {
                "uid": uid,
                "name": None,
                "student_number": "",
                "status": "미등록",
                "reason": "학생DB(rfid)에 UID가 없음",
                "menu_date": today,
                "menu_name": None,
                "student_allergy_names": "",
                "hit_codes": "",
                "hit_names": "",
                "unsafe_menus": [],
                "matched_allergies": [],
                "risk_explanation": "",
                "alternative_recommendation": "",
                "needs_staff_review": False,
            },
        }

    menus = get_menu_rows_by_date(today, include_trash=False)
    if not menus:
        write_scan_log(uid, student.get("name"), "메뉴없음", [], today)
        return {
            "ok": True,
            "scan": {
                "uid": uid,
                "name": student.get("name"),
                "student_number": student.get("student_number"),
                "status": "메뉴없음",
                "reason": f"점심DB(food_menu)에 {today} 날짜 메뉴가 없음",
                "menu_date": today,
                "menu_name": None,
                "student_allergy_names": code_string_to_names(student.get("allergy_codes"), empty_text="없음"),
                "hit_codes": "",
                "hit_names": "",
                "unsafe_menus": [],
                "matched_allergies": [],
                "risk_explanation": "",
                "alternative_recommendation": "",
                "needs_staff_review": False,
            },
        }

    student_name = safe_str(student.get("name"))
    student_codes = parse_codes(student.get("allergy_codes"))
    student_allergy_names_text = code_string_to_names(student.get("allergy_codes"), empty_text="없음")
    student_allergy_names = codes_to_names(sorted(student_codes))
    merged_menu_codes, menu_names = merge_menu_allergies(menus)
    merged_menu_names = ", ".join(menu_names)
    hit = sorted(list(student_codes & set(merged_menu_codes)))
    hit_names = codes_to_names(hit)

    if hit:
        menu_hit_details = get_menu_hit_details(student_codes, menus)
        unsafe_menus = [item["menu_name"] for item in menu_hit_details if safe_str(item.get("menu_name")).strip()]
        detail_lines = [f"{item['menu_name']} → {', '.join(item['hit_names'])}" for item in menu_hit_details]
        detail_text = "\n".join(detail_lines)
        ai_assist = generate_risk_assistance(
            student_name=student_name,
            student_allergies=student_allergy_names,
            unsafe_menus=unsafe_menus,
            matched_allergies=hit_names,
            all_menus=menu_names,
        )
        write_scan_log(uid, student_name, "경고", hit, today)
        return {
            "ok": True,
            "scan": {
                "uid": uid,
                "name": student_name,
                "student_number": student.get("student_number"),
                "status": "경고",
                "reason": f"겹치는 알러지: {', '.join(hit_names)}\n\n메뉴별 주의 항목:\n{detail_text}",
                "menu_date": today,
                "menu_name": merged_menu_names,
                "student_allergy_names": student_allergy_names_text,
                "hit_codes": ",".join(str(x) for x in hit),
                "hit_names": ", ".join(hit_names),
                "unsafe_menus": unsafe_menus,
                "matched_allergies": ai_assist.get("matched_allergies", hit_names),
                "risk_explanation": ai_assist.get("risk_explanation", ""),
                "alternative_recommendation": ai_assist.get("alternative_recommendation", ""),
                "needs_staff_review": bool(ai_assist.get("needs_staff_review", True)),
            },
        }

    write_scan_log(uid, student_name, "OK", [], today)
    return {
        "ok": True,
        "scan": {
            "uid": uid,
            "name": student_name,
            "student_number": student.get("student_number"),
            "status": "OK",
            "reason": f"문제 없음\n오늘 메뉴: {merged_menu_names}",
            "menu_date": today,
            "menu_name": merged_menu_names,
            "student_allergy_names": student_allergy_names_text,
            "hit_codes": "",
            "hit_names": "",
            "unsafe_menus": [],
            "matched_allergies": [],
            "risk_explanation": "",
            "alternative_recommendation": "",
            "needs_staff_review": False,
        },
    }

def get_trash_rows():
    trash_rows = []

    for row in get_all_student_rows(include_trash=True):
        if row["trash"] == "1":
            trash_rows.append({
                "kind": "학생",
                "row_index": row["row_index"],
                "main_name": row["name"],
                "sub_info": f"RFID: {row['rfid_id']} / 학번: {row['student_number']}",
                "allergy_names": row["allergy_names"],
                "created_by": row["created_by"],
                "created_at": row["created_at"],
            })

    for row in get_all_menu_rows(include_trash=True):
        if row["trash"] == "1":
            trash_rows.append({
                "kind": "급식",
                "row_index": row["row_index"],
                "main_name": row["menu_name"],
                "sub_info": f"급식 날짜: {row['date']}",
                "allergy_names": row["allergy_names"],
                "created_by": row["created_by"],
                "created_at": row["created_at"],
            })

    return trash_rows

# =========================
# 로그인 페이지
# =========================
LOGIN_HTML = """
<!doctype html>
<html lang="ko">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width,initial-scale=1">
  <title>로그인</title>
  <style>
    * { box-sizing: border-box; }
    body {
      margin: 0;
      min-height: 100vh;
      display: flex;
      align-items: center;
      justify-content: center;
      background: #0d0f12;
      color: #f1f5f9;
      font-family: system-ui, "Segoe UI", "Apple SD Gothic Neo", "Malgun Gothic", sans-serif;
    }
    .login-card {
      width: 100%;
      max-width: 430px;
      background: #181b20;
      border: 1px solid #2a2f37;
      border-radius: 22px;
      padding: 28px;
      box-shadow: 0 18px 40px rgba(0,0,0,0.35);
    }
    .login-title {
      font-size: 28px;
      font-weight: 900;
      margin-bottom: 8px;
    }
    .login-sub {
      color: #9aa4b2;
      margin-bottom: 24px;
      font-size: 14px;
    }
    .form-group { margin-bottom: 16px; }
    .form-group label {
      display: block;
      margin-bottom: 8px;
      font-weight: 800;
    }
    .form-group input {
      width: 100%;
      padding: 12px 14px;
      border-radius: 14px;
      border: 1px solid #2a2f37;
      background: #111418;
      color: #f1f5f9;
      font-size: 15px;
      outline: none;
    }
    .login-btn {
      width: 100%;
      border: none;
      border-radius: 14px;
      padding: 13px 14px;
      font-weight: 900;
      font-size: 15px;
      background: #f3f4f6;
      color: #111827;
      cursor: pointer;
    }
    .error-box {
      margin-bottom: 14px;
      padding: 12px 14px;
      border-radius: 12px;
      background: rgba(239,68,68,0.12);
      color: #fca5a5;
      border: 1px solid rgba(239,68,68,0.24);
      font-size: 14px;
      font-weight: 700;
    }
    .helper-row{display:flex;align-items:center;gap:10px;flex-wrap:wrap;margin-top:8px;}
  .helper-text{font-size:12px;color:#9fb0c7;}
</style>
</head>
<body>
  <div class="login-card">
    <div class="login-title">로그인</div>
    <div class="login-sub">관리자 또는 학생 계정으로 로그인할 수 있어.</div>

    {% if error %}
      <div class="error-box">{{ error }}</div>
    {% endif %}

    <form method="post" action="{{ url_for('login_page') }}">
      <div class="form-group">
        <label>아이디</label>
        <input type="text" name="login_id" autocomplete="username">
      </div>

      <div class="form-group">
        <label>비밀번호</label>
        <input type="password" name="login_pw" autocomplete="current-password">
      </div>

      <button type="submit" class="login-btn">로그인</button>
    </form>
  </div>
</body>
</html>
"""

# =========================
# 관리자 페이지 HTML
# =========================
BASE_HTML = r"""
<!doctype html>
<html lang="ko">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width,initial-scale=1">
  <title>{{ page_title }}</title>
  <style>
    * { box-sizing: border-box; }

    :root {
      --bg: #0d0f12;
      --sidebar: #121417;
      --topbar: #17191d;
      --card: #181b20;
      --card-2: #1d2127;
      --line: #2a2f37;
      --line-soft: rgba(255,255,255,0.06);
      --text: #f1f5f9;
      --muted: #9aa4b2;
      --hover: #232831;
      --active: #2d333c;
      --primary: #f3f4f6;
      --danger: #ef4444;
      --danger-bg: rgba(239,68,68,0.12);
      --input: #111418;
      --warn-bg: rgba(245,158,11,0.12);
      --content-glow: none;
      --card-shadow: none;
      --card-shadow-soft: none;
      --stat-neutral-bg: #171b21;
      --stat-green-bg: #13281f;
      --stat-red-bg: #2a1317;
      --stat-blue-bg: #102130;
      --toolbar-bg: #14181e;
      --section-soft-bg: #12161d;
      --warn-row-bg: rgba(127, 29, 29, 0.46);
      --warn-row-border: rgba(239, 68, 68, 0.22);
      --log-header-bg: #171b21;
      --log-header-text: #f8fafc;
      --log-chip-bg: #1a2027;
      --log-chip-text: #e5edf7;
      --log-chip-border: #2f3742;
    }

    body.light-mode {
      --bg: #edf0f3;
      --sidebar: #ffffff;
      --topbar: #ffffff;
      --card: #ffffff;
      --card-2: #f3f4f6;
      --line: #d7dee8;
      --line-soft: rgba(15,23,42,0.08);
      --text: #1e293b;
      --muted: #64748b;
      --hover: #f3f4f6;
      --active: #e5e7eb;
      --primary: #111827;
      --danger: #dc2626;
      --danger-bg: rgba(220,38,38,0.10);
      --input: #ffffff;
      --warn-bg: rgba(245,158,11,0.10);
      --content-glow: none;
      --card-shadow: none;
      --card-shadow-soft: none;
      --stat-neutral-bg: #ffffff;
      --stat-green-bg: #ffffff;
      --stat-red-bg: #ffffff;
      --stat-blue-bg: #ffffff;
      --toolbar-bg: #ffffff;
      --section-soft-bg: #ffffff;
      --warn-row-bg: rgba(220, 38, 38, 0.08);
      --warn-row-border: rgba(220, 38, 38, 0.14);
      --log-header-bg: #f8fafc;
      --log-header-text: #111827;
      --log-chip-bg: #ffffff;
      --log-chip-text: #1e293b;
      --log-chip-border: #d7dee8;
    }

    body {
      margin: 0;
      background: var(--bg);
      color: var(--text);
      font-family: system-ui, "Segoe UI", "Apple SD Gothic Neo", "Malgun Gothic", sans-serif;
    }

    .layout {
      display: grid;
      grid-template-columns: 260px minmax(0, 1fr);
      min-height: 100vh;
      width: 100%;
      max-width: 100%;
      overflow: hidden;
    }

    .sidebar {
      background: var(--sidebar);
      border-right: 1px solid var(--line);
      padding: 22px 18px;
      position: sticky;
      top: 0;
      height: 100vh;
    }

    .brand {
      font-size: 22px;
      font-weight: 900;
      margin-bottom: 28px;
      letter-spacing: -0.02em;
    }

    .section-title {
      color: var(--muted);
      font-size: 12px;
      font-weight: 800;
      margin: 22px 10px 8px 10px;
      text-transform: uppercase;
      letter-spacing: 0.08em;
    }

    .menu-list {
      display: flex;
      flex-direction: column;
      gap: 6px;
    }

    .menu-link {
      display: block;
      padding: 12px 14px;
      border-radius: 12px;
      text-decoration: none;
      color: var(--text);
      font-weight: 700;
      transition: 0.18s ease;
      border: 1px solid transparent;
    }

    .menu-link:hover {
      background: var(--hover);
      border-color: var(--line);
    }

    .menu-link.active {
      background: var(--active);
      border-color: var(--line);
    }

    .menu-link.disabled {
      opacity: 0.45;
      cursor: not-allowed;
      pointer-events: none;
    }

    .main {
      display: flex;
      flex-direction: column;
      min-width: 0;
      max-width: 100%;
      overflow-x: hidden;
    }

    .topbar {
      height: 74px;
      background: var(--topbar);
      border-bottom: 1px solid var(--line);
      display: flex;
      align-items: center;
      justify-content: space-between;
      padding: 0 24px;
    }

    .topbar-left {
      display: flex;
      align-items: center;
      gap: 14px;
      min-width: 0;
    }

    .topbar-logo {
      display: flex;
      align-items: center;
      justify-content: center;
      flex: 0 0 auto;
      min-width: 42px;
    }

    .topbar-logo img {
      height: 34px;
      width: auto;
      display: block;
      object-fit: contain;
    }

    .topbar-title-wrap {
      display: flex;
      flex-direction: column;
      gap: 4px;
      min-width: 0;
    }

    .page-title {
      font-size: 24px;
      font-weight: 900;
    }

    .page-sub {
      color: var(--muted);
      font-size: 13px;
    }

    .topbar-right {
      display: flex;
      gap: 10px;
      align-items: center;
    }

    .ghost-btn,
    .primary-btn,
    .secondary-btn,
    .danger-btn,
    .close-btn {
      border: 1px solid var(--line);
      cursor: pointer;
      border-radius: 12px;
      padding: 10px 14px;
      font-weight: 800;
      font-size: 14px;
      transition: 0.18s ease;
      text-decoration: none;
      display: inline-flex;
      align-items: center;
      justify-content: center;
      color: var(--text);
    }

    .ghost-btn {
      background: var(--card-2);
    }

    .ghost-btn:hover {
      background: var(--hover);
    }

    .theme-toggle-btn {
      background: transparent !important;
      border-color: transparent !important;
      box-shadow: none !important;
      padding-left: 6px;
      padding-right: 6px;
      color: var(--muted);
    }

    .theme-toggle-btn:hover {
      background: transparent !important;
      color: var(--text);
      border-color: transparent !important;
    }

    .primary-btn {
      background: #f3f4f6;
      color: #111827;
      border-color: transparent;
      box-shadow: none;
    }

    .secondary-btn {
      background: var(--card-2);
    }


    body.light-mode .modal-footer .secondary-btn {
      background: #f3f4f6;
      color: #111827;
      border-color: #d1d5db;
    }

    body.light-mode .modal-footer .secondary-btn:hover {
      background: #e5e7eb;
    }

    body.light-mode .modal-footer .primary-btn {
      background: #2563eb;
      color: #ffffff;
      border-color: #2563eb;
    }

    body.light-mode .modal-footer .primary-btn:hover {
      background: #1d4ed8;
      border-color: #1d4ed8;
    }

    .danger-btn {
      background: var(--danger-bg);
      color: #fca5a5;
      border-color: rgba(239,68,68,0.24);
    }

    .danger-btn:hover {
      background: rgba(239,68,68,0.18);
    }

    .content {
      padding: 24px;
      min-width: 0;
      max-width: 100%;
      overflow-x: hidden;
      background: var(--bg);
    }

    .panel {
      width: 100%;
      max-width: 1480px;
      margin: 0 auto;
      min-width: 0;
    }

    .panel-header {
      display: flex;
      justify-content: space-between;
      align-items: center;
      gap: 14px;
      margin-bottom: 18px;
    }

    .panel-header-actions {
      display: flex;
      gap: 10px;
      flex-wrap: wrap;
    }

    .student-header-stats {
      display: flex;
      gap: 8px;
      flex-wrap: wrap;
      align-items: center;
    }

    .student-header-stat {
      min-width: 112px;
      padding: 10px 12px;
      border: 1px solid var(--line-soft);
      background: var(--toolbar-bg);
      border-radius: 8px;
    }

    .student-header-stat span {
      display: block;
      color: var(--muted);
      font-size: 11px;
      font-weight: 800;
      margin-bottom: 4px;
    }

    .student-header-stat b {
      color: var(--text);
      font-size: 20px;
      line-height: 1;
    }

    .card {
      background: var(--card);
      border: 1px solid var(--line);
      border-radius: 18px;
      overflow: hidden;
      box-shadow: none;
      max-width: 100%;
      min-width: 0;
    }

    html, body {
      width: 100%;
      max-width: 100%;
      overflow-x: hidden;
    }

    .page-wrap, .layout, .main, .content, .panel {
      width: 100%;
      max-width: 100%;
      min-width: 0;
      box-sizing: border-box;
    }

    .filters-left {
      padding: 16px;
      border-top: 1px solid var(--line-soft);
      display: flex;
      gap: 8px;
      flex-wrap: wrap;
      justify-content: flex-start;
    }

    .log-scroll-wrap {
      width: 100%;
      max-width: 100%;
      max-height: min(520px, calc(100vh - 360px));
      overflow-y: auto;
      overflow-x: auto;
    }

    .log-scroll-wrap::-webkit-scrollbar {
      width: 10px;
      height: 10px;
    }

    .log-scroll-wrap::-webkit-scrollbar-thumb {
      background: var(--line);
      border-radius: 999px;
    }

    .table-scroll-wrap {
      width: 100%;
      max-width: 100%;
      min-width: 0;
      overflow-x: auto;
    }

    .table-scroll-wrap table,
    .log-scroll-wrap table {
      min-width: 100%;
      width: 100%;
      table-layout: fixed;
      max-width: 100%;
    }

    .panel, .main, .content {
      min-width: 0;
    }

    .log-scroll-wrap {
      max-height: min(520px, calc(100vh - 360px));
      overflow-y: auto;
      overflow-x: auto;
    }

    .log-scroll-wrap thead th {
      position: sticky;
      top: 0;
      z-index: 5;
      background: var(--log-header-bg);
      color: var(--log-header-text);
      font-weight: 900;
      font-size: 14px;
      letter-spacing: 0.01em;
      text-shadow: none;
      box-shadow: none;
      white-space: nowrap;
      border-bottom: 1px solid var(--line);
    }

    .row-warn {
      background: var(--warn-row-bg) !important;
    }

    .row-warn td {
      border-bottom-color: var(--warn-row-border);
    }

    .mini-stat-chip {
      padding: 10px 14px;
      border-radius: 999px;
      background: rgba(255,255,255,0.04);
      border: 1px solid var(--line-soft);
      font-size: 13px;
      color: var(--muted);
      font-weight: 700;
    }

    .mini-stat-chip strong {
      color: var(--text);
      font-size: 15px;
      margin-left: 6px;
    }

    .menu-edit-item {
      background: var(--card-2);
      border: 1px solid var(--line-soft);
      border-radius: 14px;
      padding: 14px 16px;
      margin-bottom: 12px;
    }

    .menu-date-row td {
      background: var(--section-soft-bg);
      color: var(--text);
      font-weight: 900;
      border-top: 1px solid var(--line);
      border-bottom: 1px solid var(--line);
    }

    .menu-date-pill {
      display: inline-flex;
      align-items: center;
      gap: 8px;
      padding: 8px 12px;
      border-radius: 999px;
      border: 1px solid var(--line-soft);
      background: var(--card);
      font-size: 13px;
      font-weight: 900;
    }

    .menu-item-row td {
      background: var(--card);
    }

    table {
      width: 100%;
      border-collapse: collapse;
    }

    thead {
      background: var(--card-2);
    }

    th, td {
      padding: 14px 16px;
      border-bottom: 1px solid var(--line-soft);
      text-align: left;
      vertical-align: top;
      font-size: 14px;
      color: var(--text);
    }

    th {
      color: var(--text);
      font-weight: 900;
    }

    td {
      color: var(--text);
    }

    .badge-trash {
      display: inline-flex;
      align-items: center;
      padding: 4px 10px;
      border-radius: 999px;
      font-size: 12px;
      font-weight: 900;
      background: var(--warn-bg);
      color: #fbbf24;
    }

    .modal-backdrop {
      position: fixed;
      inset: 0;
      background: rgba(0,0,0,0.6);
      display: flex;
      align-items: center;
      justify-content: center;
      padding: 24px;
      z-index: 100;
    }

    .hidden { display: none !important; }

    .modal {
      width: 100%;
      max-width: 700px;
      max-height: 80vh;
      background: var(--card-2);
      border: 1px solid var(--line);
      border-radius: 22px;
      padding: 22px;
      box-shadow: none;
      overflow-y: auto;
      overflow-x: hidden;
    }

    .modal-header {
      display: flex;
      align-items: center;
      justify-content: space-between;
      margin-bottom: 18px;
    }

    .modal-header h2 {
      margin: 0;
      font-size: 24px;
    }

    .close-btn {
      background: transparent;
      color: var(--muted);
    }

    .form-group { margin-bottom: 16px; }

    .form-group label {
      display: block;
      margin-bottom: 8px;
      font-weight: 800;
      color: var(--text);
    }

    .form-group input,
    .form-group select,
    .form-group textarea {
      width: 100%;
      padding: 12px 14px;
      border: 1px solid var(--line);
      border-radius: 14px;
      background: var(--input);
      color: var(--text);
      outline: none;
      font-size: 15px;
    }

    .allergy-row {
      display: flex;
      gap: 10px;
    }

    .tag-wrap {
      display: flex;
      flex-wrap: wrap;
      gap: 10px;
      min-height: 56px;
      padding: 12px;
      border-radius: 14px;
      border: 1px dashed var(--line);
      background: var(--bg);
      margin-top: 14px;
    }

    .modal::-webkit-scrollbar {
      width: 8px;
    }

    .modal::-webkit-scrollbar-track {
      background: transparent;
    }

    .modal::-webkit-scrollbar-thumb {
      background: var(--line);
      border-radius: 999px;
    }

    .modal::-webkit-scrollbar-thumb:hover {
      background: var(--muted);
    }

    .tag {
      display: inline-flex;
      align-items: center;
      gap: 8px;
      padding: 8px 12px;
      border-radius: 999px;
      background: var(--active);
      color: var(--text);
      font-size: 13px;
      font-weight: 800;
      border: 1px solid var(--line);
    }

    .tag button {
      border: none;
      background: transparent;
      color: var(--text);
      cursor: pointer;
      font-weight: 900;
    }

    .modal-footer {
      display: flex;
      justify-content: flex-end;
      gap: 10px;
      margin-top: 20px;
      flex-wrap: wrap;
    }

    .empty-text {
      color: var(--muted);
      font-size: 14px;
    }

    .ai-tool-controls {
      display: grid;
      grid-template-columns: minmax(150px, 210px) minmax(240px, 1fr) minmax(240px, 1fr);
      gap: 12px;
      align-items: end;
    }

    .ai-tool-controls .form-group {
      margin-bottom: 0;
    }

    .ai-safe-shell {
      display: grid;
      gap: 16px;
      background: var(--card);
      border: 1px solid var(--line);
      border-radius: 18px;
      overflow: hidden;
    }

    .ai-safe-hero {
      padding: 24px;
      display: grid;
      grid-template-columns: minmax(0, 1fr) minmax(240px, 0.55fr);
      gap: 18px;
      border-bottom: 1px solid var(--line-soft);
      background: var(--card);
    }

    .ai-safe-title {
      font-size: 30px;
      font-weight: 900;
      line-height: 1.15;
      margin-bottom: 10px;
    }

    .ai-safe-subtitle {
      color: var(--muted);
      line-height: 1.6;
      max-width: 760px;
    }

    .ai-safe-kpis {
      display: grid;
      grid-template-columns: 1fr 1fr;
      gap: 10px;
    }

    .ai-safe-kpi {
      min-height: 96px;
      padding: 15px;
      border: 1px solid var(--line-soft);
      border-radius: 16px;
      background: var(--section-soft-bg);
      display: flex;
      flex-direction: column;
      justify-content: space-between;
    }

    .ai-safe-kpi span {
      color: var(--muted);
      font-size: 12px;
      font-weight: 900;
    }

    .ai-safe-kpi b {
      font-size: 28px;
      line-height: 1;
    }

    .ai-safe-console {
      margin: 0 24px;
      padding: 18px;
      border: 1px solid var(--line-soft);
      border-radius: 16px;
      background: var(--section-soft-bg);
      display: grid;
      gap: 14px;
    }

    .ai-safe-actions {
      display: flex;
      gap: 10px;
      flex-wrap: wrap;
      align-items: center;
    }

    .ai-result-board {
      padding: 0 24px 24px;
      display: grid;
      gap: 14px;
    }

    .ai-result-card {
      border: 1px solid var(--line-soft);
      border-radius: 16px;
      background: var(--section-soft-bg);
      overflow: hidden;
    }

    .ai-result-card .section-banner {
      border-top: 0;
      background: transparent;
    }

    .ai-result-card .table-scroll-wrap {
      border-top: 1px solid var(--line-soft);
    }

    @media (max-width: 980px) {
      .ai-safe-hero {
        grid-template-columns: 1fr;
      }
      .ai-safe-kpis {
        grid-template-columns: 1fr 1fr;
      }
    }

    .ai-status {
      margin: 0 18px 18px;
      padding: 13px 15px;
      border-radius: 14px;
      background: var(--card-2);
      border: 1px solid var(--line-soft);
      color: var(--muted);
      font-weight: 800;
    }

    .ai-status.error {
      color: #fca5a5;
      border-color: rgba(239,68,68,0.26);
      background: var(--danger-bg);
    }

    .ai-chip-row {
      display: flex;
      flex-wrap: wrap;
      gap: 8px;
    }

    .ai-chip {
      display: inline-flex;
      align-items: center;
      padding: 7px 10px;
      border-radius: 999px;
      background: var(--log-chip-bg);
      color: var(--log-chip-text);
      border: 1px solid var(--log-chip-border);
      font-size: 12px;
      font-weight: 900;
      line-height: 1;
    }

    .ai-chip.warn {
      background: var(--danger-bg);
      color: #fca5a5;
      border-color: rgba(239,68,68,0.24);
    }

    .ai-chip.ok {
      background: rgba(34,197,94,0.12);
      color: #86efac;
      border-color: rgba(34,197,94,0.22);
    }

    body.light-mode .ai-chip.warn {
      color: #b91c1c;
      background: rgba(220,38,38,0.08);
    }

    body.light-mode .ai-chip.ok {
      color: #047857;
      background: rgba(16,185,129,0.10);
    }

    .meal-ai-actions {
      display: flex;
      gap: 10px;
      flex-wrap: wrap;
      align-items: center;
      justify-content: flex-end;
    }

    .meal-ai-alert {
      margin: 18px 18px 0;
      border: 1px solid rgba(239,68,68,0.28);
      background: var(--danger-bg);
      color: #fca5a5;
      border-radius: 14px;
      padding: 13px 15px;
      font-weight: 800;
      line-height: 1.5;
    }

    .meal-ai-table {
      min-width: 1180px;
    }

    .meal-ai-table input {
      width: 100%;
      padding: 10px 11px;
      border-radius: 12px;
      border: 1px solid var(--line);
      background: var(--input);
      color: var(--text);
      outline: none;
      font-size: 14px;
    }

    .meal-ai-table .mini-input {
      min-width: 88px;
    }

    .meal-ai-shell {
      background: var(--card);
      border: 1px solid var(--line);
      border-radius: 18px;
      overflow: hidden;
    }

    .meal-ai-hero {
      padding: 24px;
      display: grid;
      grid-template-columns: minmax(0, 1.2fr) minmax(260px, 0.8fr);
      gap: 18px;
      align-items: stretch;
      border-bottom: 1px solid var(--line-soft);
    }

    .meal-ai-title {
      font-size: 30px;
      font-weight: 900;
      line-height: 1.15;
      margin-bottom: 10px;
    }

    .meal-ai-subtitle {
      color: var(--muted);
      line-height: 1.6;
      max-width: 760px;
    }

    .meal-ai-strip {
      display: grid;
      grid-template-columns: repeat(3, minmax(0, 1fr));
      gap: 10px;
      margin-top: 18px;
    }

    .meal-ai-step {
      border: 1px solid var(--line-soft);
      background: rgba(255,255,255,0.035);
      border-radius: 14px;
      padding: 14px;
      min-height: 92px;
    }

    .meal-ai-step strong {
      display: block;
      font-size: 15px;
      margin-bottom: 7px;
    }

    .meal-ai-step span {
      display: block;
      color: var(--muted);
      font-size: 13px;
      line-height: 1.45;
    }

    .meal-ai-scoreboard {
      display: grid;
      grid-template-columns: 1fr 1fr;
      gap: 10px;
    }

    .meal-ai-metric {
      min-height: 114px;
      border: 1px solid var(--line-soft);
      background: var(--section-soft-bg);
      border-radius: 16px;
      padding: 16px;
      display: flex;
      flex-direction: column;
      justify-content: space-between;
    }

    .meal-ai-metric b {
      font-size: 28px;
      line-height: 1;
    }

    .meal-ai-metric span {
      color: var(--muted);
      font-size: 13px;
      font-weight: 800;
    }

    .meal-ai-stage {
      padding: 20px 24px 24px;
    }

    .meal-upload-grid {
      display: grid;
      grid-template-columns: minmax(260px, 0.95fr) minmax(0, 1.05fr);
      gap: 18px;
      align-items: stretch;
    }

    .meal-upload-card,
    .meal-flow-card,
    .meal-review-card {
      border: 1px solid var(--line-soft);
      background: var(--section-soft-bg);
      border-radius: 16px;
      padding: 18px;
    }

    .meal-upload-card input[type="file"] {
      border-style: dashed;
      min-height: 86px;
      padding: 28px 14px;
      cursor: pointer;
    }

    .meal-flow-list {
      display: grid;
      gap: 10px;
      margin-top: 12px;
    }

    .meal-flow-item {
      display: grid;
      grid-template-columns: 38px minmax(0,1fr);
      gap: 12px;
      align-items: start;
      padding: 12px;
      border-radius: 14px;
      background: rgba(255,255,255,0.035);
      border: 1px solid var(--line-soft);
    }

    .meal-flow-no {
      width: 38px;
      height: 38px;
      border-radius: 12px;
      display: inline-flex;
      align-items: center;
      justify-content: center;
      font-weight: 900;
      background: var(--active);
    }

    .meal-workbench {
      padding: 18px;
      display: grid;
      gap: 16px;
      border-top: 1px solid var(--line-soft);
    }

    .meal-review-grid {
      display: grid;
      grid-template-columns: minmax(0,1fr) minmax(300px,0.72fr);
      gap: 16px;
    }

    .meal-analysis-overlay {
      position: fixed;
      inset: 0;
      z-index: 220;
      display: flex;
      align-items: center;
      justify-content: center;
      padding: 24px;
      background: rgba(0,0,0,0.62);
    }

    .meal-analysis-dialog {
      width: min(430px, 100%);
      border: 1px solid var(--line);
      border-radius: 18px;
      background: var(--card);
      padding: 24px;
      box-shadow: 0 22px 70px rgba(0,0,0,0.38);
    }

    .meal-analysis-spinner {
      width: 38px;
      height: 38px;
      border-radius: 999px;
      border: 4px solid var(--line);
      border-top-color: var(--primary);
      animation: mealSpin 0.8s linear infinite;
      margin-bottom: 16px;
    }

    .meal-analysis-title {
      font-size: 22px;
      font-weight: 900;
      margin-bottom: 8px;
    }

    .meal-analysis-message {
      color: var(--muted);
      line-height: 1.55;
      word-break: keep-all;
    }

    .meal-analysis-dialog.error .meal-analysis-spinner {
      display: none;
    }

    .meal-analysis-dialog.error .meal-analysis-title {
      color: #fca5a5;
    }

    .meal-analysis-close {
      margin-top: 18px;
      width: 100%;
    }

    @keyframes mealSpin {
      to { transform: rotate(360deg); }
    }

    @media (max-width: 980px) {
      .meal-ai-hero,
      .meal-upload-grid,
      .meal-review-grid {
        grid-template-columns: 1fr;
      }
      .meal-ai-strip,
      .meal-ai-scoreboard {
        grid-template-columns: 1fr;
      }
    }

    .checkbox-cell { width: 46px; }
    .selection-toolbar {
      display: flex;
      gap: 10px;
      align-items: center;
    }


    .soft-panel {
      padding: 20px;
      border-bottom: 1px solid var(--line-soft);
      display: grid;
      gap: 16px;
      background: var(--section-soft-bg);
    }

    .stats-grid {
      display: grid;
      gap: 14px;
      grid-template-columns: repeat(auto-fit, minmax(170px, 1fr));
    }

    .stat-box {
      padding: 18px;
      border-radius: 20px;
      border: 1px solid var(--line-soft);
      box-shadow: none;
      backdrop-filter: none;
      -webkit-backdrop-filter: none;
    }

    .stat-box.neutral { background: var(--stat-neutral-bg); }
    .stat-box.ok { background: var(--stat-green-bg); border-color: var(--line); }
    .stat-box.warn { background: var(--stat-red-bg); border-color: var(--line); }
    .stat-box.info { background: var(--stat-blue-bg); border-color: var(--line); }

    .toolbar-soft {
      display:flex;
      gap:10px;
      flex-wrap:wrap;
      align-items:center;
      justify-content:space-between;
      padding:14px 16px;
      border-radius:18px;
      background: var(--toolbar-bg);
      border:1px solid var(--line-soft);
      box-shadow: none;
      backdrop-filter: none;
      -webkit-backdrop-filter: none;
    }

    .section-banner {
      padding:14px 18px;
      border-bottom:1px solid var(--line-soft);
      display:flex;
      gap:14px;
      flex-wrap:wrap;
      align-items:center;
      justify-content:space-between;
      background: var(--section-soft-bg);
      box-shadow: none;
    }

    @media (max-width: 1100px) {
      .layout { grid-template-columns: 1fr; }
      .sidebar {
        position: static;
        height: auto;
        border-right: none;
        border-bottom: 1px solid var(--line);
      }
      .topbar {
        height: auto;
        padding: 16px;
        flex-direction: column;
        align-items: flex-start;
        gap: 12px;
      }
      .panel-header {
        flex-direction: column;
        align-items: flex-start;
      }
      .allergy-row {
        flex-direction: column;
      }
      .ai-tool-controls {
        grid-template-columns: 1fr;
      }
    }
  </style>
</head>
<body class="{% if active_tab == 'lunch_log' %}page-lunch-log{% endif %}">
  <div class="layout">
    <aside class="sidebar">
      <div class="brand">관리자</div>

      <div class="section-title">학생 메뉴</div>
      <div class="menu-list">
        <a href="{{ url_for('student_manage_page') }}" class="menu-link {% if active_tab == 'student_manage' %}active{% endif %}">학생관리</a>
      </div>

      <div class="section-title">급식메뉴</div>
      <div class="menu-list">
        <a href="{{ url_for('menu_manage_page') }}" class="menu-link {% if active_tab == 'menu_manage' %}active{% endif %}">급식관리</a>
        <a href="{{ url_for('lunch_log_page') }}" class="menu-link {% if active_tab == 'lunch_log' %}active{% endif %}">급식명단</a>
        <a href="{{ url_for('meal_upload_page') }}" class="menu-link {% if active_tab == 'meal_ai' %}active{% endif %}">AI급식추가</a>
        <a href="{{ url_for('ai_tools_page') }}" class="menu-link {% if active_tab == 'ai_tools' %}active{% endif %}">AI안전도우미</a>
      </div>

      <div class="section-title">기타</div>
      <div class="menu-list">
        <a href="{{ url_for('trash_page') }}" class="menu-link {% if active_tab == 'trash' %}active{% endif %}">휴지통</a>
      </div>
    </aside>

    <section class="main">
      <header class="topbar">
        <div class="topbar-left">
          <div class="topbar-logo">
            <img id="adminLogo" src="{{ url_for('static', filename='logoimage_dark.png') }}" alt="MJ ENTER 로고">
          </div>
          <div class="topbar-title-wrap">
            <div class="page-title">{{ title }}</div>
            <div class="page-sub">{{ subtitle }}</div>
          </div>
        </div>

        <div class="topbar-right">
          <button class="ghost-btn theme-toggle-btn" type="button" id="themeToggleBtn" onclick="toggleAdminTheme()">화이트모드</button>
          {% if logged_in %}
            <button class="ghost-btn" type="button" style="cursor:default;">{{ login_name }}님</button>
            <button class="ghost-btn" onclick="openMyAccountModal()">내 계정관리</button>
            <a class="ghost-btn" href="{{ url_for('logout') }}">로그아웃</a>
          {% else %}
            <a class="ghost-btn" href="{{ url_for('login_page') }}">로그인</a>
          {% endif %}
          <a class="primary-btn" href="{{ url_for('kiosk_page') }}" target="_blank">RFID 페이지 열기</a>
        </div>
      </header>

      <main class="content">
        <section class="panel">
          <div class="panel-header">
            <div>
              {% if active_tab == 'student_manage' %}
              <div class="student-header-stats">
                <div class="student-header-stat">
                  <span>총 재학생</span>
                  <b>{{ student_summary.total }}</b>
                </div>
                <div class="student-header-stat">
                  <span>알레르기 학생</span>
                  <b>{{ student_summary.allergy }}</b>
                </div>
                <div class="student-header-stat">
                  <span>RFID 등록</span>
                  <b>{{ student_summary.rfid }}</b>
                </div>
                <div class="student-header-stat">
                  <span>RFID 미등록</span>
                  <b>{{ student_summary.no_rfid }}</b>
                </div>
              </div>
              {% endif %}
            </div>

            <div class="panel-header-actions">
              {% if active_tab == 'menu_manage' %}
                <button class="primary-btn" onclick="openMenuModal()">급식추가</button>
                <a class="secondary-btn" href="{{ url_for('meal_upload_page') }}">AI업로드</a>
                <div class="selection-toolbar">
                  <button class="secondary-btn" id="toggleMenuDeleteModeBtn" onclick="toggleMenuSelectionMode()">삭제</button>
                  <button class="danger-btn hidden" id="deleteSelectedMenusBtn" onclick="deleteSelectedMenus()">삭제 실행</button>
                </div>

              {% elif active_tab == 'student_manage' %}
                <button class="primary-btn" onclick="openStudentModal()">학생추가</button>
                <button class="secondary-btn" onclick="openStudentSheetModal()">시트로 추가</button>
                <div class="selection-toolbar">
                  <button class="secondary-btn" id="toggleStudentDeleteModeBtn" onclick="toggleStudentSelectionMode()">삭제</button>
                  <button class="danger-btn hidden" id="deleteSelectedStudentsBtn" onclick="deleteSelectedStudents()">삭제 실행</button>
                </div>
              {% endif %}
            </div>
          </div>

          <div class="card">
            {% if active_tab == 'menu_manage' %}
              <div class="table-scroll-wrap">
              <table>
                <thead>
                  <tr>
                    <th class="checkbox-cell hidden" id="menuCheckboxHeader"></th>
                    <th>급식 날짜</th>
                    <th>식단표</th>
                    <th>전체 알러지</th>
                    <th>추가자</th>
                    <th>등록시각</th>
                    <th>수정</th>
                  </tr>
                </thead>
                <tbody>
                  {% for group in menu_groups %}
                  <tr class="menu-date-row">
                    <td class="checkbox-cell hidden menu-checkbox-cell"></td>
                    <td colspan="6">
                      <span class="menu-date-pill">{{ group.date or "-" }} · 메뉴 {{ group.menu_count }}개</span>
                      <span class="empty-text" style="margin-left:10px;">전체 알러지 {{ group.allergy_names or "없음" }}</span>
                    </td>
                  </tr>
                    {% for menu in group.menus %}
                    <tr class="menu-item-row">
                      <td class="checkbox-cell hidden menu-checkbox-cell">
                        <input type="checkbox" class="menu-row-checkbox" value='[{{ menu.row_index }}]'>
                      </td>
                      <td>{{ group.date or "-" }}</td>
                      <td>{{ menu.menu_name or "-" }}</td>
                      <td>{{ menu.allergy_names or "없음" }}</td>
                      <td>{{ menu.created_by or "-" }}</td>
                      <td>{{ menu.created_at or "-" }}</td>
                      <td>
                        <button class="secondary-btn" onclick='openMenuEditModal([{{ menu.row_index }}])'>수정</button>
                      </td>
                    </tr>
                    {% endfor %}
                  {% endfor %}
                </tbody>
              </table>
              </div>

            {% elif active_tab == 'meal_ai' %}
              <div class="meal-ai-shell">
              <div class="meal-ai-hero">
                <div>
                  <div class="meal-ai-title">AI 급식표 정리실</div>
                  <div class="meal-ai-subtitle">이미지, 엑셀, CSV 급식표를 올리면 AI가 날짜별 메뉴와 알레르기 번호를 정리하고 저장 전 검토 화면으로 넘겨줘.</div>
                  <div class="meal-ai-strip">
                    <div class="meal-ai-step">
                      <strong>파일 업로드</strong>
                      <span>급식표 이미지나 스프레드시트 원본을 그대로 넣어.</span>
                    </div>
                    <div class="meal-ai-step">
                      <strong>AI 정리</strong>
                      <span>표 구조와 메뉴명을 food_menu 형식으로 맞춰.</span>
                    </div>
                    <div class="meal-ai-step">
                      <strong>확인 후 저장</strong>
                      <span>관리자가 마지막으로 수정하고 Google Sheets에 저장해.</span>
                    </div>
                  </div>
                </div>
                <div class="meal-ai-scoreboard">
                  <div class="meal-ai-metric">
                    <span>현재 단계</span>
                    <b>{% if meal_mode == 'preview' %}검토{% else %}입력{% endif %}</b>
                  </div>
                  <div class="meal-ai-metric">
                    <span>분석 행</span>
                    <b style="color:#60a5fa;">{{ meal_preview.rows|length if meal_preview else 0 }}</b>
                  </div>
                  <div class="meal-ai-metric">
                    <span>검토 필요</span>
                    <b style="color:#ef4444;">{{ meal_preview.warning_count if meal_preview else 0 }}</b>
                  </div>
                  <div class="meal-ai-metric">
                    <span>저장 대상</span>
                    <b style="color:#22c55e; font-size:23px;">Sheets</b>
                  </div>
                </div>
              </div>

              {% if meal_error %}
                <div class="meal-ai-alert">{{ meal_error }}</div>
              {% endif %}

              {% if meal_mode == 'preview' and meal_preview %}
              <form method="post" action="{{ url_for('meal_save_page') }}">
                <input type="hidden" name="row_count" id="mealAiRowCount" value="{{ meal_preview.rows|length }}">
                <textarea class="hidden" name="ocr_text">{{ meal_ocr_text }}</textarea>
                <textarea class="hidden" name="warnings_json">{{ meal_preview.warnings|tojson }}</textarea>

                <div class="meal-workbench">
                <div class="section-banner">
                  <div>
                    <div style="font-weight:800; font-size:18px;">시트 row 미리보기</div>
                    <div class="empty-text" style="margin-top:4px;">날짜, 메뉴명, 알러지 번호를 확인한 뒤 그대로 저장할 수 있어.</div>
                  </div>
                  <div class="meal-ai-actions">
                    <button class="secondary-btn" type="button" onclick="addMealAiRow()">행 추가</button>
                    <button class="secondary-btn" type="button" onclick="downloadMealAiCsv()">CSV 다운로드</button>
                  </div>
                </div>

                <div class="meal-review-card">
                  <div class="form-group" style="max-width:320px; margin-bottom:14px;">
                    <label>기본 작성자</label>
                    <input id="mealAiCreatedByInput" type="text" name="created_by" value="{{ meal_created_by }}" required>
                  </div>

                  <div class="table-scroll-wrap">
                    <table class="meal-ai-table">
                      <thead>
                        <tr>
                          <th style="width:110px;">date</th>
                          <th style="width:220px;">menu_name</th>
                          <th style="width:150px;">allergy_codes</th>
                          <th style="width:130px;">created_by</th>
                          <th style="width:140px;">created_at</th>
                          <th style="width:70px;">trash</th>
                          <th style="width:90px;">confidence</th>
                          <th style="width:110px;">검토</th>
                          <th style="width:260px;">review_reason</th>
                          <th style="width:78px;">삭제</th>
                        </tr>
                      </thead>
                      <tbody id="mealAiRowPreviewBody">
                        {% for row in meal_preview.rows %}
                        <tr class="meal-ai-preview-row">
                          <td><input class="mini-input" type="text" name="row_date_{{ loop.index0 }}" value="{{ row.date }}" placeholder="20260418" required></td>
                          <td><input type="text" name="row_menu_name_{{ loop.index0 }}" value="{{ row.menu_name }}" required></td>
                          <td><input type="text" name="row_allergy_codes_{{ loop.index0 }}" value="{{ row.allergy_codes_text }}" placeholder="1,2,5,6"></td>
                          <td><input type="text" name="row_created_by_{{ loop.index0 }}" value="{{ row.created_by or meal_created_by }}"></td>
                          <td><input type="text" name="row_created_at_{{ loop.index0 }}" value="{{ row.created_at }}" placeholder="저장 시 자동입력" readonly></td>
                          <td><input class="mini-input" type="text" name="row_trash_{{ loop.index0 }}" value="{{ row.trash }}"></td>
                          <td><input class="mini-input" type="number" name="row_confidence_{{ loop.index0 }}" value="{{ row.confidence }}" step="0.01" min="0" max="1"></td>
                          <td>
                            <label style="display:inline-flex; gap:8px; align-items:center; font-weight:800; white-space:nowrap;">
                              <input type="checkbox" name="row_needs_review_{{ loop.index0 }}" {% if row.needs_review %}checked{% endif %}>
                              필요
                            </label>
                          </td>
                          <td><input type="text" name="row_review_reason_{{ loop.index0 }}" value="{{ row.review_reason }}"></td>
                          <td><button class="danger-btn" style="padding:9px 10px;" type="button" onclick="removeMealAiRow(this)">삭제</button></td>
                        </tr>
                        {% endfor %}
                      </tbody>
                    </table>
                  </div>
                </div>

                <div class="section-banner">
                  <div>
                    <div style="font-weight:800; font-size:18px;">AI 경고와 OCR 원문</div>
                    <div class="empty-text" style="margin-top:4px;">저장 전 검토가 필요한 부분과 OCR 텍스트를 함께 확인해.</div>
                  </div>
                </div>

                <div class="meal-review-grid">
                  <div class="table-scroll-wrap">
                    <table>
                      <thead>
                        <tr>
                          <th style="width:110px;">날짜</th>
                          <th style="width:180px;">메뉴</th>
                          <th>검토 메시지</th>
                        </tr>
                      </thead>
                      <tbody>
                        {% if meal_preview.warnings %}
                          {% for warning in meal_preview.warnings %}
                          <tr class="row-warn">
                            <td>{{ warning.date }}</td>
                            <td>{{ warning.menu_name }}</td>
                            <td style="white-space:normal; word-break:keep-all;">{{ warning.message }}</td>
                          </tr>
                          {% endfor %}
                        {% else %}
                          <tr>
                            <td colspan="3" class="empty-text" style="padding:22px; text-align:center;">추가 경고 없음</td>
                          </tr>
                        {% endif %}
                      </tbody>
                    </table>
                  </div>

                  <div>
                    <div class="form-group" style="margin-bottom:0;">
                      <label>OCR 텍스트</label>
                      <textarea readonly style="min-height:260px;">{{ meal_ocr_text }}</textarea>
                    </div>
                  </div>
                </div>

                <div style="display:flex; gap:10px; flex-wrap:wrap; justify-content:flex-end;">
                  <a class="secondary-btn" href="{{ url_for('meal_upload_page') }}">취소</a>
                  <button class="primary-btn" type="submit">Google Sheets 저장</button>
                </div>
                </div>
              </form>
              {% else %}
              <form id="mealAnalyzeForm" method="post" action="{{ url_for('meal_analyze_page') }}" enctype="multipart/form-data">
                <div class="meal-ai-stage">
                  <div class="meal-upload-grid">
                    <div class="meal-upload-card">
                      <div style="font-weight:900; font-size:20px; margin-bottom:12px;">원본 파일</div>
                    <div class="form-group">
                        <label>급식표 파일</label>
                        <input type="file" name="meal_file" accept="image/*,.xlsx,.csv,.tsv" required>
                        <div class="empty-text" style="margin-top:8px;">PNG, JPG, JPEG, WEBP, BMP, XLSX, CSV, TSV</div>
                    </div>
                    <div class="form-group">
                      <label>등록자 이름</label>
                      <input type="text" name="created_by" value="{{ meal_created_by }}" required>
                    </div>
                    <div class="form-group">
                        <label>AI 보조 메모</label>
                        <textarea name="ocr_hint" rows="4" placeholder="예: 5월 중식표야. 날짜는 표 상단 월/일을 기준으로 잡아줘."></textarea>
                      </div>
                      <div style="display:flex; gap:10px; flex-wrap:wrap; justify-content:flex-end;">
                        <button class="primary-btn" type="submit">분석 시작</button>
                      </div>
                    </div>

                    <div class="meal-flow-card">
                      <div style="font-weight:900; font-size:20px;">정리 흐름</div>
                      <div class="empty-text" style="margin-top:6px;">급식명단 표 화면과 다르게, 여기는 파일을 넣고 AI가 정리하는 작업 공간으로 분리했어.</div>
                      <div class="meal-flow-list">
                        <div class="meal-flow-item">
                          <div class="meal-flow-no">1</div>
                          <div><strong>원본 읽기</strong><div class="empty-text" style="margin-top:4px;">이미지는 OCR, 엑셀/CSV는 셀 내용을 그대로 추출해.</div></div>
                        </div>
                        <div class="meal-flow-item">
                          <div class="meal-flow-no">2</div>
                          <div><strong>메뉴 분리</strong><div class="empty-text" style="margin-top:4px;">여러 날짜와 여러 메뉴를 각각 저장 행으로 나눠.</div></div>
                        </div>
                        <div class="meal-flow-item">
                          <div class="meal-flow-no">3</div>
                          <div><strong>알레르기 코드 추론</strong><div class="empty-text" style="margin-top:4px;">등록된 사전과 메뉴 힌트를 기준으로 번호를 채워.</div></div>
                        </div>
                        <div class="meal-flow-item">
                          <div class="meal-flow-no">4</div>
                          <div><strong>저장 전 검토</strong><div class="empty-text" style="margin-top:4px;">불확실한 행은 표시하고, 바로 수정한 뒤 저장해.</div></div>
                        </div>
                      </div>
                    </div>
                  </div>
                </div>
              </form>
              {% endif %}
              </div>

              <div id="mealAnalysisOverlay" class="meal-analysis-overlay hidden" role="alertdialog" aria-modal="true" aria-labelledby="mealAnalysisTitle">
                <div id="mealAnalysisDialog" class="meal-analysis-dialog">
                  <div class="meal-analysis-spinner" aria-hidden="true"></div>
                  <div id="mealAnalysisTitle" class="meal-analysis-title">분석중입니다</div>
                  <div id="mealAnalysisMessage" class="meal-analysis-message">급식표를 읽고 AI가 메뉴와 알레르기 번호를 정리하고 있어. 잠시만 기다려줘.</div>
                  <button id="mealAnalysisCloseBtn" class="secondary-btn meal-analysis-close hidden" type="button" onclick="hideMealAnalysisOverlay()">확인</button>
                </div>
              </div>

            {% elif active_tab == 'lunch_log' %}
              <div class="soft-panel">
                <div class="stats-grid">
                  <div class="stat-box neutral">
                    <div class="empty-text" style="margin-bottom:8px;">오늘 총 인원</div>
                    <div style="font-size:40px; font-weight:900; line-height:1;">{{ today_stats.total_count }}</div>
                    <div class="empty-text" style="margin-top:10px;">등록 학생 기준</div>
                  </div>
                  <div class="stat-box ok">
                    <div class="empty-text" style="margin-bottom:8px;">오늘 통과</div>
                    <div style="font-size:40px; font-weight:900; line-height:1; color:#22c55e;">{{ today_stats.ok_count }}</div>
                    <div class="empty-text" style="margin-top:10px;">안전 확인 완료</div>
                  </div>
                  <div class="stat-box warn">
                    <div class="empty-text" style="margin-bottom:8px;">오늘 경고</div>
                    <div style="font-size:40px; font-weight:900; line-height:1; color:#ef4444;">{{ today_stats.warning_count }}</div>
                    <div class="empty-text" style="margin-top:10px; color:#fca5a5;">즉시 확인 필요</div>
                  </div>
                  <div class="stat-box info">
                    <div class="empty-text" style="margin-bottom:8px;">오늘 이용률</div>
                    <div style="font-size:40px; font-weight:900; line-height:1; color:#60a5fa;">{{ today_stats.utilization_rate }}%</div>
                    <div class="empty-text" style="margin-top:10px;">{{ today_stats.total_count }} / {{ today_stats.total_students }}명</div>
                  </div>
                </div>

                <div class="toolbar-soft">
                  <form method="get" action="{{ url_for('lunch_log_page') }}" style="display:flex; gap:10px; flex-wrap:wrap; align-items:center;">
                    <input type="date" name="date" value="{{ selected_date_input }}" style="padding:10px 12px; border-radius:12px; border:1px solid var(--line); background:var(--input); color:var(--text);">
                    <input type="hidden" name="class_key" value="{{ selected_class }}">
                    <button class="secondary-btn" type="submit">보기</button>
                    <a class="ghost-btn" href="{{ url_for('lunch_log_page', class_key=selected_class) }}">날짜 초기화</a>
                  </form>
                  <div style="display:flex; gap:12px; flex-wrap:wrap; align-items:center; justify-content:flex-start;">
                    <div class="empty-text">현재 목록 {{ lunch_log_rows|length }}건</div>
                    <div class="empty-text">전체 스캔 {{ today_stats.scan_count }}회</div>
                  </div>
                </div>
              </div>
              <div class="section-banner">
                <div>
                  <div style="font-weight:800; font-size:18px;">실시간 급식 로그</div>
                  <div class="empty-text" style="margin-top:4px;">경고 학생은 빨간색으로 강조돼.</div>
                </div>
                <div style="display:flex; gap:8px; flex-wrap:wrap; align-items:center;">
                  <span style="padding:8px 12px; border-radius:999px; background:var(--log-chip-bg); color:var(--log-chip-text); border:1px solid var(--log-chip-border); font-size:13px; font-weight:700;">경고 {{ today_stats.warning_count }}</span>
                  <span style="padding:8px 12px; border-radius:999px; background:var(--log-chip-bg); color:var(--log-chip-text); border:1px solid var(--log-chip-border); font-size:13px; font-weight:700;">통과 {{ today_stats.ok_count }}</span>
                </div>
              </div>
              <div style="width:100%; max-width:100%; overflow:hidden;">
              <div class="log-scroll-wrap" style="width:100%; max-width:100%;">
                <table style="width:100%; max-width:100%; table-layout:fixed;">
                  <thead>
                    <tr>
                      <th style="width:180px;">스캔 시각</th>
                      <th style="width:110px;">메뉴 날짜</th>
                      <th style="width:90px;">학번</th>
                      <th style="width:60px;">학년</th>
                      <th style="width:60px;">반</th>
                      <th style="width:60px;">번호</th>
                      <th style="width:100px;">이름</th>
                      <th style="width:90px;">결과</th>
                      <th style="width:180px;">걸린 알러지</th>
                    </tr>
                  </thead>
                  <tbody>
                    {% for row in lunch_log_rows %}
                    <tr class="{% if row.result == '경고' %}row-warn{% elif row.result == 'OK' %}row-ok{% endif %}">
                      <td>{{ row.scanned_at }}</td>
                      <td>{{ row.date or "-" }}</td>
                      <td>{{ row.student_number or "-" }}</td>
                      <td>{{ row.grade or "-" }}</td>
                      <td>{{ row.class_no or "-" }}</td>
                      <td>{{ row.number or "-" }}</td>
                      <td>{{ row.name or "-" }}</td>
                      <td>
                        {% if row.result == '경고' %}
                          <span style="color:#ef4444; font-weight:700;">⚠ 경고</span>
                        {% elif row.result == 'OK' %}
                          <span style="color:#22c55e; font-weight:700;">✓ OK</span>
                        {% else %}
                          <span style="color:#f59e0b;">{{ row.result }}</span>
                        {% endif %}
                      </td>
                      <td style="white-space:normal; word-break:keep-all;">{{ row.hit_names or "-" }}</td>
                    </tr>
                    {% endfor %}
                  </tbody>
                </table>
              </div>
              </div>
              <div class="filters-left" style="justify-content:flex-start;">
                {% for class_item in available_classes %}
                <a class="{% if selected_class == class_item.key %}primary-btn{% else %}ghost-btn{% endif %}" href="{{ url_for('lunch_log_page', class_key=class_item.key, date=selected_date_input) }}">{{ class_item.label }}</a>
                {% endfor %}
              </div>

              <div style="padding:16px 18px 20px; border-top:1px solid var(--line-soft); display:grid; gap:14px;">
                <div style="display:flex; gap:10px; flex-wrap:wrap; align-items:center; justify-content:space-between;">
                  <div>
                    <div style="font-weight:800; font-size:17px;">{{ active_lunch_date[:4] }}-{{ active_lunch_date[4:6] }}-{{ active_lunch_date[6:8] }} 미식사 학생</div>
                    <div class="empty-text" style="margin-top:4px;">해당 날짜에 OK 또는 경고로 기록되지 않은 학생 목록이야.</div>
                  </div>
                  <div style="display:flex; gap:8px; flex-wrap:wrap; align-items:center; justify-content:flex-start;">
                    <span class="mini-stat-chip">전체 대상<strong>{{ attendance_summary.total_students }}</strong></span>
                    <span class="mini-stat-chip">식사 완료<strong>{{ attendance_summary.eaten_count }}</strong></span>
                    <span class="mini-stat-chip">미식사<strong>{{ attendance_summary.not_eaten_count }}</strong></span>
                  </div>
                </div>
                <div class="table-scroll-wrap" style="max-height:280px;">
                  <table style="table-layout:fixed; width:100%;">
                    <thead>
                      <tr>
                        <th style="width:100px;">학번</th>
                        <th style="width:60px;">학년</th>
                        <th style="width:60px;">반</th>
                        <th style="width:60px;">번호</th>
                        <th style="width:120px;">이름</th>
                        <th>알러지명</th>
                      </tr>
                    </thead>
                    <tbody>
                      {% if attendance_summary.not_eaten_rows %}
                        {% for row in attendance_summary.not_eaten_rows %}
                        <tr>
                          <td>{{ row.student_number or '-' }}</td>
                          <td>{{ row.grade or '-' }}</td>
                          <td>{{ row.class_no or '-' }}</td>
                          <td>{{ row.number or '-' }}</td>
                          <td>{{ row.name or '-' }}</td>
                          <td style="white-space:normal; word-break:keep-all;">{{ row.allergy_names or '없음' }}</td>
                        </tr>
                        {% endfor %}
                      {% else %}
                        <tr>
                          <td colspan="6" class="empty-text" style="padding:22px; text-align:center;">전원이 식사 완료했어.</td>
                        </tr>
                      {% endif %}
                    </tbody>
                  </table>
                </div>
              </div>

            {% elif active_tab == 'student_manage' %}
              <div class="filters-left" style="border-bottom:1px solid var(--line-soft); border-top:none;">
                <input id="studentSearchInput" type="search" placeholder="학생검색: 이름, 학번, RFID, 알러지" oninput="filterStudentRows()" style="min-width:260px; padding:10px 12px; border-radius:12px; border:1px solid var(--line); background:var(--input); color:var(--text);">
                {% for class_item in available_classes %}
                <a class="{% if selected_class == class_item.key %}primary-btn{% else %}ghost-btn{% endif %}" href="{{ url_for('student_manage_page', class_key=class_item.key) }}">{{ class_item.label }}</a>
                {% endfor %}
              </div>
              <div class="table-scroll-wrap">
              <table>
                <thead>
                  <tr>
                    <th class="checkbox-cell hidden" id="studentCheckboxHeader"></th>
                    <th>RFID ID</th>
                    <th>이름</th>
                    <th>학년</th>
                    <th>반</th>
                    <th>번호</th>
                    <th>학번</th>
                    <th>알러지명</th>
                    <th>추가자</th>
                    <th>등록시각</th>
                    <th>수정</th>
                  </tr>
                </thead>
                <tbody>
                  {% for row in student_rows %}
                  <tr class="student-data-row" data-search="{{ (row.rfid_id ~ ' ' ~ row.name ~ ' ' ~ row.student_number ~ ' ' ~ row.allergy_names ~ ' ' ~ row.class_label ~ ' ' ~ row.student_label)|lower }}">
                    <td class="checkbox-cell hidden student-checkbox-cell">
                      <input type="checkbox" class="student-row-checkbox" value="{{ row.row_index }}">
                    </td>
                    <td>{{ row.rfid_id }}</td>
                    <td>{{ row.name }}</td>
                    <td>{{ row.grade or "-" }}</td>
                    <td>{{ row.class_no or "-" }}</td>
                    <td>{{ row.number or "-" }}</td>
                    <td>{{ row.student_number }}</td>
                    <td>{{ row.allergy_names }}</td>
                    <td>{{ row.created_by }}</td>
                    <td>{{ row.created_at }}</td>
                    <td>
                      <button class="secondary-btn" onclick='openStudentEditModal({{ row.row_index }}, {{ row.rfid_id|tojson }}, {{ row.name|tojson }}, {{ row.student_number|tojson }}, {{ row.allergy_codes|tojson }})'>수정</button>
                    </td>
                  </tr>
                  {% endfor %}
                </tbody>
              </table>
              </div>

            {% elif active_tab == 'ai_tools' %}
              <div class="ai-safe-shell">
                <div class="ai-safe-hero">
                  <div>
                    <div class="ai-safe-title">AI 안전 상황실</div>
                    <div class="ai-safe-subtitle">학생, 날짜, 현장 메모를 기준으로 대체급식 계획과 사고 대응 절차를 즉시 정리하는 화면이야.</div>
                  </div>
                  <div class="ai-safe-kpis">
                    <div class="ai-safe-kpi">
                      <span>등록 학생</span>
                      <b>{{ student_rows|length }}</b>
                    </div>
                    <div class="ai-safe-kpi">
                      <span>급식 날짜</span>
                      <b style="color:#60a5fa;">{{ ai_date_options|length }}</b>
                    </div>
                    <div class="ai-safe-kpi">
                      <span>오늘 경고</span>
                      <b style="color:#ef4444;">{{ today_stats.warning_count }}</b>
                    </div>
                    <div class="ai-safe-kpi">
                      <span>AI 기능</span>
                      <b style="color:#22c55e;">3</b>
                    </div>
                  </div>
                </div>

                <div class="ai-safe-console">
                  <div class="ai-tool-controls" style="width:100%;">
                    <div class="form-group">
                      <label>급식 날짜</label>
                      <select id="aiDateSelect">
                        {% for item in ai_date_options %}
                          <option value="{{ item }}" {% if item == today_sheet %}selected{% endif %}>{{ item[:4] }}-{{ item[4:6] }}-{{ item[6:8] }}</option>
                        {% endfor %}
                      </select>
                    </div>
                    <div class="form-group">
                      <label>학생 선택</label>
                      <select id="aiStudentSelect">
                        <option value="">학생 선택</option>
                        {% for row in student_rows %}
                          <option value="{{ row.student_number }}">{{ row.class_label }} {{ row.student_label }} {{ row.name }} · {{ row.allergy_names or '알러지 없음' }}</option>
                        {% endfor %}
                      </select>
                    </div>
                    <div class="form-group">
                      <label>상황 메모</label>
                      <textarea id="aiContextInput" rows="1" placeholder="예: 우유 알레르기 학생 대체식 먼저 확인"></textarea>
                    </div>
                  </div>
                  <div class="ai-safe-actions">
                    <button class="primary-btn" type="button" onclick="loadAiSafetyPlan()">학생 안전계획</button>
                    <button class="secondary-btn" type="button" onclick="loadAiDailyBrief()">오늘 브리핑</button>
                    <button class="secondary-btn" type="button" onclick="loadAiMenuReview()">메뉴 코드점검</button>
                    <button class="ghost-btn" type="button" onclick="clearAiResults()">결과 지우기</button>
                  </div>
                </div>

              <div id="aiStatusBox" class="ai-status hidden"></div>

              <div class="ai-result-board">
              <div class="ai-result-card">
              <div class="section-banner">
                <div>
                  <div style="font-weight:800; font-size:18px;">학생 안전계획</div>
                  <div class="empty-text" style="margin-top:4px;">선택 학생의 위험 메뉴, 대체급식 후보, 전달 메시지를 표 형태로 확인해.</div>
                </div>
              </div>
              <div id="aiSafetyResult" class="table-scroll-wrap">
                <table>
                  <tbody>
                    <tr>
                      <td class="empty-text" style="padding:22px; text-align:center;">학생을 선택하고 학생 안전계획을 실행해줘.</td>
                    </tr>
                  </tbody>
                </table>
              </div>
              </div>

              <div class="ai-result-card">
              <div class="section-banner">
                <div>
                  <div style="font-weight:800; font-size:18px;">사고 발생 시 대응</div>
                  <div class="empty-text" style="margin-top:4px;">섭취 의심 상황에서 현장 교직원이 바로 볼 수 있는 단계별 대응안이 표시돼.</div>
                </div>
              </div>
              <div id="aiEmergencyResult" class="table-scroll-wrap">
                <table>
                  <tbody>
                    <tr>
                      <td class="empty-text" style="padding:22px; text-align:center;">안전계획 결과와 함께 표시됩니다.</td>
                    </tr>
                  </tbody>
                </table>
              </div>
              </div>

              <div class="ai-result-card">
              <div class="section-banner">
                <div>
                  <div style="font-weight:800; font-size:18px;">오늘 브리핑</div>
                  <div class="empty-text" style="margin-top:4px;">선택 날짜의 전체 주의 학생과 급식실 우선 조치를 확인할 수 있어.</div>
                </div>
              </div>
              <div id="aiBriefResult" class="table-scroll-wrap">
                <table>
                  <tbody>
                    <tr>
                      <td class="empty-text" style="padding:22px; text-align:center;">오늘 브리핑을 실행해줘.</td>
                    </tr>
                  </tbody>
                </table>
              </div>
              </div>

              <div class="ai-result-card">
              <div class="section-banner">
                <div>
                  <div style="font-weight:800; font-size:18px;">메뉴 코드점검</div>
                  <div class="empty-text" style="margin-top:4px;">선택 날짜 급식 메뉴의 알레르기 코드 누락 가능성을 점검해.</div>
                </div>
              </div>
              <div id="aiReviewResult" class="table-scroll-wrap">
                <table>
                  <tbody>
                    <tr>
                      <td class="empty-text" style="padding:22px; text-align:center;">메뉴 코드점검을 실행해줘.</td>
                    </tr>
                  </tbody>
                </table>
              </div>
              </div>
              </div>
              </div>

            {% elif active_tab == 'trash' %}
              <div class="table-scroll-wrap">
              <table>
                <thead>
                  <tr>
                    <th>종류</th>
                    <th>이름/메뉴명</th>
                    <th>보조 정보</th>
                    <th>알러지명</th>
                    <th>추가자</th>
                    <th>등록시각</th>
                    <th>복구</th>
                    <th>완전 삭제</th>
                  </tr>
                </thead>
                <tbody>
                  {% for row in trash_rows %}
                  <tr>
                    <td><span class="badge-trash">{{ row.kind }}</span></td>
                    <td>{{ row.main_name }}</td>
                    <td>{{ row.sub_info }}</td>
                    <td>{{ row.allergy_names }}</td>
                    <td>{{ row.created_by }}</td>
                    <td>{{ row.created_at }}</td>
                    <td>
                      {% if row.kind == '학생' %}
                      <button class="secondary-btn" onclick="restoreStudent({{ row.row_index }})">복구</button>
                      {% else %}
                      <button class="secondary-btn" onclick="restoreMenu({{ row.row_index }})">복구</button>
                      {% endif %}
                    </td>
                    <td>
                      {% if row.kind == '학생' %}
                      <button class="danger-btn" onclick="hardDeleteStudent({{ row.row_index }})">완전 삭제</button>
                      {% else %}
                      <button class="danger-btn" onclick="hardDeleteMenu({{ row.row_index }})">완전 삭제</button>
                      {% endif %}
                    </td>
                  </tr>
                  {% endfor %}
                </tbody>
              </table>
            {% endif %}
          </div>
        </section>
      </main>
    </section>
  </div>

  <!-- 내 계정관리 -->
  <div id="myAccountModal" class="modal-backdrop hidden">
    <div class="modal">
      <div class="modal-header">
        <h2>내 계정관리</h2>
        <button class="close-btn" onclick="closeMyAccountModal()">×</button>
      </div>

      <div class="form-group">
        <label>아이디</label>
        <input type="text" id="myAccountId" value="{{ login_id }}">
      </div>

      <div class="form-group">
        <label>이름</label>
        <input type="text" id="myAccountName" value="{{ login_name }}">
      </div>

      <div class="form-group">
        <label>새 비밀번호</label>
        <input type="text" id="myAccountPw" placeholder="변경할 비밀번호를 입력해줘">
      </div>

      <div class="modal-footer">
        <button class="secondary-btn" onclick="resetMyPasswordToDefault()">기본번호(1234)로 변경</button>
        <button class="secondary-btn" onclick="closeMyAccountModal()">취소</button>
        <button class="primary-btn" onclick="saveMyAccount()">저장</button>
      </div>
    </div>
  </div>

  <!-- 급식 추가 -->
  <div id="menuModal" class="modal-backdrop hidden">
    <div class="modal">
      <div class="modal-header">
        <h2>급식추가</h2>
        <button class="close-btn" onclick="closeMenuModal()">×</button>
      </div>

      <div class="form-group">
        <label>급식 날짜</label>
        <input type="date" id="menuDate">
      </div>

      <div class="form-group">
        <label>추가자</label>
        <input type="text" id="menuCreatedBy" value="{{ default_admin_name }}">
      </div>

      <div class="form-group">
        <label>식단표</label>
        <button type="button" class="secondary-btn" onclick="addMenuItem()">메뉴 추가</button>
      </div>

      <div id="menuAddList"></div>

      <div class="modal-footer">
        <button class="secondary-btn" onclick="closeMenuModal()">취소</button>
        <button class="primary-btn" onclick="submitMenu()">저장</button>
      </div>
    </div>
  </div>

  <!-- 급식 수정 -->
  <div id="menuEditModal" class="modal-backdrop hidden">
    <div class="modal">
      <div class="modal-header">
        <h2>급식 수정</h2>
        <button class="close-btn" onclick="closeMenuEditModal()">×</button>
      </div>

      <input type="hidden" id="editMenuDateOriginal">
      <input type="hidden" id="editMenuRowIndexes">

      <div class="form-group">
        <label>급식 날짜</label>
        <input type="date" id="editMenuDate">
      </div>

      <div class="form-group">
        <label>식단표</label>
        <button type="button" class="secondary-btn" onclick="addEditMenuItem()">메뉴 추가</button>
      </div>

      <div id="editMenuList"></div>

      <div class="modal-footer">
        <button class="secondary-btn" onclick="closeMenuEditModal()">취소</button>
        <button class="primary-btn" onclick="saveMenuEdit()">저장</button>
      </div>
    </div>
  </div>

  <!-- 학생 추가 -->
  <div id="studentModal" class="modal-backdrop hidden">
    <div class="modal">
      <div class="modal-header">
        <h2>학생추가</h2>
        <button class="close-btn" onclick="closeStudentModal()">×</button>
      </div>

      <div class="form-group">
        <label>RFID ID</label>
        <input type="text" id="studentRfidId" placeholder="예: 3004727299">
        <div class="helper-row">
          <button type="button" class="secondary-btn" onclick="fillLatestRfidToAdd()">최근 태그 불러오기</button>
          <span id="latestRfidAddHint" class="helper-text">카드를 태그한 뒤 버튼을 누르면 자동으로 들어가.</span>
        </div>
      </div>

      <div class="form-group">
        <label>학생 이름</label>
        <input type="text" id="studentName" placeholder="예: 황병기">
      </div>

      <div class="form-group">
        <label>학년 / 반 / 번호</label>
        <div class="allergy-row">
          <input type="number" id="studentGrade" min="1" max="9" placeholder="학년" oninput="updateStudentNumberPreview()">
          <input type="number" id="studentClassNo" min="1" max="99" placeholder="반" value="1" oninput="updateStudentNumberPreview()">
          <input type="number" id="studentSeq" min="1" max="99" placeholder="번호" oninput="updateStudentNumberPreview()">
        </div>
      </div>

      <div class="form-group">
        <label>생성될 학번</label>
        <input type="text" id="studentNoPreview" placeholder="예: 10101" readonly>
      </div>

      <div class="form-group">
        <label>초기 비밀번호</label>
        <input type="text" value="1234" readonly>
      </div>

      <div class="form-group">
        <label>알러지 선택 (없으면 선택 안 해도 돼)</label>
        <div class="allergy-row">
          <select id="studentAllergySelect">
            <option value="">알러지를 선택해줘</option>
            <option value="none">없음</option>
            {% for code, name in allergy_map.items() %}
            <option value="{{ code }}">{{ code }}. {{ name }}</option>
            {% endfor %}
          </select>
          <button type="button" class="secondary-btn" onclick="addSelectedStudentAllergy()">알러지 추가</button>
        </div>
      </div>

      <div class="form-group">
        <label>선택된 알러지</label>
        <div id="selectedStudentAllergyList" class="tag-wrap"></div>
      </div>

      <div class="modal-footer">
        <button class="secondary-btn" onclick="closeStudentModal()">취소</button>
        <button class="primary-btn" onclick="submitStudent()">저장</button>
      </div>
    </div>
  </div>

  <!-- 학생 수정 -->
  <div id="studentEditModal" class="modal-backdrop hidden">
    <div class="modal">
      <div class="modal-header">
        <h2>학생 수정</h2>
        <button class="close-btn" onclick="closeStudentEditModal()">×</button>
      </div>

      <input type="hidden" id="editStudentRowIndex">

      <div class="form-group">
        <label>RFID ID</label>
        <input type="text" id="editStudentRfidId" placeholder="예: 3004727299">
        <div class="helper-row">
          <button type="button" class="secondary-btn" onclick="fillLatestRfidToEdit()">최근 태그 불러오기</button>
          <span id="latestRfidEditHint" class="helper-text">현재 학생 카드가 바뀌었으면 새 카드 태그 후 불러오면 돼.</span>
        </div>
      </div>

      <div class="form-group">
        <label>이름</label>
        <input type="text" id="editStudentName">
      </div>

      <div class="form-group">
        <label>학년 / 반 / 번호</label>
        <div class="allergy-row">
          <input type="number" id="editStudentGrade" min="1" max="9" placeholder="학년" oninput="updateEditStudentNumberPreview()">
          <input type="number" id="editStudentClassNo" min="1" max="99" placeholder="반" oninput="updateEditStudentNumberPreview()">
          <input type="number" id="editStudentSeq" min="1" max="99" placeholder="번호" oninput="updateEditStudentNumberPreview()">
        </div>
      </div>

      <div class="form-group">
        <label>학번</label>
        <input type="text" id="editStudentId" readonly>
      </div>

      <div class="form-group">
        <label>새 비밀번호</label>
        <input type="text" id="editStudentPw" placeholder="변경할 비밀번호">
      </div>

      <div class="form-group">
        <label>알러지 선택</label>
        <div class="allergy-row">
          <select id="editStudentAllergySelect">
            <option value="">알러지를 선택해줘</option>
            <option value="none">없음</option>
            {% for code, name in allergy_map.items() %}
            <option value="{{ code }}">{{ code }}. {{ name }}</option>
            {% endfor %}
          </select>
          <button type="button" class="secondary-btn" onclick="addEditStudentAllergy()">알러지 추가</button>
        </div>
      </div>

      <div class="form-group">
        <label>선택된 알러지</label>
        <div id="editStudentAllergyList" class="tag-wrap"></div>
      </div>

      <div class="modal-footer">
        <button class="secondary-btn" onclick="resetStudentPasswordToDefault()">기본번호(1234)로 변경</button>
        <button class="secondary-btn" onclick="closeStudentEditModal()">취소</button>
        <button class="primary-btn" onclick="saveStudentEdit()">저장</button>
      </div>
    </div>
  </div>

  <!-- 시트로 학생 추가 -->
  <div id="studentSheetModal" class="modal-backdrop hidden">
    <div class="modal" style="max-width:1100px;">
      <div class="modal-header">
        <h2>시트로 추가</h2>
        <button class="close-btn" onclick="closeStudentSheetModal()">×</button>
      </div>

      <form id="studentSheetForm" enctype="multipart/form-data">
        <div class="form-group">
          <label>학생 명단 시트</label>
          <input type="file" name="student_sheet" accept=".xlsx,.csv,.tsv" required>
          <div class="empty-text" style="margin-top:8px;">XLSX, CSV, TSV 파일을 지원해. rfid_id, name, allergy_codes, student_number 형식이면 바로 추가돼.</div>
        </div>

        <div id="studentSheetResult" class="meal-ai-alert hidden"></div>
        <div id="studentSheetPreviewWrap" class="table-scroll-wrap hidden" style="margin-top:14px; max-height:420px;"></div>

        <div class="modal-footer">
          <button class="secondary-btn" type="button" onclick="closeStudentSheetModal()">취소</button>
          <button class="secondary-btn hidden" id="studentSheetClearPreviewBtn" type="button" onclick="clearStudentSheetPreview()">다시 선택</button>
          <button class="primary-btn" id="studentSheetPreviewBtn" type="submit">미리보기</button>
          <button class="primary-btn hidden" id="studentSheetConfirmBtn" type="button" onclick="confirmStudentSheetAdd()">확인 후 추가</button>
        </div>
      </form>
    </div>
  </div>

  <script>
    const ALLERGY_MAP = {{ allergy_map | tojson }};
    const DEFAULT_ADMIN_NAME = {{ default_admin_name | tojson }};
    const DEFAULT_STUDENT_PASSWORD = {{ default_student_password | tojson }};

    let selectedStudentAllergies = [];
    let studentSelectionMode = false;
    let menuSelectionMode = false;
    let editStudentAllergies = [];
    let editMenuItems = [];
    let addMenuItems = [];

    function renderTags(targetId, selectedArray, removeFnName) {
      const wrap = document.getElementById(targetId);
      if (!wrap) return;

      wrap.innerHTML = "";

      if (selectedArray.length === 0) {
        wrap.innerHTML = `<div class="empty-text">알러지 없음</div>`;
        return;
      }

      selectedArray.forEach(code => {
        const name = ALLERGY_MAP[code] ?? `알수없음(${code})`;
        const tag = document.createElement("div");
        tag.className = "tag";
        tag.innerHTML = `
          <span>${code}. ${name}</span>
          <button type="button" onclick="${removeFnName}(${code})">×</button>
        `;
        wrap.appendChild(tag);
      });
    }

    function renderStudentAllergyState(isNone = false) {
      const wrap = document.getElementById("selectedStudentAllergyList");
      if (!wrap) return;

      if (isNone) {
        wrap.innerHTML = `<div class="tag"><span>없음</span></div>`;
        return;
      }

      renderTags("selectedStudentAllergyList", selectedStudentAllergies, "removeStudentAllergy");
    }

    function renderEditStudentTags(isNone = false) {
      const wrap = document.getElementById("editStudentAllergyList");
      if (!wrap) return;

      if (isNone) {
        wrap.innerHTML = `<div class="tag"><span>없음</span></div>`;
        return;
      }

      wrap.innerHTML = "";
      if (editStudentAllergies.length === 0) {
        wrap.innerHTML = `<div class="empty-text">알러지 없음</div>`;
        return;
      }

      editStudentAllergies.forEach(code => {
        const name = ALLERGY_MAP[code] ?? `알수없음(${code})`;
        const tag = document.createElement("div");
        tag.className = "tag";
        tag.innerHTML = `
          <span>${code}. ${name}</span>
          <button type="button" onclick="removeEditStudentAllergy(${code})">×</button>
        `;
        wrap.appendChild(tag);
      });
    }

    function renderEditMenuAllergyTags(index) {
      const wrap = document.getElementById(`editMenuAllergyList_${index}`);
      if (!wrap) return;

      const item = editMenuItems[index];
      const codes = item?.allergyCodes || [];

      wrap.innerHTML = "";
      if (codes.length === 0) {
        wrap.innerHTML = `<div class="empty-text">알러지 없음</div>`;
        return;
      }

      codes.forEach(code => {
        const name = ALLERGY_MAP[code] ?? `알수없음(${code})`;
        const tag = document.createElement("div");
        tag.className = "tag";
        tag.innerHTML = `
          <span>${code}. ${name}</span>
          <button type="button" onclick="removeEditMenuAllergy(${index}, ${code})">×</button>
        `;
        wrap.appendChild(tag);
      });
    }

    function renderEditMenuItems() {
      const list = document.getElementById("editMenuList");
      if (!list) return;

      list.innerHTML = "";

      if (!editMenuItems || editMenuItems.length === 0) {
        list.innerHTML = `<div class="empty-text">메뉴가 아직 없어. "메뉴 추가" 버튼을 눌러줘.</div>`;
        return;
      }

      editMenuItems.forEach((item, index) => {
        const wrapper = document.createElement("div");
        wrapper.className = "menu-edit-item";
        const nameValue = (item.menu_name || "").replace(/"/g, "&quot;");

        wrapper.innerHTML = `
          <div class="form-group">
            <label>메뉴명</label>
            <input type="text" class="edit-menu-name" data-index="${index}" value="${nameValue}">
          </div>
          <div class="form-group">
            <label>알러지 선택</label>
            <div class="allergy-row">
              <select id="editMenuAllergySelect_${index}">
                <option value="">알러지를 선택해줘</option>
                ${Object.entries(ALLERGY_MAP).map(([code, name]) => `<option value="${code}">${code}. ${name}</option>`).join("")}
              </select>
              <button type="button" class="secondary-btn" onclick="addEditMenuAllergy(${index})">알러지 추가</button>
              <button type="button" class="danger-btn" onclick="removeEditMenuItem(${index})">메뉴 삭제</button>
            </div>
            <div id="editMenuAllergyList_${index}" class="tag-wrap"></div>
          </div>
        `;

        list.appendChild(wrapper);
        renderEditMenuAllergyTags(index);
      });
    }

    function renderMenuAddAllergyTags(index) {
      const wrap = document.getElementById(`menuAddAllergyList_${index}`);
      if (!wrap) return;

      const item = addMenuItems[index];
      const codes = item?.allergyCodes || [];

      wrap.innerHTML = "";
      if (codes.length === 0) {
        wrap.innerHTML = `<div class="empty-text">알러지 없음</div>`;
        return;
      }

      codes.forEach(code => {
        const name = ALLERGY_MAP[code] ?? `알수없음(${code})`;
        const tag = document.createElement("div");
        tag.className = "tag";
        tag.innerHTML = `
          <span>${code}. ${name}</span>
          <button type="button" onclick="removeMenuItemAllergy(${index}, ${code})">×</button>
        `;
        wrap.appendChild(tag);
      });
    }

    function renderMenuAddItems() {
      const list = document.getElementById("menuAddList");
      if (!list) return;

      list.innerHTML = "";

      if (!addMenuItems || addMenuItems.length === 0) {
        list.innerHTML = `<div class="empty-text">메뉴가 아직 없어. "메뉴 추가" 버튼을 눌러줘.</div>`;
        return;
      }

      addMenuItems.forEach((item, index) => {
        const wrapper = document.createElement("div");
        wrapper.className = "menu-edit-item";
        const nameValue = (item.menu_name || "").replace(/"/g, "&quot;");

        wrapper.innerHTML = `
          <div class="form-group">
            <label>메뉴명</label>
            <input type="text" class="add-menu-name" data-index="${index}" value="${nameValue}" placeholder="예: 돈가스">
          </div>
          <div class="form-group">
            <label>알러지 선택</label>
            <div class="allergy-row">
              <select id="menuAddAllergySelect_${index}">
                <option value="">알러지를 선택해줘</option>
                ${Object.entries(ALLERGY_MAP).map(([code, name]) => `<option value="${code}">${code}. ${name}</option>`).join("")}
              </select>
              <button type="button" class="secondary-btn" onclick="addMenuItemAllergy(${index})">알러지 추가</button>
              <button type="button" class="danger-btn" onclick="removeMenuItem(${index})">메뉴 삭제</button>
            </div>
            <div id="menuAddAllergyList_${index}" class="tag-wrap"></div>
          </div>
        `;

        list.appendChild(wrapper);
        renderMenuAddAllergyTags(index);
      });
    }

    function addMenuItem() {
      addMenuItems.push({
        menu_name: "",
        allergyCodes: []
      });
      renderMenuAddItems();
    }

    function removeMenuItem(index) {
      if (index < 0 || index >= addMenuItems.length) return;
      addMenuItems.splice(index, 1);
      renderMenuAddItems();
    }

    function addMenuItemAllergy(index) {
      const select = document.getElementById(`menuAddAllergySelect_${index}`);
      if (!select) return;

      const value = select.value;
      if (!value) return;

      const num = Number(value);
      if (!num) return;

      if (!addMenuItems[index]) return;
      if (!addMenuItems[index].allergyCodes) {
        addMenuItems[index].allergyCodes = [];
      }

      if (!addMenuItems[index].allergyCodes.includes(num)) {
        addMenuItems[index].allergyCodes.push(num);
        addMenuItems[index].allergyCodes.sort((a, b) => a - b);
      }

      renderMenuAddAllergyTags(index);
      select.value = "";
    }

    function removeMenuItemAllergy(index, code) {
      if (!addMenuItems[index] || !addMenuItems[index].allergyCodes) return;
      addMenuItems[index].allergyCodes = addMenuItems[index].allergyCodes.filter(x => x !== code);
      renderMenuAddAllergyTags(index);
    }

    function addEditMenuItem() {
      editMenuItems.push({
        menu_name: "",
        allergyCodes: []
      });
      renderEditMenuItems();
    }

    function removeEditMenuItem(index) {
      if (index < 0 || index >= editMenuItems.length) return;
      editMenuItems.splice(index, 1);
      renderEditMenuItems();
    }

    function addEditMenuAllergy(index) {
      const select = document.getElementById(`editMenuAllergySelect_${index}`);
      if (!select) return;

      const value = select.value;
      if (!value) return;

      const num = Number(value);
      if (!num) return;

      if (!editMenuItems[index]) return;
      if (!editMenuItems[index].allergyCodes) {
        editMenuItems[index].allergyCodes = [];
      }

      if (!editMenuItems[index].allergyCodes.includes(num)) {
        editMenuItems[index].allergyCodes.push(num);
        editMenuItems[index].allergyCodes.sort((a, b) => a - b);
      }

      renderEditMenuAllergyTags(index);
      select.value = "";
    }

    function removeEditMenuAllergy(index, code) {
      if (!editMenuItems[index] || !editMenuItems[index].allergyCodes) return;
      editMenuItems[index].allergyCodes = editMenuItems[index].allergyCodes.filter(x => x !== code);
      renderEditMenuAllergyTags(index);
    }

    function resetMenuForm() {
      addMenuItems = [];
      renderMenuAddItems();

      const today = new Date().toISOString().split("T")[0];
      const dateEl = document.getElementById("menuDate");
      if (dateEl) dateEl.value = today;

      const createdByEl = document.getElementById("menuCreatedBy");
      if (createdByEl) createdByEl.value = DEFAULT_ADMIN_NAME;
    }

    function resetStudentForm() {
      selectedStudentAllergies = [];
      renderStudentAllergyState(false);

      const rfidEl = document.getElementById("studentRfidId");
      if (rfidEl) rfidEl.value = "";

      const nameEl = document.getElementById("studentName");
      if (nameEl) nameEl.value = "";

      const gradeEl = document.getElementById("studentGrade");
      if (gradeEl) gradeEl.value = "";

      const classEl = document.getElementById("studentClassNo");
      if (classEl) classEl.value = "1";

      const seqEl = document.getElementById("studentSeq");
      if (seqEl) seqEl.value = "";

      const noEl = document.getElementById("studentNoPreview");
      if (noEl) noEl.value = "";

      const selectEl = document.getElementById("studentAllergySelect");
      if (selectEl) selectEl.value = "";
    }

    function openMyAccountModal() {
      document.getElementById("myAccountModal")?.classList.remove("hidden");
    }

    function closeMyAccountModal() {
      document.getElementById("myAccountModal")?.classList.add("hidden");
      const pwEl = document.getElementById("myAccountPw");
      if (pwEl) pwEl.value = "";
    }

    async function saveMyAccount() {
      const newId = document.getElementById("myAccountId")?.value.trim() || "";
      const newName = document.getElementById("myAccountName")?.value.trim() || "";
      const newPw = document.getElementById("myAccountPw")?.value.trim() || "";

      const res = await fetch("/api/my-account/update", {
        method: "POST",
        headers: {"Content-Type": "application/json"},
        body: JSON.stringify({
          new_id: newId,
          new_name: newName,
          new_pw: newPw
        })
      });

      const data = await res.json();
      if (!data.ok) {
        alert(data.error || "계정 수정 실패");
        return;
      }

      alert("내 계정이 수정됐어.");
      window.location.reload();
    }

    async function resetMyPasswordToDefault() {
      if (!confirm("비밀번호를 기본번호 1234로 바꿀까?")) return;

      const res = await fetch("/api/my-account/reset-password", {
        method: "POST",
        headers: {"Content-Type": "application/json"}
      });

      const data = await res.json();
      if (!data.ok) {
        alert(data.error || "비밀번호 초기화 실패");
        return;
      }

      alert("비밀번호가 1234로 변경됐어.");
      window.location.reload();
    }

    function openMenuModal() {
      resetMenuForm();
      document.getElementById("menuModal")?.classList.remove("hidden");
    }

    function closeMenuModal() {
      document.getElementById("menuModal")?.classList.add("hidden");
      resetMenuForm();
    }

    async function openMenuEditModal(rowIndexes) {
      const res = await fetch("/api/menu/group-detail", {
        method: "POST",
        headers: {"Content-Type": "application/json"},
        body: JSON.stringify({ row_indexes: rowIndexes })
      });

      const data = await res.json();
      if (!data.ok) {
        alert(data.error || "급식 정보를 불러오지 못했어");
        return;
      }

      document.getElementById("editMenuRowIndexes").value = JSON.stringify(rowIndexes);
      document.getElementById("editMenuDateOriginal").value = data.date;

      const formattedDate = data.date.length === 8
        ? `${data.date.slice(0,4)}-${data.date.slice(4,6)}-${data.date.slice(6,8)}`
        : data.date;

      document.getElementById("editMenuDate").value = formattedDate;
      const items = data.items || [];
      editMenuItems = items.map(item => {
        const name = (item.menu_name || "").trim();
        const codesRaw = item.allergy_codes || [];
        const cleanedCodes = [];
        for (const c of codesRaw) {
          const n = Number(c);
          if (n >= 1 && n <= 19 && !cleanedCodes.includes(n)) {
            cleanedCodes.push(n);
          }
        }
        cleanedCodes.sort((a, b) => a - b);
        return {
          menu_name: name,
          allergyCodes: cleanedCodes
        };
      });
      if (!editMenuItems.length && Array.isArray(data.menus)) {
        editMenuItems = data.menus.map(name => ({
          menu_name: String(name || "").trim(),
          allergyCodes: []
        }));
      }
      renderEditMenuItems();
      document.getElementById("menuEditModal")?.classList.remove("hidden");
    }

    function closeMenuEditModal() {
      document.getElementById("menuEditModal")?.classList.add("hidden");
      document.getElementById("editMenuRowIndexes").value = "";
      document.getElementById("editMenuDateOriginal").value = "";
      document.getElementById("editMenuDate").value = "";
      editMenuItems = [];
      const list = document.getElementById("editMenuList");
      if (list) list.innerHTML = "";
    }

    async function saveMenuEdit() {
      const originalDate = document.getElementById("editMenuDateOriginal").value;
      const newDate = (document.getElementById("editMenuDate").value || "").replaceAll("-", "");

      const nameInputs = document.querySelectorAll(".edit-menu-name");
      nameInputs.forEach(input => {
        const idx = Number(input.dataset.index);
        if (!Number.isNaN(idx) && editMenuItems[idx]) {
          editMenuItems[idx].menu_name = (input.value || "").trim();
        }
      });

      const cleanedItems = [];
      for (const item of editMenuItems) {
        const name = (item.menu_name || "").trim();
        if (!name) continue;

        const codes = Array.isArray(item.allergyCodes) ? item.allergyCodes : [];
        const cleanedCodes = [];
        for (const c of codes) {
          const n = Number(c);
          if (n >= 1 && n <= 19 && !cleanedCodes.includes(n)) {
            cleanedCodes.push(n);
          }
        }
        cleanedCodes.sort((a, b) => a - b);

        cleanedItems.push({
          menu_name: name,
          allergy_codes: cleanedCodes
        });
      }

      if (!cleanedItems.length) {
        alert("메뉴를 1개 이상 입력해야 해.");
        return;
      }

      const rowIndexes = JSON.parse(document.getElementById("editMenuRowIndexes").value || "[]");

      const res = await fetch("/api/menu/group-update", {
        method: "POST",
        headers: {"Content-Type": "application/json"},
        body: JSON.stringify({
          original_date: originalDate,
          new_date: newDate,
          row_indexes: rowIndexes,
          menu_items: cleanedItems
        })
      });

      const data = await res.json();
      if (!data.ok) {
        alert(data.error || "급식 수정 실패");
        return;
      }

      alert("급식 식단표가 수정됐어.");
      window.location.reload();
    }

    function openStudentModal() {
      document.getElementById("studentModal")?.classList.remove("hidden");
      updateStudentNumberPreview();
    }

    function closeStudentModal() {
      document.getElementById("studentModal")?.classList.add("hidden");
      resetStudentForm();
    }

    function openStudentEditModal(rowIndex, rfidId, name, studentId, allergyCodesText) {
      const modal = document.getElementById("studentEditModal");
      if (!modal) return;

      const parts = splitStudentNumber(studentId);

      const normalizeCodesForEditModal = (value) => {
        if (Array.isArray(value)) {
          return value
            .map(v => String(v ?? "").trim())
            .filter(v => /^\d+$/.test(v) && v !== "0");
        }

        return String(value || "")
          .split(/[\s,]+/)
          .map(v => v.trim())
          .filter(v => /^\d+$/.test(v) && v !== "0");
      };

      document.getElementById("editStudentRowIndex").value = rowIndex;
      document.getElementById("editStudentRfidId").value = rfidId || "";
      document.getElementById("editStudentName").value = name || "";
      document.getElementById("editStudentGrade").value = parts.grade || "";
      document.getElementById("editStudentClassNo").value = parts.classNo || "";
      document.getElementById("editStudentSeq").value = parts.seq || "";
      document.getElementById("editStudentPw").value = "";

      editStudentAllergies = normalizeCodesForEditModal(allergyCodesText);
      renderEditStudentTags(editStudentAllergies.length === 0);

      updateEditStudentNumberPreview();
      modal.classList.remove("hidden");
    }

    async function loadLatestRfid() {
      const res = await fetch("/api/latest-rfid");
      const data = await res.json();
      if (!data.ok) {
        throw new Error(data.error || "최근 RFID를 불러오지 못했어.");
      }
      if (!data.uid) {
        throw new Error("최근에 태그된 RFID가 아직 없어. 카드를 먼저 태그해줘.");
      }
      return data;
    }

    async function fillLatestRfidToAdd() {
      try {
        const data = await loadLatestRfid();
        document.getElementById("studentRfidId").value = data.uid || "";
        const hint = document.getElementById("latestRfidAddHint");
        if (hint) hint.innerText = `최근 태그: ${data.uid} (${data.scanned_at || "시간 정보 없음"})`;
      } catch (err) {
        alert(err.message || "최근 RFID를 불러오지 못했어.");
      }
    }

    async function fillLatestRfidToEdit() {
      try {
        const data = await loadLatestRfid();
        document.getElementById("editStudentRfidId").value = data.uid || "";
        const hint = document.getElementById("latestRfidEditHint");
        if (hint) hint.innerText = `최근 태그: ${data.uid} (${data.scanned_at || "시간 정보 없음"})`;
      } catch (err) {
        alert(err.message || "최근 RFID를 불러오지 못했어.");
      }
    }

    function closeStudentEditModal() {
      document.getElementById("studentEditModal")?.classList.add("hidden");
      editStudentAllergies = [];
    }

    function openStudentSheetModal() {
      document.getElementById("studentSheetModal")?.classList.remove("hidden");
      clearStudentSheetPreview();
    }

    function closeStudentSheetModal() {
      document.getElementById("studentSheetModal")?.classList.add("hidden");
      document.getElementById("studentSheetForm")?.reset();
      clearStudentSheetPreview();
    }

    function studentSheetEscape(value) {
      return String(value ?? "")
        .replace(/&/g, "&amp;")
        .replace(/</g, "&lt;")
        .replace(/>/g, "&gt;")
        .replace(/"/g, "&quot;")
        .replace(/'/g, "&#039;");
    }

    function setStudentSheetResult(message, isHidden = false) {
      const result = document.getElementById("studentSheetResult");
      if (!result) return;
      result.classList.toggle("hidden", Boolean(isHidden));
      result.innerText = message || "";
    }

    function clearStudentSheetPreview() {
      const preview = document.getElementById("studentSheetPreviewWrap");
      if (preview) {
        preview.classList.add("hidden");
        preview.innerHTML = "";
      }
      setStudentSheetResult("", true);
      document.getElementById("studentSheetPreviewBtn")?.classList.remove("hidden");
      document.getElementById("studentSheetConfirmBtn")?.classList.add("hidden");
      document.getElementById("studentSheetClearPreviewBtn")?.classList.add("hidden");
    }

    function removeStudentSheetPreviewRow(button) {
      button.closest("tr")?.remove();
      const rows = document.querySelectorAll(".student-sheet-preview-row");
      if (!rows.length) {
        setStudentSheetResult("미리보기 행이 없어. 다시 선택해줘.");
      }
    }

    function renderStudentSheetPreview(rows, skipped = [], warnings = []) {
      const preview = document.getElementById("studentSheetPreviewWrap");
      if (!preview) return;
      const safeRows = Array.isArray(rows) ? rows : [];
      const skippedHtml = skipped.length
        ? `<div class="empty-text" style="margin:10px 0;">건너뛸 행 ${skipped.length}개: ${studentSheetEscape(skipped.slice(0, 5).map(x => `${x.name || x.student_number || x.index}: ${x.reason}`).join(" / "))}${skipped.length > 5 ? " ..." : ""}</div>`
        : "";
      const warningHtml = warnings.length
        ? `<div class="empty-text" style="margin:10px 0;">참고: ${studentSheetEscape(warnings.join(" / "))}</div>`
        : "";
      preview.innerHTML = `
        <div class="empty-text" style="margin-bottom:8px;">값을 수정하거나 필요 없는 행은 삭제한 뒤 확인 후 추가를 눌러줘.</div>
        ${warningHtml}
        ${skippedHtml}
        <table style="table-layout:fixed;">
          <thead>
            <tr>
              <th style="width:150px;">RFID ID</th>
              <th style="width:150px;">이름</th>
              <th style="width:110px;">학번</th>
              <th>알러지 번호</th>
              <th style="width:88px;">작업</th>
            </tr>
          </thead>
          <tbody>
            ${safeRows.map((row) => `
              <tr class="student-sheet-preview-row">
                <td><input class="sheet-rfid" type="text" value="${studentSheetEscape(row.rfid_id || "")}" placeholder="비워도 됨"></td>
                <td><input class="sheet-name" type="text" value="${studentSheetEscape(row.name || "")}"></td>
                <td><input class="sheet-student-number" type="text" value="${studentSheetEscape(row.student_number || "")}"></td>
                <td><input class="sheet-allergy-codes" type="text" value="${studentSheetEscape(Array.isArray(row.allergy_codes) ? row.allergy_codes.join(",") : (row.allergy_codes || ""))}" placeholder="예: 1,6"></td>
                <td><button class="ghost-btn" type="button" onclick="removeStudentSheetPreviewRow(this)">삭제</button></td>
              </tr>
            `).join("")}
          </tbody>
        </table>
      `;
      preview.classList.remove("hidden");
      document.getElementById("studentSheetPreviewBtn")?.classList.add("hidden");
      document.getElementById("studentSheetConfirmBtn")?.classList.remove("hidden");
      document.getElementById("studentSheetClearPreviewBtn")?.classList.remove("hidden");
    }

    function collectStudentSheetPreviewRows() {
      return Array.from(document.querySelectorAll(".student-sheet-preview-row")).map(row => ({
        rfid_id: row.querySelector(".sheet-rfid")?.value.trim() || "",
        name: row.querySelector(".sheet-name")?.value.trim() || "",
        student_number: row.querySelector(".sheet-student-number")?.value.trim() || "",
        allergy_codes: row.querySelector(".sheet-allergy-codes")?.value.trim() || ""
      }));
    }

    async function confirmStudentSheetAdd() {
      const rows = collectStudentSheetPreviewRows();
      if (!rows.length) {
        setStudentSheetResult("추가할 행이 없어.");
        return;
      }
      setStudentSheetResult("확인한 학생을 추가하는 중...");
      try {
        const res = await fetch("/api/student/sheet-confirm", {
          method: "POST",
          headers: {"Content-Type": "application/json"},
          body: JSON.stringify({ rows })
        });
        const contentType = res.headers.get("content-type") || "";
        const data = contentType.includes("application/json")
          ? await res.json()
          : { ok: false, error: (await res.text()).replace(/<[^>]*>/g, " ").replace(/\s+/g, " ").trim() || "서버 오류가 발생했어." };
        if (!data.ok) throw new Error(data.error || "학생 추가 실패");
        alert(`학생 ${data.added_count || 0}명이 추가됐어.` + (data.skipped_count ? ` 건너뜀 ${data.skipped_count}건.` : ""));
        window.location.reload();
      } catch (err) {
        setStudentSheetResult(err.message || "학생 추가에 실패했어.");
      }
    }

    function filterStudentRows() {
      const query = (document.getElementById("studentSearchInput")?.value || "").trim().toLowerCase();
      document.querySelectorAll(".student-data-row").forEach(row => {
        const haystack = row.dataset.search || "";
        row.style.display = !query || haystack.includes(query) ? "" : "none";
      });
    }

    function addSelectedMenuAllergy() {
      const select = document.getElementById("menuAllergySelect");
      const value = Number(select.value);
      if (!value) return;

      if (!selectedMenuAllergies.includes(value)) {
        selectedMenuAllergies.push(value);
        selectedMenuAllergies.sort((a, b) => a - b);
      }

      renderTags("selectedMenuAllergyList", selectedMenuAllergies, "removeMenuAllergy");
      select.value = "";
    }

    function removeMenuAllergy(code) {
      selectedMenuAllergies = selectedMenuAllergies.filter(x => x !== code);
      renderTags("selectedMenuAllergyList", selectedMenuAllergies, "removeMenuAllergy");
    }

    function addSelectedStudentAllergy() {
      const select = document.getElementById("studentAllergySelect");
      const value = select.value;

      if (!value) return;

      if (value === "none") {
        selectedStudentAllergies = [];
        renderStudentAllergyState(true);
        select.value = "";
        return;
      }

      const numValue = Number(value);
      if (!numValue) return;

      if (!selectedStudentAllergies.includes(numValue)) {
        selectedStudentAllergies.push(numValue);
        selectedStudentAllergies.sort((a, b) => a - b);
      }

      renderStudentAllergyState(false);
      select.value = "";
    }

    function removeStudentAllergy(code) {
      selectedStudentAllergies = selectedStudentAllergies.filter(x => x !== code);
      renderStudentAllergyState(false);
    }

    function addEditStudentAllergy() {
      const select = document.getElementById("editStudentAllergySelect");
      const value = select.value;

      if (!value) return;

      if (value === "none") {
        editStudentAllergies = [];
        renderEditStudentTags(true);
        select.value = "";
        return;
      }

      const num = Number(value);
      if (!num) return;

      if (!editStudentAllergies.includes(num)) {
        editStudentAllergies.push(num);
        editStudentAllergies.sort((a, b) => a - b);
      }

      renderEditStudentTags(false);
      select.value = "";
    }

    function removeEditStudentAllergy(code) {
      editStudentAllergies = editStudentAllergies.filter(x => x !== code);
      renderEditStudentTags(false);
    }

    async function submitMenu() {
      const menuDate = document.getElementById("menuDate")?.value || "";
      const createdBy = document.getElementById("menuCreatedBy")?.value.trim() || "";

      const nameInputs = document.querySelectorAll(".add-menu-name");
      nameInputs.forEach(input => {
        const idx = Number(input.dataset.index);
        if (!Number.isNaN(idx) && addMenuItems[idx]) {
          addMenuItems[idx].menu_name = (input.value || "").trim();
        }
      });

      const cleanedItems = [];
      for (const item of addMenuItems) {
        const name = (item.menu_name || "").trim();
        if (!name) continue;

        const codes = Array.isArray(item.allergyCodes) ? item.allergyCodes : [];
        const cleanedCodes = [];
        for (const c of codes) {
          const n = Number(c);
          if (n >= 1 && n <= 19 && !cleanedCodes.includes(n)) {
            cleanedCodes.push(n);
          }
        }
        cleanedCodes.sort((a, b) => a - b);

        cleanedItems.push({
          menu_name: name,
          allergy_codes: cleanedCodes
        });
      }

      if (!menuDate || !createdBy) {
        alert("날짜와 추가자를 입력해야 해.");
        return;
      }

      if (!cleanedItems.length) {
        alert("메뉴를 1개 이상 입력해야 해.");
        return;
      }

      for (const item of cleanedItems) {
        const res = await fetch("/api/menu", {
          method: "POST",
          headers: {"Content-Type": "application/json"},
          body: JSON.stringify({
            date: menuDate,
            menu_name: item.menu_name,
            created_by: createdBy,
            allergy_codes: item.allergy_codes
          })
        });

        const data = await res.json();
        if (!data.ok) {
          alert(data.error || "급식 추가 실패");
          return;
        }
      }

      alert("급식이 추가됐어.");
      window.location.reload();
    }

    async function submitStudent() {
      const rfidId = document.getElementById("studentRfidId")?.value.trim() || "";
      const studentName = document.getElementById("studentName")?.value.trim() || "";
      const studentNo = updateStudentNumberPreview();
      const studentGrade = document.getElementById("studentGrade")?.value.trim() || "";
      const studentClassNo = document.getElementById("studentClassNo")?.value.trim() || "1";
      const studentSeq = document.getElementById("studentSeq")?.value.trim() || "";

      const res = await fetch("/api/student", {
        method: "POST",
        headers: {"Content-Type": "application/json"},
        body: JSON.stringify({
          rfid_id: rfidId,
          name: studentName,
          student_number: studentNo,
          grade: studentGrade,
          class_no: studentClassNo,
          student_seq: studentSeq,
          password: DEFAULT_STUDENT_PASSWORD,
          allergy_codes: selectedStudentAllergies
        })
      });

      const data = await res.json();
      if (!data.ok) {
        alert(data.error || "학생 추가 실패");
        return;
      }

      alert("학생과 로그인 계정이 추가됐어. 초기 비밀번호는 1234야.");
      window.location.reload();
    }

    async function saveStudentEdit() {
      const rowIndex = Number(document.getElementById("editStudentRowIndex").value);
      const newRfidId = document.getElementById("editStudentRfidId").value.trim();
      const newName = document.getElementById("editStudentName").value.trim();
      const newId = updateEditStudentNumberPreview();
      const newPw = document.getElementById("editStudentPw").value.trim();
      const grade = document.getElementById("editStudentGrade").value.trim();
      const classNo = document.getElementById("editStudentClassNo").value.trim();
      const studentSeq = document.getElementById("editStudentSeq").value.trim();

      const res = await fetch("/api/student/update", {
        method: "POST",
        headers: {"Content-Type": "application/json"},
        body: JSON.stringify({
          row_index: rowIndex,
          new_rfid_id: newRfidId,
          new_name: newName,
          new_id: newId,
          grade: grade,
          class_no: classNo,
          student_seq: studentSeq,
          new_pw: newPw,
          allergy_codes: editStudentAllergies
        })
      });

      const data = await res.json();
      if (!data.ok) {
        alert(data.error || "학생 수정 실패");
        return;
      }

      alert("학생 정보가 수정됐어.");
      window.location.reload();
    }

    async function resetStudentPasswordToDefault() {
      const rowIndex = Number(document.getElementById("editStudentRowIndex").value);

      if (!confirm("이 학생 비밀번호를 1234로 바꿀까?")) return;

      const res = await fetch("/api/student/reset-password", {
        method: "POST",
        headers: {"Content-Type": "application/json"},
        body: JSON.stringify({ row_index: rowIndex })
      });

      const data = await res.json();
      if (!data.ok) {
        alert(data.error || "비밀번호 초기화 실패");
        return;
      }

      alert("학생 비밀번호가 1234로 변경됐어.");
      window.location.reload();
    }

    function toggleStudentSelectionMode() {
      studentSelectionMode = !studentSelectionMode;

      const header = document.getElementById("studentCheckboxHeader");
      const cells = document.querySelectorAll(".student-checkbox-cell");
      const deleteBtn = document.getElementById("deleteSelectedStudentsBtn");
      const toggleBtn = document.getElementById("toggleStudentDeleteModeBtn");

      if (studentSelectionMode) {
        header?.classList.remove("hidden");
        cells.forEach(el => el.classList.remove("hidden"));
        deleteBtn?.classList.remove("hidden");
        if (toggleBtn) toggleBtn.innerText = "취소";
      } else {
        header?.classList.add("hidden");
        cells.forEach(el => el.classList.add("hidden"));
        deleteBtn?.classList.add("hidden");
        document.querySelectorAll(".student-row-checkbox").forEach(cb => cb.checked = false);
        if (toggleBtn) toggleBtn.innerText = "삭제";
      }
    }

    function toggleMenuSelectionMode() {
      menuSelectionMode = !menuSelectionMode;

      const header = document.getElementById("menuCheckboxHeader");
      const cells = document.querySelectorAll(".menu-checkbox-cell");
      const deleteBtn = document.getElementById("deleteSelectedMenusBtn");
      const toggleBtn = document.getElementById("toggleMenuDeleteModeBtn");

      if (menuSelectionMode) {
        header?.classList.remove("hidden");
        cells.forEach(el => el.classList.remove("hidden"));
        deleteBtn?.classList.remove("hidden");
        if (toggleBtn) toggleBtn.innerText = "취소";
      } else {
        header?.classList.add("hidden");
        cells.forEach(el => el.classList.add("hidden"));
        deleteBtn?.classList.add("hidden");
        document.querySelectorAll(".menu-row-checkbox").forEach(cb => cb.checked = false);
        if (toggleBtn) toggleBtn.innerText = "삭제";
      }
    }

    async function deleteSelectedStudents() {
      const checked = Array.from(document.querySelectorAll(".student-row-checkbox:checked"))
        .map(cb => Number(cb.value));

      if (checked.length === 0) {
        alert("삭제할 학생을 먼저 선택해줘.");
        return;
      }

      if (!confirm("선택한 학생을 휴지통으로 보낼까?")) return;

      const res = await fetch("/api/student/delete-selected", {
        method: "POST",
        headers: {"Content-Type": "application/json"},
        body: JSON.stringify({ row_indexes: checked })
      });

      const data = await res.json();
      if (!data.ok) {
        alert(data.error || "삭제 실패");
        return;
      }

      alert("선택한 학생을 휴지통으로 이동했어.");
      window.location.reload();
    }

    async function deleteSelectedMenus() {
      const checked = Array.from(document.querySelectorAll(".menu-row-checkbox:checked"))
        .map(cb => {
          try {
            return JSON.parse(cb.value);
          } catch {
            return [];
          }
        })
        .flat()
        .map(x => Number(x))
        .filter(x => x >= 2);

      const uniqueRows = [...new Set(checked)].sort((a, b) => a - b);

      if (uniqueRows.length === 0) {
        alert("삭제할 급식 식단표를 먼저 선택해줘.");
        return;
      }

      if (!confirm("선택한 급식 식단표를 휴지통으로 보낼까?")) return;

      const res = await fetch("/api/menu/delete-selected", {
        method: "POST",
        headers: {"Content-Type": "application/json"},
        body: JSON.stringify({ row_indexes: uniqueRows })
      });

      const data = await res.json();
      if (!data.ok) {
        alert(data.error || "삭제 실패");
        return;
      }

      alert("선택한 급식 식단표를 휴지통으로 이동했어.");
      window.location.reload();
    }

    async function restoreMenu(rowIndex) {
      if (!confirm("이 급식 메뉴를 복구할까?")) return;

      const res = await fetch("/api/trash/menu/restore", {
        method: "POST",
        headers: {"Content-Type": "application/json"},
        body: JSON.stringify({ row_index: rowIndex })
      });

      const data = await res.json();
      if (!data.ok) {
        alert(data.error || "복구 실패");
        return;
      }

      alert("복구됐어.");
      window.location.reload();
    }

    async function restoreStudent(rowIndex) {
      if (!confirm("이 학생 정보를 복구할까?")) return;

      const res = await fetch("/api/trash/student/restore", {
        method: "POST",
        headers: {"Content-Type": "application/json"},
        body: JSON.stringify({ row_index: rowIndex })
      });

      const data = await res.json();
      if (!data.ok) {
        alert(data.error || "복구 실패");
        return;
      }

      alert("복구됐어.");
      window.location.reload();
    }

    async function hardDeleteMenu(rowIndex) {
      if (!confirm("휴지통에서 완전히 삭제할까? 되돌릴 수 없어.")) return;

      const res = await fetch("/api/trash/menu/hard-delete", {
        method: "POST",
        headers: {"Content-Type": "application/json"},
        body: JSON.stringify({ row_index: rowIndex })
      });

      const data = await res.json();
      if (!data.ok) {
        alert(data.error || "완전 삭제 실패");
        return;
      }

      alert("완전히 삭제됐어.");
      window.location.reload();
    }

    async function hardDeleteStudent(rowIndex) {
      if (!confirm("휴지통에서 완전히 삭제할까? 되돌릴 수 없어.")) return;

      const res = await fetch("/api/trash/student/hard-delete", {
        method: "POST",
        headers: {"Content-Type": "application/json"},
        body: JSON.stringify({ row_index: rowIndex })
      });

      const data = await res.json();
      if (!data.ok) {
        alert(data.error || "완전 삭제 실패");
        return;
      }

      alert("완전히 삭제됐어.");
      window.location.reload();
    }

    function aiEscape(value) {
      return String(value ?? "")
        .replaceAll("&", "&amp;")
        .replaceAll("<", "&lt;")
        .replaceAll(">", "&gt;")
        .replaceAll('"', "&quot;")
        .replaceAll("'", "&#039;");
    }

    function aiDateValue() {
      return document.getElementById("aiDateSelect")?.value || "";
    }

    function aiStudentValue() {
      return document.getElementById("aiStudentSelect")?.value || "";
    }

    function aiContextValue() {
      return document.getElementById("aiContextInput")?.value || "";
    }

    function setAiStatus(message, isError = false) {
      const box = document.getElementById("aiStatusBox");
      if (!box) return;
      box.textContent = message;
      box.className = isError ? "ai-status error" : "ai-status";
    }

    function hideAiStatus() {
      const box = document.getElementById("aiStatusBox");
      if (!box) return;
      box.textContent = "";
      box.className = "ai-status hidden";
    }

    async function postAiJson(url, payload) {
      const res = await fetch(url, {
        method: "POST",
        headers: {"Content-Type": "application/json"},
        body: JSON.stringify(payload || {})
      });
      const data = await res.json();
      if (!res.ok || !data.ok) {
        throw new Error(data.error || "AI 요청에 실패했어.");
      }
      return data;
    }

    function aiChips(values, kind = "") {
      const list = Array.isArray(values) ? values.filter(Boolean) : [];
      if (!list.length) return `<span class="ai-chip">없음</span>`;
      return list.map(item => `<span class="ai-chip ${kind}">${aiEscape(item)}</span>`).join("");
    }

    function aiList(values) {
      const list = Array.isArray(values) ? values.filter(Boolean) : [];
      if (!list.length) return `<span class="empty-text">표시할 항목이 없어.</span>`;
      return `<ul style="margin:0; padding-left:18px; display:grid; gap:7px;">${list.map(item => `<li style="white-space:normal; word-break:keep-all;">${aiEscape(item)}</li>`).join("")}</ul>`;
    }

    function aiEmptyTable(message) {
      return `
        <table>
          <tbody>
            <tr>
              <td class="empty-text" style="padding:22px; text-align:center;">${aiEscape(message)}</td>
            </tr>
          </tbody>
        </table>
      `;
    }

    function renderAiSafety(data) {
      const plan = data.plan || {};
      const student = data.student || {};
      const alternatives = Array.isArray(plan.alternative_meals) ? plan.alternative_meals : [];
      const messageItems = [
        ["보호자 안내", plan.guardian_message],
        ["보건실 전달", plan.nurse_message],
        ["급식실 전달", plan.kitchen_message],
      ];

      const safetyTarget = document.getElementById("aiSafetyResult");
      if (safetyTarget) {
        const alternativeHtml = alternatives.length
          ? alternatives.map(item => `
              <div style="display:grid; gap:6px;">
                <div style="font-weight:900;">${aiEscape(item.name)}</div>
                <div>${aiEscape(item.components)}</div>
                <div class="empty-text">${aiEscape(item.preparation_notes)}</div>
                <div>${aiChips(item.avoids || [], "warn")}</div>
                <div>${aiEscape(item.reason)}</div>
                <div><strong>확인:</strong> ${aiEscape(item.staff_check)}</div>
              </div>
            `).join("<hr style=\"border:0; border-top:1px solid var(--line-soft); margin:10px 0;\">")
          : `<span class="empty-text">대체급식 후보가 없어.</span>`;

        safetyTarget.innerHTML = `
          <table style="table-layout:fixed;">
            <thead>
              <tr>
                <th style="width:180px;">항목</th>
                <th>내용</th>
              </tr>
            </thead>
            <tbody>
              <tr class="${data.unsafe_menus?.length ? "row-warn" : ""}">
                <td>학생</td>
                <td>${aiEscape(student.name || "학생")} · ${aiEscape(student.student_number || "")}</td>
              </tr>
              <tr>
                <td>위험 요약</td>
                <td style="white-space:normal; word-break:keep-all;">${aiEscape(plan.risk_summary || "")}</td>
              </tr>
              <tr>
                <td>위험도</td>
                <td><span class="ai-chip ${plan.risk_level === "높음" ? "warn" : "ok"}">위험도 ${aiEscape(plan.risk_level || "-")}</span></td>
              </tr>
              <tr>
                <td>일치 알레르기</td>
                <td>${aiChips(data.matched_allergies || [], "warn")}</td>
              </tr>
              <tr>
                <td>주의 메뉴</td>
                <td>${aiChips(data.unsafe_menus || [], "warn")}</td>
              </tr>
              <tr>
                <td>대체급식 후보</td>
                <td style="white-space:normal; word-break:keep-all;">${alternativeHtml}</td>
              </tr>
              <tr>
                <td>예방 체크</td>
                <td>${aiList(plan.prevention_checklist || [])}</td>
              </tr>
              <tr>
                <td>관찰 포인트</td>
                <td>${aiList(plan.monitoring_points || [])}</td>
              </tr>
              ${messageItems.map(([title, text]) => `
                <tr>
                  <td>${aiEscape(title)}</td>
                  <td style="white-space:normal; word-break:keep-all;">${aiEscape(text || "")}</td>
                </tr>
              `).join("")}
            </tbody>
          </table>
        `;
      }

      const emergencyTarget = document.getElementById("aiEmergencyResult");
      const steps = Array.isArray(plan.emergency_steps) ? plan.emergency_steps : [];
      if (emergencyTarget) {
        emergencyTarget.innerHTML = steps.length
          ? `
            <table style="table-layout:fixed;">
              <thead>
                <tr>
                  <th style="width:70px;">순서</th>
                  <th style="width:180px;">단계</th>
                  <th>조치</th>
                  <th style="width:150px;">담당</th>
                  <th style="width:110px;">긴급도</th>
                </tr>
              </thead>
              <tbody>
                ${steps.map(step => `
                  <tr class="row-warn">
                    <td>${aiEscape(step.order || "")}</td>
                    <td>${aiEscape(step.title || "")}</td>
                    <td style="white-space:normal; word-break:keep-all;">${aiEscape(step.action || "")}</td>
                    <td>${aiEscape(step.owner || "")}</td>
                    <td><span class="ai-chip warn">${aiEscape(step.urgency || "")}</span></td>
                  </tr>
                `).join("")}
              </tbody>
            </table>
          `
          : aiEmptyTable("대응 단계가 없어.");
      }
    }

    function renderAiBrief(data) {
      const brief = data.brief || {};
      const watch = Array.isArray(brief.watch_students) ? brief.watch_students : [];
      const target = document.getElementById("aiBriefResult");
      if (!target) return;

      target.innerHTML = `
        <table style="table-layout:fixed;">
          <thead>
            <tr>
              <th style="width:160px;">구분</th>
              <th>내용</th>
            </tr>
          </thead>
          <tbody>
            <tr class="${watch.length ? "row-warn" : ""}">
              <td>${aiEscape(data.date)} 요약</td>
              <td style="white-space:normal; word-break:keep-all;">${aiEscape(brief.summary || "")}</td>
            </tr>
            <tr>
              <td>집계</td>
              <td>
                <span class="ai-chip warn">주의 ${aiEscape(data.risk_count || 0)}명</span>
                <span class="ai-chip">메뉴 ${aiEscape((data.menus || []).length)}개</span>
              </td>
            </tr>
            <tr>
              <td>우선 조치</td>
              <td>${aiList(brief.priority_actions || [])}</td>
            </tr>
            <tr>
              <td>급식실 메모</td>
              <td>${aiList(brief.kitchen_notes || [])}</td>
            </tr>
          </tbody>
        </table>
        <table style="table-layout:fixed; margin-top:0;">
          <thead>
            <tr>
              <th style="width:170px;">학생</th>
              <th style="width:210px;">알레르기</th>
              <th style="width:210px;">주의 메뉴</th>
              <th>조치</th>
            </tr>
          </thead>
          <tbody>
            ${watch.length ? watch.map(item => `
              <tr>
                <td>${aiEscape(item.student_name || "")}<br><span class="empty-text">${aiEscape(item.student_number || "")}</span></td>
                <td>${aiChips(item.allergies || [], "warn")}</td>
                <td>${aiChips(item.unsafe_menus || [], "warn")}</td>
                <td style="white-space:normal; word-break:keep-all;">${aiEscape(item.action || "")}</td>
              </tr>
            `).join("") : `<tr><td colspan="4" class="empty-text" style="padding:22px; text-align:center;">직접 충돌 학생이 없어.</td></tr>`}
          </tbody>
        </table>
      `;
    }

    function renderAiReview(data) {
      const review = data.review || {};
      const rows = Array.isArray(review.reviewed_menus) ? review.reviewed_menus : [];
      const target = document.getElementById("aiReviewResult");
      if (!target) return;

      target.innerHTML = `
        <table style="table-layout:fixed;">
          <thead>
            <tr>
              <th style="width:220px;">메뉴</th>
              <th style="width:230px;">현재 코드</th>
              <th style="width:230px;">누락 후보</th>
              <th>검토 사유</th>
            </tr>
          </thead>
          <tbody>
            ${rows.length ? rows.map(item => `
              <tr class="${(item.suspected_missing_codes || []).length ? "row-warn" : ""}">
                <td>${aiEscape(item.menu_name || "")}</td>
                <td>${aiChips((item.current_codes || []).map(code => `${code}. ${ALLERGY_MAP[code] || "알수없음"}`))}</td>
                <td>${aiChips((item.suspected_missing_codes || []).map(code => `${code}. ${ALLERGY_MAP[code] || "알수없음"}`), "warn")}</td>
                <td style="white-space:normal; word-break:keep-all;">${aiEscape(item.review_reason || "")}</td>
              </tr>
            `).join("") : `<tr><td colspan="4" class="empty-text" style="padding:22px; text-align:center;">점검 결과가 없어.</td></tr>`}
          </tbody>
        </table>
        <table style="table-layout:fixed;">
          <tbody>
            <tr class="${(review.warnings || []).length ? "row-warn" : ""}">
              <td style="width:160px;">점검 요약</td>
              <td>
                <span class="ai-chip">점검 ${rows.length}개</span>
                <span class="ai-chip ${(review.warnings || []).length ? "warn" : "ok"}">검토 ${(review.warnings || []).length}개</span>
              </td>
            </tr>
            <tr>
              <td>경고</td>
              <td>${aiList(review.warnings || [])}</td>
            </tr>
          </tbody>
        </table>
      `;
    }

    async function loadAiSafetyPlan() {
      const studentNumber = aiStudentValue();
      if (!studentNumber) {
        setAiStatus("학생을 먼저 선택해줘.", true);
        return;
      }
      setAiStatus("학생 안전계획, 오늘 브리핑, 메뉴 코드점검 생성 중...");
      try {
        const date = aiDateValue();
        const context = aiContextValue();
        const results = await Promise.allSettled([
          postAiJson("{{ url_for('api_ai_safety_plan') }}", {
            date,
            student_number: studentNumber,
            context
          }),
          postAiJson("{{ url_for('api_ai_daily_brief') }}", { date }),
          postAiJson("{{ url_for('api_ai_menu_review') }}", { date })
        ]);

        const errors = [];
        if (results[0].status === "fulfilled") renderAiSafety(results[0].value);
        else errors.push(`학생 안전계획: ${results[0].reason?.message || results[0].reason || "실패"}`);

        if (results[1].status === "fulfilled") renderAiBrief(results[1].value);
        else errors.push(`오늘 브리핑: ${results[1].reason?.message || results[1].reason || "실패"}`);

        if (results[2].status === "fulfilled") renderAiReview(results[2].value);
        else errors.push(`메뉴 코드점검: ${results[2].reason?.message || results[2].reason || "실패"}`);

        if (errors.length) {
          setAiStatus(errors.join(" / "), true);
        } else {
          hideAiStatus();
        }
      } catch (err) {
        setAiStatus(err.message, true);
      }
    }

    async function loadAiDailyBrief() {
      setAiStatus("오늘 브리핑 생성 중...");
      try {
        const data = await postAiJson("{{ url_for('api_ai_daily_brief') }}", { date: aiDateValue() });
        renderAiBrief(data);
        hideAiStatus();
      } catch (err) {
        setAiStatus(err.message, true);
      }
    }

    async function loadAiMenuReview() {
      setAiStatus("메뉴 알레르기 코드 점검 중...");
      try {
        const data = await postAiJson("{{ url_for('api_ai_menu_review') }}", { date: aiDateValue() });
        renderAiReview(data);
        hideAiStatus();
      } catch (err) {
        setAiStatus(err.message, true);
      }
    }

    function clearAiResults() {
      hideAiStatus();
      const safety = document.getElementById("aiSafetyResult");
      const emergency = document.getElementById("aiEmergencyResult");
      const brief = document.getElementById("aiBriefResult");
      const review = document.getElementById("aiReviewResult");
      if (safety) safety.innerHTML = aiEmptyTable("학생을 선택하고 학생 안전계획을 실행해줘.");
      if (emergency) emergency.innerHTML = aiEmptyTable("안전계획 결과와 함께 표시됩니다.");
      if (brief) brief.innerHTML = aiEmptyTable("오늘 브리핑을 실행해줘.");
      if (review) review.innerHTML = aiEmptyTable("메뉴 코드점검을 실행해줘.");
    }

    function renumberMealAiRows() {
      const rows = Array.from(document.querySelectorAll(".meal-ai-preview-row"));
      rows.forEach((row, index) => {
        row.querySelectorAll("input").forEach((input) => {
          if (!input.name) return;
          input.name = input.name.replace(/_(\d+)$/, `_${index}`);
        });
      });
      const rowCount = document.getElementById("mealAiRowCount");
      if (rowCount) rowCount.value = rows.length;
    }

    function addMealAiRow() {
      const createdBy = document.getElementById("mealAiCreatedByInput")?.value || "";
      const body = document.getElementById("mealAiRowPreviewBody");
      if (!body) return;

      const row = document.createElement("tr");
      row.className = "meal-ai-preview-row";
      row.innerHTML = `
        <td><input class="mini-input" type="text" name="row_date_0" placeholder="20260418" required></td>
        <td><input type="text" name="row_menu_name_0" required></td>
        <td><input type="text" name="row_allergy_codes_0" placeholder="1,2,5,6"></td>
        <td><input type="text" name="row_created_by_0" value="${aiEscape(createdBy)}"></td>
        <td><input type="text" name="row_created_at_0" value="" placeholder="저장 시 자동입력" readonly></td>
        <td><input class="mini-input" type="text" name="row_trash_0" value="0"></td>
        <td><input class="mini-input" type="number" name="row_confidence_0" value="0.5" step="0.01" min="0" max="1"></td>
        <td><label style="display:inline-flex; gap:8px; align-items:center; font-weight:800; white-space:nowrap;"><input type="checkbox" name="row_needs_review_0" checked> 필요</label></td>
        <td><input type="text" name="row_review_reason_0" value="관리자가 직접 추가한 행입니다."></td>
        <td><button class="danger-btn" style="padding:9px 10px;" type="button" onclick="removeMealAiRow(this)">삭제</button></td>
      `;
      body.appendChild(row);
      renumberMealAiRows();
    }

    function removeMealAiRow(button) {
      const row = button.closest(".meal-ai-preview-row");
      if (row) row.remove();
      renumberMealAiRows();
    }

    function mealAiCsvEscape(value) {
      const text = String(value ?? "");
      if (/[",\n]/.test(text)) {
        return `"${text.replace(/"/g, '""')}"`;
      }
      return text;
    }

    function downloadMealAiCsv() {
      const rows = Array.from(document.querySelectorAll(".meal-ai-preview-row"));
      const headers = ["date", "menu_name", "allergy_codes", "created_by", "created_at", "trash"];
      const lines = [headers.join(",")];

      rows.forEach((row) => {
        const values = [
          row.querySelector('[name^="row_date_"]')?.value || "",
          row.querySelector('[name^="row_menu_name_"]')?.value || "",
          row.querySelector('[name^="row_allergy_codes_"]')?.value || "",
          row.querySelector('[name^="row_created_by_"]')?.value || "",
          row.querySelector('[name^="row_created_at_"]')?.value || "",
          row.querySelector('[name^="row_trash_"]')?.value || "0"
        ];
        lines.push(values.map(mealAiCsvEscape).join(","));
      });

      const blob = new Blob(["\uFEFF" + lines.join("\n")], { type: "text/csv;charset=utf-8;" });
      const link = document.createElement("a");
      link.href = URL.createObjectURL(blob);
      link.download = "meal_rows_preview.csv";
      document.body.appendChild(link);
      link.click();
      link.remove();
      URL.revokeObjectURL(link.href);
    }

    function showMealAnalysisOverlay(message, isError = false) {
      const overlay = document.getElementById("mealAnalysisOverlay");
      const dialog = document.getElementById("mealAnalysisDialog");
      const title = document.getElementById("mealAnalysisTitle");
      const body = document.getElementById("mealAnalysisMessage");
      const closeBtn = document.getElementById("mealAnalysisCloseBtn");
      if (!overlay || !dialog || !title || !body || !closeBtn) return;
      overlay.classList.remove("hidden");
      dialog.classList.toggle("error", Boolean(isError));
      title.textContent = isError ? "분석에 실패했습니다" : "분석중입니다";
      body.textContent = message || (isError ? "원인을 확인할 수 없어." : "급식표를 읽고 AI가 메뉴와 알레르기 번호를 정리하고 있어. 잠시만 기다려줘.");
      closeBtn.classList.toggle("hidden", !isError);
    }

    function hideMealAnalysisOverlay() {
      const overlay = document.getElementById("mealAnalysisOverlay");
      if (overlay) overlay.classList.add("hidden");
    }

    
    function buildStudentNumberFromFields(gradeId, classId, seqId, targetId) {
      const grade = document.getElementById(gradeId)?.value.trim() || "";
      const classNo = document.getElementById(classId)?.value.trim() || "";
      const seq = document.getElementById(seqId)?.value.trim() || "";
      let value = "";
      if (grade && classNo && seq) {
        value = `${grade}${String(Number(classNo)).padStart(2, "0")}${String(Number(seq)).padStart(2, "0")}`;
      }
      const target = document.getElementById(targetId);
      if (target) target.value = value;
      return value;
    }

    function updateStudentNumberPreview() {
      return buildStudentNumberFromFields("studentGrade", "studentClassNo", "studentSeq", "studentNoPreview");
    }

    function updateEditStudentNumberPreview() {
      return buildStudentNumberFromFields("editStudentGrade", "editStudentClassNo", "editStudentSeq", "editStudentId");
    }

    function splitStudentNumber(studentNumber) {
      const s = String(studentNumber || "").replace(/\D/g, "");
      if (s.length !== 5) return {grade:"", classNo:"1", seq:""};
      return {
        grade: s.slice(0,1),
        classNo: String(Number(s.slice(1,3)) || 1),
        seq: String(Number(s.slice(3,5)) || "")
      };
    }

    let isDarkMode = true;

    function syncAdminTheme() {
      document.body.classList.toggle("light-mode", !isDarkMode);
      const btn = document.getElementById("themeToggleBtn");
      if (btn) {
        btn.textContent = isDarkMode ? "화이트모드" : "다크모드";
      }
      const logo = document.getElementById("adminLogo");
      if (logo) {
        logo.src = isDarkMode
          ? "{{ url_for('static', filename='logoimage_dark.png') }}"
          : "{{ url_for('static', filename='logoimage_white.png') }}";
      }
      try {
        localStorage.setItem("adminDarkMode", isDarkMode ? "true" : "false");
      } catch (e) {}
    }

    function toggleAdminTheme() {
      isDarkMode = !isDarkMode;
      syncAdminTheme();
    }

    document.addEventListener("DOMContentLoaded", () => {
      try {
        const saved = localStorage.getItem("adminDarkMode");
        if (saved !== null) {
          isDarkMode = saved === "true";
        }
      } catch (e) {}

      syncAdminTheme();
      resetMenuForm();
      resetStudentForm();
      updateStudentNumberPreview();
      updateEditStudentNumberPreview();
      renumberMealAiRows();

      const mealAnalyzeForm = document.getElementById("mealAnalyzeForm");
      if (mealAnalyzeForm) {
        mealAnalyzeForm.addEventListener("submit", () => {
          showMealAnalysisOverlay("급식표를 읽고 AI가 메뉴와 알레르기 번호를 정리하고 있어. 잠시만 기다려줘.");
        });
      }

      const mealAnalysisError = {{ meal_error|tojson }};
      if (document.getElementById("mealAnalysisOverlay") && mealAnalysisError) {
        showMealAnalysisOverlay(mealAnalysisError, true);
      }

      const studentSheetForm = document.getElementById("studentSheetForm");
      if (studentSheetForm) {
        studentSheetForm.addEventListener("submit", async (event) => {
          event.preventDefault();
          setStudentSheetResult("학생 시트를 읽고 미리보기를 만드는 중...");
          try {
            const res = await fetch("/api/student/sheet-add", {
              method: "POST",
              body: new FormData(studentSheetForm)
            });
            const contentType = res.headers.get("content-type") || "";
            const data = contentType.includes("application/json")
              ? await res.json()
              : { ok: false, error: (await res.text()).replace(/<[^>]*>/g, " ").replace(/\s+/g, " ").trim() || "서버 오류가 발생했어." };
            if (!data.ok) throw new Error(data.error || "시트로 추가 실패");
            renderStudentSheetPreview(data.rows || [], data.skipped || [], data.warnings || []);
            setStudentSheetResult(`미리보기 ${data.preview_count || 0}명을 불러왔어. 확인 후 추가를 눌러줘.`);
          } catch (err) {
            setStudentSheetResult(err.message || "시트로 추가에 실패했어.");
          }
        });
      }
    });
  </script>
</body>
</html>
"""

# =========================
# 학생 페이지 HTML
# =========================
STUDENT_HTML = r"""
<!doctype html>
<html lang="ko">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width,initial-scale=1">
  <title>학생 페이지</title>
  <style>
    * { box-sizing: border-box; }
    body {
      margin: 0;
      background: #0d0f12;
      color: #f1f5f9;
      font-family: system-ui, "Segoe UI", "Apple SD Gothic Neo", "Malgun Gothic", sans-serif;
    }
    .topbar {
      height: 74px;
      background: #17191d;
      border-bottom: 1px solid #2a2f37;
      display: flex;
      align-items: center;
      justify-content: space-between;
      padding: 0 24px;
    }
    .title-wrap {
      display: flex;
      flex-direction: column;
      gap: 4px;
    }
    .page-title {
      font-size: 24px;
      font-weight: 900;
    }
    .page-sub {
      color: #9aa4b2;
      font-size: 13px;
    }
    .top-actions {
      display: flex;
      gap: 10px;
      align-items: center;
    }
    .btn {
      border: 1px solid #2a2f37;
      cursor: pointer;
      border-radius: 12px;
      padding: 10px 14px;
      font-weight: 800;
      font-size: 14px;
      text-decoration: none;
      display: inline-flex;
      align-items: center;
      justify-content: center;
      color: #f1f5f9;
      background: #1a1d22;
    }
    .btn-primary {
      background: #f3f4f6;
      color: #111827;
      border-color: transparent;
    }
    .page {
      max-width: 1480px;
      margin: 0 auto;
      padding: 24px;
    }
    .grid {
      display: grid;
      grid-template-columns: 1fr 1fr;
      gap: 20px;
      margin-bottom: 20px;
    }
    .card {
      background: #181b20;
      border: 1px solid #2a2f37;
      border-radius: 18px;
      padding: 20px;
      box-shadow: 0 18px 40px rgba(0,0,0,0.25);
    }
    .card h2 {
      margin: 0 0 16px 0;
      font-size: 22px;
    }
    .card h3 {
      margin: 0 0 12px 0;
      font-size: 18px;
    }
    .menu-item, .alert-item {
      padding: 14px;
      border: 1px solid rgba(255,255,255,0.06);
      border-radius: 14px;
      background: #1d2127;
      margin-bottom: 12px;
    }
    .danger {
      color: #fca5a5;
      font-weight: 800;
    }
    .safe {
      color: #86efac;
      font-weight: 800;
    }
    .muted {
      color: #9aa4b2;
      font-size: 14px;
    }
    .form-group { margin-bottom: 14px; }
    .form-group label {
      display: block;
      margin-bottom: 8px;
      font-weight: 800;
    }
    .form-group input, .form-group select {
      width: 100%;
      padding: 12px 14px;
      border-radius: 14px;
      border: 1px solid #2a2f37;
      background: #111418;
      color: #f1f5f9;
      font-size: 15px;
      outline: none;
    }
    .allergy-row {
      display: flex;
      gap: 10px;
    }
    .tag-wrap {
      display: flex;
      flex-wrap: wrap;
      gap: 10px;
      min-height: 56px;
      padding: 12px;
      border-radius: 14px;
      border: 1px dashed #38404c;
      background: rgba(255,255,255,0.02);
      margin-top: 8px;
    }
    .tag {
      display: inline-flex;
      align-items: center;
      gap: 8px;
      padding: 8px 12px;
      border-radius: 999px;
      background: #2a313a;
      color: white;
      font-size: 13px;
      font-weight: 800;
    }
    .tag button {
      border: none;
      background: transparent;
      color: white;
      cursor: pointer;
      font-weight: 900;
    }
    .row-actions {
      display: flex;
      gap: 10px;
      flex-wrap: wrap;
      margin-top: 16px;
    }

    .notification-settings {
      margin-top: 20px;
      padding: 18px;
      border: 1px solid #2a2f37;
      border-radius: 16px;
      background: #14171b;
    }
    .notification-status {
      display: flex;
      align-items: center;
      gap: 10px;
      margin: 12px 0;
      font-weight: 800;
    }
    .status-dot {
      width: 10px;
      height: 10px;
      border-radius: 999px;
      background: #94a3b8;
      flex: 0 0 auto;
    }
    .status-dot.granted { background: #22c55e; }
    .status-dot.default { background: #f59e0b; }
    .status-dot.denied { background: #ef4444; }
    .notification-help {
      margin-top: 10px;
      color: #9aa4b2;
      font-size: 13px;
      line-height: 1.6;
      white-space: pre-line;
    }

    .account-row {
      display: flex;
      gap: 12px;
      flex-wrap: wrap;
    }

    .account-field {
      display: flex;
      flex-direction: column;
      gap: 6px;
      flex: 1 1 0;
      min-width: 0;
    }

    .account-label {
      font-size: 13px;
      color: #9aa4b2;
      font-weight: 700;
    }

    .account-value {
      padding: 10px 12px;
      border-radius: 14px;
      border: 1px solid #2a2f37;
      background: #111418;
      color: #f1f5f9;
      font-size: 15px;
      white-space: nowrap;
      text-overflow: ellipsis;
      overflow: hidden;
    }
    @media (max-width: 1000px) {
      .grid {
        grid-template-columns: 1fr;
      }
      .topbar {
        height: auto;
        padding: 16px;
        flex-direction: column;
        align-items: flex-start;
        gap: 12px;
      }
      .allergy-row {
        flex-direction: column;
      }
    }
  </style>
</head>
<body>
  <header class="topbar">
    <div class="title-wrap">
      <div class="page-title">학생 페이지</div>
      <div class="page-sub">{{ login_name }}님, 오늘 메뉴와 내 알러지 현황을 확인할 수 있어.</div>
    </div>

    <div class="top-actions">
      <button class="btn" type="button" style="cursor:default;">{{ login_name }}님</button>
      <button class="btn" onclick="openStudentMyAccountModal()">내 계정관리</button>
      <a class="btn" href="{{ url_for('logout') }}">로그아웃</a>
      <a class="btn btn-primary" href="{{ url_for('kiosk_page') }}" target="_blank">RFID 페이지 열기</a>
    </div>
  </header>

  <main class="page">
    <section class="grid">
      <div class="card">
        <h2>오늘 메뉴 알러지 현황</h2>
        {% if menu_alerts %}
          {% for item in menu_alerts %}
          <div class="alert-item">
            <h3>{{ item.menu_name }}</h3>
            <div>메뉴 알러지: {{ item.allergy_names }}</div>
            {% if item.has_conflict %}
              <div class="danger" style="margin-top:8px;">주의: {{ item.matched_names }}</div>
            {% else %}
              <div class="safe" style="margin-top:8px;">해당 없음</div>
            {% endif %}
          </div>
          {% endfor %}
        {% else %}
          <div class="muted">오늘 등록된 메뉴가 없어.</div>
        {% endif %}
      </div>

      <div class="card">
        <h2>오늘 메뉴</h2>
        {% if today_menus %}
          {% for menu in today_menus %}
          <div class="menu-item">
            <h3>{{ menu.menu_name }}</h3>
            <div>알러지명: {{ menu.allergy_names }}</div>
          </div>
          {% endfor %}
        {% else %}
          <div class="muted">오늘 등록된 메뉴가 없어.</div>
        {% endif %}
      </div>
    </section>

    <section class="card">
      <h2>내 계정 관리 / 내 알러지 설정</h2>

      <div class="form-group">
        <label>계정 정보</label>
        <div class="account-row">
          <div class="account-field">
            <span class="account-label">아이디</span>
            <span class="account-value">{{ login_id }}</span>
          </div>
          <div class="account-field">
            <span class="account-label">이름</span>
            <span class="account-value">{{ login_name }}</span>
          </div>
          <div class="account-field">
            <span class="account-label">새 비밀번호</span>
            <input type="text" id="studentMyPw" placeholder="변경할 비밀번호를 입력해줘">
          </div>
        </div>
        <input type="hidden" id="studentMyId" value="{{ login_id }}">
        <input type="hidden" id="studentMyName" value="{{ login_name }}">
      </div>

      <div class="form-group">
        <label>내 알러지 선택</label>
        <div class="allergy-row">
          <select id="studentMyAllergySelect">
            <option value="">알러지를 선택해줘</option>
            <option value="none">없음</option>
            {% for code, name in allergy_map.items() %}
            <option value="{{ code }}">{{ code }}. {{ name }}</option>
            {% endfor %}
          </select>
          <button type="button" class="btn" onclick="addMyAllergy()">알러지 추가</button>
        </div>
        <div id="studentMyAllergyList" class="tag-wrap"></div>
      </div>

      <div class="notification-settings">
        <h3>알림센터 수신 설정</h3>
        <div class="muted">등교 전 급식 비교 결과를 푸시알림이나 이메일로 받을 수 있어.</div>

        <div style="margin-top:18px;font-weight:900;">푸시알림</div>
        <div class="notification-status">
          <span id="notificationStatusDot" class="status-dot"></span>
          <span id="notificationStatusText">확인 중...</span>
        </div>
        <div class="row-actions">
          <button id="notificationPermissionButton" class="btn" type="button" onclick="handleNotificationPermission()">알림 상태 확인</button>
          <button id="notificationTestButton" class="btn btn-primary" type="button" onclick="sendNotificationTest()" style="display:none;">푸시 테스트</button>
        </div>
        <div id="notificationHelp" class="notification-help"></div>

        <div style="height:1px;background:#2a2f37;margin:22px 0;"></div>
        <div style="font-weight:900;margin-bottom:12px;">이메일 알림</div>
        <div class="account-row">
          <div class="account-field" style="flex:2 1 320px;">
            <span class="account-label">받을 이메일</span>
            <input type="email" id="studentNotificationEmail" value="{{ notification_email }}" placeholder="student@example.com" autocomplete="email">
          </div>
          <div class="account-field" style="flex:0 1 180px;">
            <span class="account-label">알림 시간</span>
            <input type="time" id="studentNotificationTime" value="{{ notification_time }}">
          </div>
        </div>
        <label style="display:flex;align-items:center;gap:9px;margin-top:14px;font-weight:800;cursor:pointer;">
          <input type="checkbox" id="studentEmailNotifications" style="width:auto;" {% if email_notifications %}checked{% endif %}>
          설정한 시간에 오늘 급식 이메일 받기
        </label>
        <div class="row-actions">
          <button id="emailTestButton" class="btn btn-primary" type="button" onclick="sendEmailTest()" {% if not email_delivery_configured %}disabled{% endif %}>오늘 급식 테스트 이메일</button>
        </div>
        <div id="emailTestResult" class="notification-help">{% if not email_delivery_configured %}서버에 이메일 API 키와 발신 주소를 설정해야 테스트할 수 있어.{% endif %}</div>
      </div>

      <div class="row-actions">
        <button class="btn btn-primary" onclick="saveStudentMyAccount()">저장</button>
      </div>
    </section>
  </main>

  <script>
    const ALLERGY_MAP = {{ allergy_map | tojson }};
    let myAllergies = {{ my_allergy_codes | tojson }};

    function renderMyAllergies(isNone = false) {
      const wrap = document.getElementById("studentMyAllergyList");
      if (!wrap) return;

      if (isNone) {
        wrap.innerHTML = `<div class="tag"><span>없음</span></div>`;
        return;
      }

      wrap.innerHTML = "";
      if (myAllergies.length === 0) {
        wrap.innerHTML = `<div class="muted">현재 등록된 알러지가 없어.</div>`;
        return;
      }

      myAllergies.forEach(code => {
        const name = ALLERGY_MAP[code] ?? `알수없음(${code})`;
        const tag = document.createElement("div");
        tag.className = "tag";
        tag.innerHTML = `
          <span>${code}. ${name}</span>
          <button type="button" onclick="removeMyAllergy(${code})">×</button>
        `;
        wrap.appendChild(tag);
      });
    }

    function openStudentMyAccountModal() {
      window.scrollTo({ top: document.body.scrollHeight, behavior: "smooth" });
    }

    function addMyAllergy() {
      const select = document.getElementById("studentMyAllergySelect");
      const value = select.value;

      if (!value) return;

      if (value === "none") {
        myAllergies = [];
        renderMyAllergies(true);
        select.value = "";
        return;
      }

      const num = Number(value);
      if (!num) return;

      if (!myAllergies.includes(num)) {
        myAllergies.push(num);
        myAllergies.sort((a, b) => a - b);
      }

      renderMyAllergies(false);
      select.value = "";
    }

    function removeMyAllergy(code) {
      myAllergies = myAllergies.filter(x => x !== code);
      renderMyAllergies(false);
    }

    async function saveStudentMyAccount() {
      const newId = document.getElementById("studentMyId")?.value.trim() || "";
      const newName = document.getElementById("studentMyName")?.value.trim() || "";
      const newPw = document.getElementById("studentMyPw")?.value.trim() || "";
      const notificationEmail = document.getElementById("studentNotificationEmail")?.value.trim() || "";
      const emailNotifications = Boolean(document.getElementById("studentEmailNotifications")?.checked);
      const notificationTime = document.getElementById("studentNotificationTime")?.value || "07:30";

      const res = await fetch("/api/student/my-account/update", {
        method: "POST",
        headers: {"Content-Type": "application/json"},
        body: JSON.stringify({
          new_id: newId,
          new_name: newName,
          new_pw: newPw,
          allergy_codes: myAllergies,
          email: notificationEmail,
          email_notifications: emailNotifications,
          notification_time: notificationTime
        })
      });

      const data = await res.json();
      if (!data.ok) {
        alert(data.error || "저장 실패");
        return;
      }

      alert("내 정보가 수정됐어.");
      window.location.reload();
    }

    
    function buildStudentNumberFromFields(gradeId, classId, seqId, targetId) {
      const grade = document.getElementById(gradeId)?.value.trim() || "";
      const classNo = document.getElementById(classId)?.value.trim() || "";
      const seq = document.getElementById(seqId)?.value.trim() || "";
      let value = "";
      if (grade && classNo && seq) {
        value = `${grade}${String(Number(classNo)).padStart(2, "0")}${String(Number(seq)).padStart(2, "0")}`;
      }
      const target = document.getElementById(targetId);
      if (target) target.value = value;
      return value;
    }

    function updateStudentNumberPreview() {
      return buildStudentNumberFromFields("studentGrade", "studentClassNo", "studentSeq", "studentNoPreview");
    }

    function updateEditStudentNumberPreview() {
      return buildStudentNumberFromFields("editStudentGrade", "editStudentClassNo", "editStudentSeq", "editStudentId");
    }

    function splitStudentNumber(studentNumber) {
      const s = String(studentNumber || "").replace(/\D/g, "");
      if (s.length !== 5) return {grade:"", classNo:"1", seq:""};
      return {
        grade: s.slice(0,1),
        classNo: String(Number(s.slice(1,3)) || 1),
        seq: String(Number(s.slice(3,5)) || "")
      };
    }

    async function sendEmailTest() {
      const button = document.getElementById("emailTestButton");
      const result = document.getElementById("emailTestResult");
      const email = document.getElementById("studentNotificationEmail")?.value.trim() || "";

      if (!email || !email.includes("@")) {
        if (result) result.textContent = "받을 이메일 주소를 먼저 입력해줘.";
        return;
      }

      if (button) button.disabled = true;
      if (result) result.textContent = "오늘 급식 비교 결과를 이메일로 보내는 중...";
      try {
        const response = await fetch("/api/student/notification/email-test", {
          method: "POST",
          headers: {"Content-Type": "application/json"},
          body: JSON.stringify({ email })
        });
        const data = await response.json();
        if (!response.ok || !data.ok) {
          throw new Error(data.error || "테스트 이메일 발송에 실패했어.");
        }
        if (result) result.textContent = `테스트 이메일 발송 완료 · 상태: ${data.status || "완료"}`;
      } catch (error) {
        if (result) result.textContent = error.message || "테스트 이메일 발송에 실패했어.";
      } finally {
        if (button) button.disabled = false;
      }
    }

    function notificationPermissionState() {
      if (!("Notification" in window)) return "unsupported";
      return Notification.permission;
    }

    function renderNotificationPermissionState() {
      const state = notificationPermissionState();
      const dot = document.getElementById("notificationStatusDot");
      const textEl = document.getElementById("notificationStatusText");
      const button = document.getElementById("notificationPermissionButton");
      const testButton = document.getElementById("notificationTestButton");
      const help = document.getElementById("notificationHelp");
      if (!dot || !textEl || !button || !testButton || !help) return;

      dot.className = "status-dot";
      testButton.style.display = "none";

      if (state === "granted") {
        dot.classList.add("granted");
        textEl.textContent = "알림 권한이 허용되어 있어.";
        button.textContent = "권한 다시 확인";
        testButton.style.display = "inline-flex";
        help.textContent = "이 기기에서 급식 푸시알림을 받을 수 있어.";
        return;
      }

      if (state === "default") {
        dot.classList.add("default");
        textEl.textContent = "아직 알림 권한을 선택하지 않았어.";
        button.textContent = "알림 허용하기";
        help.textContent = "버튼을 누르면 브라우저의 알림 권한 요청창이 표시돼.";
        return;
      }

      if (state === "denied") {
        dot.classList.add("denied");
        textEl.textContent = "알림 권한이 차단되어 있어.";
        button.textContent = "차단 해제 방법 보기";
        help.textContent = "주소창 옆 사이트 설정 아이콘을 누른 뒤, 알림을 ‘허용’으로 변경해줘. 변경 후 이 버튼을 다시 눌러 상태를 확인하면 돼.";
        return;
      }

      textEl.textContent = "이 브라우저는 푸시알림을 지원하지 않아.";
      button.textContent = "지원되지 않음";
      button.disabled = true;
      help.textContent = "Chrome, Edge, Safari 등 알림을 지원하는 최신 브라우저에서 다시 접속해줘.";
    }

    async function handleNotificationPermission() {
      const state = notificationPermissionState();
      if (state === "unsupported") {
        renderNotificationPermissionState();
        return;
      }

      if (state === "default") {
        try {
          await Notification.requestPermission();
        } catch (error) {
          alert("알림 권한 요청 중 문제가 생겼어.");
        }
        renderNotificationPermissionState();
        return;
      }

      if (state === "denied") {
        alert("알림이 차단되어 있어. 주소창 옆 사이트 설정에서 알림을 허용으로 바꾼 뒤 다시 확인해줘.");
        renderNotificationPermissionState();
        return;
      }

      renderNotificationPermissionState();
    }

    function sendNotificationTest() {
      if (notificationPermissionState() !== "granted") {
        renderNotificationPermissionState();
        return;
      }
      new Notification("급식 안전 알림 테스트", {
        body: "푸시알림 권한이 정상적으로 설정됐어.",
        tag: "meal-safety-test"
      });
    }

    document.addEventListener("DOMContentLoaded", () => {
      renderMyAllergies(false);
      renderNotificationPermissionState();
      window.addEventListener("focus", renderNotificationPermissionState);
    });
  </script>
</body>
</html>
"""

# =========================
# 렌더
# =========================
def render_admin_page(
    title,
    subtitle,
    active_tab,
    meal_mode="upload",
    meal_preview=None,
    meal_ocr_text="",
    meal_error="",
    meal_created_by=None,
):
    selected_class = safe_str(request.args.get("class_key", "all")).strip() or "all"
    selected_date = safe_str(request.args.get("date", "")).replace("-", "").strip()

    student_rows = get_all_student_rows()
    lunch_log_rows = get_all_lunch_log_rows()
    available_class_keys = get_available_class_keys(student_rows=student_rows, lunch_log_rows=lunch_log_rows)
    available_classes = [{"key": "all", "label": "전체"}] + [
        {"key": key, "label": f"{key.split('-')[0]}학년 {int(key.split('-')[1])}반"}
        for key in available_class_keys if key != "unknown"
    ]

    if selected_class != "all":
        filtered_student_rows = [x for x in student_rows if x.get("class_key") == selected_class]
    else:
        filtered_student_rows = list(student_rows)

    student_summary = {
        "total": len(filtered_student_rows),
        "allergy": sum(1 for row in filtered_student_rows if parse_codes(row.get("allergy_codes"))),
        "rfid": sum(1 for row in filtered_student_rows if safe_str(row.get("rfid_id")).strip()),
        "no_rfid": sum(1 for row in filtered_student_rows if not safe_str(row.get("rfid_id")).strip()),
    }

    if selected_date:
        # 실제 스캔 날짜(scan_date)로 필터
        filtered_lunch_rows = [x for x in lunch_log_rows if x.get("scan_date") == selected_date]
    else:
        filtered_lunch_rows = list(lunch_log_rows)

    if selected_class != "all":
        filtered_lunch_rows = [x for x in filtered_lunch_rows if x.get("class_key") == selected_class]

    lunch_log_date_options = sorted(list({x["scan_date"] for x in lunch_log_rows if x.get("scan_date")}), reverse=True)
    active_lunch_date = selected_date or today_sheet_str()
    today_rows = [x for x in lunch_log_rows if x.get("scan_date") == today_sheet_str()]
    date_rows_for_attendance = [x for x in lunch_log_rows if x.get("scan_date") == active_lunch_date]
    today_stats = build_lunch_log_stats(today_rows, all_students=student_rows)
    attendance_summary = build_not_eaten_students(student_rows, date_rows_for_attendance, selected_class=selected_class)
    selected_date_input = ""
    if selected_date and len(selected_date) == 8:
        selected_date_input = f"{selected_date[:4]}-{selected_date[4:6]}-{selected_date[6:8]}"
    elif not selected_date and len(active_lunch_date) == 8:
        selected_date_input = f"{active_lunch_date[:4]}-{active_lunch_date[4:6]}-{active_lunch_date[6:8]}"

    today_sheet = today_sheet_str()
    ai_date_options = sorted(
        {
            compact_date_value(row.get("date"))
            for row in get_all_menu_rows(include_trash=False)
            if compact_date_value(row.get("date"))
        },
        reverse=True,
    )
    if today_sheet not in ai_date_options:
        ai_date_options.insert(0, today_sheet)

    return render_template_string(
        BASE_HTML,
        page_title=title,
        title=title,
        subtitle=subtitle,
        active_tab=active_tab,
        menu_rows=get_all_menu_rows(),
        menu_groups=get_menu_groups(),
        student_rows=filtered_student_rows,
        student_summary=student_summary,
        student_class_groups=build_class_groups(filtered_student_rows),
        lunch_log_rows=filtered_lunch_rows,
        lunch_log_class_groups=build_class_groups(filtered_lunch_rows),
        trash_rows=get_trash_rows(),
        allergy_map=ALLERGY_MAP,
        default_admin_name=DEFAULT_ADMIN_NAME,
        default_student_password=DEFAULT_STUDENT_PASSWORD,
        rfid_dashboard_url=RFID_DASHBOARD_URL,
        logged_in=is_logged_in(),
        login_name=session.get("login_name", "관리자"),
        login_id=session.get("login_id", ""),
        available_classes=available_classes,
        selected_class=selected_class,
        selected_date=selected_date,
        selected_date_input=selected_date_input,
        lunch_log_date_options=lunch_log_date_options,
        today_stats=today_stats,
        attendance_summary=attendance_summary,
        active_lunch_date=active_lunch_date,
        today_sheet=today_sheet,
        ai_date_options=ai_date_options[:80],
        meal_mode=meal_mode,
        meal_preview=meal_preview or {"rows": [], "warnings": [], "warning_count": 0},
        meal_ocr_text=meal_ocr_text,
        meal_error=meal_error,
        meal_created_by=meal_created_by or session.get("login_name", DEFAULT_ADMIN_NAME),
        next_class_no=max([int(key.split('-')[1]) for key in available_class_keys if key != 'unknown'] + [1]) + 1,
        public_base_url=load_public_base_url()
    )



@app.errorhandler(Exception)
def handle_unexpected_error(e):
    logger.exception("처리되지 않은 서버 오류")
    message = "처리되지 않은 서버 오류" if IS_PRODUCTION else (safe_str(e) or "처리되지 않은 서버 오류")
    if request.path.startswith("/api/"):
        return jsonify({"ok": False, "error": message}), 500
    if IS_PRODUCTION:
        return "<h1>Internal Server Error</h1>", 500
    return f"<h1>Internal Server Error</h1><pre style='white-space:pre-wrap'>{message}</pre>", 500

# =========================
# 로그인 / 홈
# =========================
@app.route("/login", methods=["GET", "POST"])
def login_page():
    if request.method == "GET":
        if is_logged_in():
            if is_admin():
                return redirect(url_for("student_manage_page"))
            if is_student():
                return redirect(url_for("student_home_page"))
        return render_template_string(LOGIN_HTML, error="")

    login_id = safe_str(request.form.get("login_id")).strip()
    login_pw = safe_str(request.form.get("login_pw")).strip()

    if not login_id or not login_pw:
        return render_template_string(LOGIN_HTML, error="아이디와 비밀번호를 입력해줘.")

    for row in get_all_login_rows():
        if row["id"] == login_id and verify_password(row["pw"], login_pw):
            if not is_password_hash(row["pw"]):
                update_cell_by_header(id_ws, row["row_index"], "pw", hash_password(login_pw))
            session["logged_in"] = True
            session.permanent = True
            session["login_id"] = row["id"]
            session["login_name"] = row["name"]
            session["role"] = row["role"]

            if row["role"] == "a":
                return redirect(url_for("student_manage_page"))
            if row["role"] == "s":
                return redirect(url_for("student_home_page"))

    return render_template_string(LOGIN_HTML, error="아이디 또는 비밀번호가 틀렸어.")

@app.get("/logout")
def logout():
    session.clear()
    return redirect(url_for("login_page"))

KIOSK_DISPLAY_HTML = r"""
<!DOCTYPE html>
<html lang="ko">
<head>
  <meta charset="UTF-8" />
  <meta name="viewport" content="width=device-width, initial-scale=1.0" />
  <title>급식 알러지 체크 시스템</title>
  <style>
    * { box-sizing: border-box; }
    :root {
      --shadow: 0 24px 60px rgba(0,0,0,0.14);
    }
    body {
      margin: 0;
      min-height: 100vh;
      display: flex;
      align-items: center;
      justify-content: center;
      font-family: 'Malgun Gothic', 'Apple SD Gothic Neo', sans-serif;
      transition: background .35s ease, color .35s ease;
      padding: 20px;
    }
    body.waiting {
      background: linear-gradient(135deg, #eef2f7, #dfe7f1);
      color: #334155;
    }
    body.detected {
      background: linear-gradient(135deg, #fff7ed, #ffedd5);
      color: #7c2d12;
    }
    body.ok {
      background: linear-gradient(135deg, #ecfdf5, #dcfce7);
      color: #14532d;
    }
    body.notfound {
      background: linear-gradient(135deg, #fff1f2, #ffe4e6);
      color: #881337;
    }
    .wrap {
      width: min(94vw, 860px);
      text-align: center;
      position: relative;
    }
    .admin-link {
      position: absolute;
      top: -6px;
      right: 0;
      display: inline-flex;
      align-items: center;
      justify-content: center;
      text-decoration: none;
      border: none;
      border-radius: 999px;
      padding: 12px 18px;
      font-size: 0.98rem;
      font-weight: 800;
      cursor: pointer;
      box-shadow: 0 8px 20px rgba(0,0,0,0.08);
      background: rgba(255,255,255,0.9);
      color: #334155;
      z-index: 2;
    }
    .card {
      display: none;
      background: rgba(255,255,255,0.97);
      border-radius: 26px;
      box-shadow: var(--shadow);
      animation: fadeIn .25s ease;
    }
    .card.active { display: block; }

    .waiting-card {
      border: 4px solid #6b7280;
      padding: 48px 28px;
    }
    .badge {
      display: inline-block;
      padding: 10px 18px;
      border-radius: 999px;
      background: #e5e7eb;
      color: #4b5563;
      font-weight: 700;
      margin-bottom: 20px;
      letter-spacing: 0.02em;
    }
    h1 {
      margin: 0 0 12px;
      font-size: clamp(2rem, 4vw, 3rem);
    }
    p {
      margin: 0;
      font-size: clamp(1.05rem, 2vw, 1.3rem);
      line-height: 1.7;
    }
    .waiting-desc { color: #64748b; }
    .rfid {
      margin: 34px auto 0;
      width: 130px;
      height: 130px;
      border-radius: 50%;
      border: 6px solid #94a3b8;
      position: relative;
      animation: pulse 1.8s infinite;
      background: radial-gradient(circle at 30% 30%, #ffffff, #e2e8f0);
    }
    .rfid::before,
    .rfid::after {
      content: '';
      position: absolute;
      border: 5px solid #94a3b8;
      border-left-color: transparent;
      border-bottom-color: transparent;
      border-radius: 50%;
      transform: rotate(45deg);
    }
    .rfid::before {
      width: 42px;
      height: 42px;
      top: 34px;
      left: 26px;
    }
    .rfid::after {
      width: 72px;
      height: 72px;
      top: 19px;
      left: 11px;
    }
    .sub {
      margin-top: 26px;
      font-size: 0.98rem;
      color: #94a3b8;
    }

    .detected-card {
      border: 5px solid #f97316;
      padding: 34px 30px;
      text-align: left;
      box-shadow: 0 24px 60px rgba(249,115,22,0.18);
    }
    .top {
      display: flex;
      justify-content: space-between;
      align-items: center;
      gap: 16px;
      flex-wrap: wrap;
      margin-bottom: 20px;
    }
    .title {
      font-size: clamp(1.8rem, 4vw, 2.8rem);
      font-weight: 800;
      margin: 0;
    }
    .status {
      padding: 10px 18px;
      border-radius: 999px;
      background: #ffedd5;
      color: #c2410c;
      font-weight: 800;
      font-size: 1rem;
    }
    .name {
      font-size: clamp(2.1rem, 5vw, 3.6rem);
      font-weight: 900;
      margin: 8px 0 10px;
      color: #9a3412;
    }
    .meta {
      display: grid;
      grid-template-columns: repeat(auto-fit, minmax(180px, 1fr));
      gap: 14px;
      margin: 24px 0;
    }
    .box {
      background: #fff7ed;
      border: 1px solid #fdba74;
      border-radius: 18px;
      padding: 16px 18px;
    }
    .label {
      font-size: 0.95rem;
      color: #9a3412;
      margin-bottom: 6px;
      font-weight: 700;
    }
    .value {
      font-size: 1.2rem;
      font-weight: 800;
      color: #7c2d12;
      word-break: keep-all;
    }
    .warn {
      margin-top: 10px;
      background: #fff1f2;
      border: 2px solid #fb7185;
      color: #be123c;
      border-radius: 18px;
      padding: 18px 20px;
      font-size: 1.1rem;
      line-height: 1.7;
      font-weight: 700;
    }
    .warn ul {
      margin: 10px 0 0 20px;
      padding: 0;
    }
    .warn li { margin: 6px 0; }

    body.ok .detected-card {
      border-color: #22c55e;
      box-shadow: 0 24px 60px rgba(34,197,94,0.18);
    }
    body.ok .status {
      background: #dcfce7;
      color: #15803d;
    }
    body.ok .name {
      color: #166534;
    }
    body.ok .box {
      background: #f0fdf4;
      border-color: #86efac;
    }
    body.ok .label {
      color: #166534;
    }
    body.ok .value {
      color: #14532d;
    }
    body.ok .warn {
      background: #ecfdf5;
      border-color: #6ee7b7;
      color: #047857;
    }

    .notfound-card {
      border: 5px solid #e11d48;
      padding: 44px 28px;
      box-shadow: 0 24px 60px rgba(225,29,72,0.18);
    }
    .icon {
      width: 110px;
      height: 110px;
      margin: 0 auto 22px;
      border-radius: 50%;
      background: #ffe4e6;
      display: grid;
      place-items: center;
      font-size: 3.2rem;
      font-weight: 900;
      color: #e11d48;
      border: 4px solid #fb7185;
    }
    .desc {
      color: #9f1239;
    }
    .code {
      margin: 24px auto 0;
      display: inline-block;
      padding: 12px 18px;
      background: #fff1f2;
      border: 1px solid #fda4af;
      border-radius: 14px;
      font-weight: 800;
      color: #be123c;
    }
    .help {
      margin-top: 18px;
      font-size: 0.96rem;
      color: #9f1239;
      white-space: pre-line;
    }
    @keyframes pulse {
      0% { transform: scale(1); box-shadow: 0 0 0 0 rgba(148,163,184,0.35); }
      70% { transform: scale(1.03); box-shadow: 0 0 0 24px rgba(148,163,184,0); }
      100% { transform: scale(1); box-shadow: 0 0 0 0 rgba(148,163,184,0); }
    }
    @keyframes fadeIn {
      from { opacity: 0; transform: translateY(8px); }
      to { opacity: 1; transform: translateY(0); }
    }
  </style>
</head>
<body class="waiting">
  <div class="wrap">
    <a class="admin-link" href="/admin">관리자 페이지</a>

    <section id="waiting" class="card waiting-card active">
      <div class="badge">대기중</div>
      <h1>학생증을 태그해 주세요</h1>
      <p class="waiting-desc">RFID 리더기에 학생증을 가까이 대면<br>알러지 위험도를 바로 확인할 수 있어요.</p>
      <div class="rfid" aria-hidden="true"></div>
      <div class="sub">급식 알러지 체크 시스템</div>
    </section>

    <section id="detected" class="card detected-card">
      <div class="top">
        <h1 class="title">학생 정보 감지됨</h1>
        <div class="status" id="detectedStatus">주의</div>
      </div>

      <div class="name" id="detectedName">카리나</div>

      <div class="meta">
        <div class="box">
          <div class="label">학번</div>
          <div class="value" id="detectedStudentNo">20315</div>
        </div>
        <div class="box">
          <div class="label">오늘 메뉴</div>
          <div class="value" id="detectedMenu">카레라이스 / 계란국 / 새우튀김</div>
        </div>
        <div class="box">
          <div class="label">학생 알러지</div>
          <div class="value" id="detectedAllergyState">있음</div>
        </div>
        <div class="box">
          <div class="label">위험 메뉴</div>
          <div class="value" id="detectedUnsafeMenus">돈까스</div>
        </div>
        <div class="box" id="detectedRiskBox">
          <div class="label">알러지 위험도</div>
          <div class="value" id="detectedRisk">주의</div>
        </div>
      </div>

      <div class="warn" id="detectedWarnBox">
        <span id="detectedWarnIntro">감지된 알러지 유발 성분이 있어요.</span>
        <ul id="detectedWarnList">
          <li>난류(계란)</li>
          <li>갑각류(새우)</li>
        </ul>
      </div>
    </section>

    <section id="notfound" class="card notfound-card">
      <div class="icon">!</div>
      <h1 id="notFoundTitle">등록되지 않은 학생입니다</h1>
      <p class="desc" id="notFoundDesc">
        태그는 감지됐지만 학생 정보가 시스템에 등록되어 있지 않아요.<br>
        관리자 페이지에서 RFID와 학생 정보를 먼저 등록해 주세요.
      </p>
      <div class="code" id="notFoundCode">예시 태그 ID: 04A7C921B3</div>
      <div class="help" id="notFoundHelp">학번 조회 실패 (태그 미등록)</div>
    </section>
  </div>

  <script>
    function setView(view) {
      document.body.className = view;
      document.querySelectorAll('.card').forEach(card => card.classList.remove('active'));
      const target = document.getElementById(view);
      if (target) target.classList.add('active');
    }

    function setText(id, value, fallback='-') {
      const el = document.getElementById(id);
      if (el) el.textContent = (value === undefined || value === null || value === '') ? fallback : value;
    }

    function escapeHtml(value) {
      return String(value ?? '')
        .replace(/&/g, '&amp;')
        .replace(/</g, '&lt;')
        .replace(/>/g, '&gt;')
        .replace(/"/g, '&quot;')
        .replace(/'/g, '&#39;');
    }

    function renderDetected(data) {
      setView('detected');
      const isSafe = data.view === 'ok';
      document.body.className = isSafe ? 'ok' : 'detected';
      const statusText = data.status_text || (isSafe ? '안전' : '주의');
      const warningNames = Array.isArray(data.warning_names) ? data.warning_names : [];
      const studentAllergyNames = Array.isArray(data.student_allergy_names) ? data.student_allergy_names : [];
      const studentAllergyState = data.student_allergy_state || (studentAllergyNames.length ? '있음' : '없음');
      const unsafeMenus = Array.isArray(data.unsafe_menus) ? data.unsafe_menus : [];
      const matchedAllergies = Array.isArray(data.matched_allergies) ? data.matched_allergies : [];
      setText('detectedStatus', statusText, isSafe ? '안전' : '주의');
      setText('detectedName', data.name, '-');
      setText('detectedStudentNo', data.student_number, '-');
      setText('detectedMenu', data.menu_name, '-');
      setText('detectedAllergyState', studentAllergyState, studentAllergyNames.length ? '있음' : '없음');
      setText('detectedUnsafeMenus', unsafeMenus.length ? unsafeMenus.join(', ') : '해당 없음', '해당 없음');
      setText('detectedRisk', statusText, isSafe ? '안전' : '주의');

      const warnBox = document.getElementById('detectedWarnBox');
      const intro = document.getElementById('detectedWarnIntro');
      const list = document.getElementById('detectedWarnList');
      const riskBox = document.getElementById('detectedRiskBox');
      list.innerHTML = '';

      if (isSafe) {
        if (riskBox) riskBox.style.display = '';
        intro.textContent = studentAllergyNames.length
          ? `학생 알러지 있음: ${studentAllergyNames.join(', ')}. 오늘 메뉴와는 겹치지 않아요.`
          : '학생 알러지 없음. 안심하고 급식을 확인해도 돼요.';
        if (warnBox) {
          warnBox.style.background = '#ecfdf5';
          warnBox.style.borderColor = '#6ee7b7';
          warnBox.style.color = '#047857';
        }
      } else {
        if (riskBox) riskBox.style.display = 'none';
        intro.textContent = `위험 메뉴: ${unsafeMenus.join(', ') || '확인 필요'}`;
        if (warnBox) {
          warnBox.style.background = '#fff1f2';
          warnBox.style.borderColor = '#fb7185';
          warnBox.style.color = '#be123c';
        }
        const listSource = warningNames.length ? warningNames : matchedAllergies;
        if (listSource.length) {
          listSource.forEach(item => {
            const li = document.createElement('li');
            li.textContent = item;
            list.appendChild(li);
          });
        } else if (data.reason) {
          const li = document.createElement('li');
          li.textContent = data.reason;
          list.appendChild(li);
        }
      }
    }

    function renderNotFound(data) {
      setView('notfound');
      setText('notFoundTitle', data.title || '등록되지 않은 학생입니다');
      const desc = document.getElementById('notFoundDesc');
      if (desc) {
        desc.innerHTML = escapeHtml(
          data.subtitle || `태그는 감지됐지만 학생 정보가 시스템에 등록되어 있지 않아요.
관리자 페이지에서 RFID와 학생 정보를 먼저 등록해 주세요.`
        ).replace(/\n/g, '<br>');
      }
      setText('notFoundCode', `태그 ID: ${data.uid || '-'}`, '태그 ID: -');
      setText('notFoundHelp', data.reason || '', '');
    }

    async function refreshState() {
      try {
        const res = await fetch('/api/last_scan', { cache: 'no-store' });
        const data = await res.json();
        const view = data.view || 'waiting';

        if (view === 'waiting') {
          setView('waiting');
          return;
        }
        if (view === 'ok' || view === 'warning') {
          renderDetected(data);
          return;
        }
        renderNotFound(data);
      } catch (err) {
        renderNotFound({
          title: '서버 연결 오류',
          subtitle: `키오스크 상태를 불러오는 중 문제가 생겼어요.
잠시 후 다시 시도해 주세요.`,
          uid: '-',
          reason: String(err || ''),
        });
      }
    }

    refreshState();
    setInterval(refreshState, 1000);
  </script>
</body>
</html>
"""

@app.get("/")
def root_home():
    return redirect(url_for("home"))

@app.get("/kiosk")
def kiosk_page():
    return render_template_string(KIOSK_DISPLAY_HTML, auto_reset_seconds=AUTO_RESET_SECONDS)


@app.get("/favicon.ico")
def favicon():
    return "", 204


@app.get("/admin")
def home():
    if is_logged_in():
        if is_admin():
            return redirect(url_for("student_manage_page"))
        if is_student():
            return redirect(url_for("student_home_page"))
    return redirect(url_for("login_page"))

# =========================
# 관리자 페이지 라우트
# =========================
@app.get("/student-manage")
def student_manage_page():
    ok, response = require_admin()
    if not ok:
        return response
    return render_admin_page("학생관리", "학생 등록, 수정, 삭제를 한 화면에서 할 수 있어.", "student_manage")

@app.get("/student-add")
def student_add_page():
    ok, response = require_admin()
    if not ok:
        return response
    class_key = safe_str(request.args.get("class_key", "")).strip()
    if class_key:
        return redirect(url_for("student_manage_page", class_key=class_key))
    return redirect(url_for("student_manage_page"))

@app.get("/menu-add")
def menu_add_page():
    ok, response = require_admin()
    if not ok:
        return response
    return redirect(url_for("menu_manage_page"))

@app.get("/menu-manage")
def menu_manage_page():
    ok, response = require_admin()
    if not ok:
        return response
    return render_admin_page("급식관리", "급식 추가, 수정, 삭제를 한 화면에서 할 수 있어.", "menu_manage")

@app.get("/lunch-log")
def lunch_log_page():
    ok, response = require_admin()
    if not ok:
        return response
    return render_admin_page("급식명단", "날짜별 급식 기록을 확인할 수 있어.", "lunch_log")

@app.get("/trash")
def trash_page():
    ok, response = require_admin()
    if not ok:
        return response
    return render_admin_page("휴지통", "삭제된 데이터를 복구하거나 완전히 삭제할 수 있어.", "trash")

@app.get("/admin/ai-tools")
def ai_tools_page():
    ok, response = require_admin()
    if not ok:
        return response
    return render_admin_page("AI안전도우미", "대체급식, 사고 대응, 오늘 위험 브리핑을 한 화면에서 생성할 수 있어.", "ai_tools")

@app.get("/admin/meal/upload")
def meal_upload_page():
    ok, response = require_admin()
    if not ok:
        return response
    return render_admin_page(
        "AI급식추가",
        "급식표 이미지와 스프레드시트를 AI로 분석하고 저장 전 검토할 수 있어.",
        "meal_ai",
        meal_mode="upload",
        meal_error="",
        meal_created_by=session.get("login_name", DEFAULT_ADMIN_NAME),
    )

@app.post("/admin/meal/analyze")
def meal_analyze_page():
    ok, response = require_admin()
    if not ok:
        return response

    created_by = safe_str(request.form.get("created_by")).strip() or session.get("login_name", DEFAULT_ADMIN_NAME)
    ocr_hint = safe_str(request.form.get("ocr_hint")).strip()
    upload = request.files.get("meal_file") or request.files.get("meal_image")

    if not upload or not safe_str(upload.filename).strip():
        return render_admin_page(
            "AI급식추가",
            "급식표 이미지와 스프레드시트를 AI로 분석하고 저장 전 검토할 수 있어.",
            "meal_ai",
            meal_mode="upload",
            meal_error="급식표 이미지나 스프레드시트 파일을 먼저 선택해줘.",
            meal_created_by=created_by,
        ), 400

    if not is_allowed_meal_upload(upload.filename):
        return render_admin_page(
            "AI급식추가",
            "급식표 이미지와 스프레드시트를 AI로 분석하고 저장 전 검토할 수 있어.",
            "meal_ai",
            meal_mode="upload",
            meal_error="PNG, JPG, JPEG, WEBP, BMP, XLSX, CSV, TSV 파일만 업로드할 수 있어.",
            meal_created_by=created_by,
        ), 400

    safe_name = secure_filename(upload.filename) or "meal_upload"
    saved_path = MEAL_UPLOAD_DIR / f"{datetime.now().strftime('%Y%m%d_%H%M%S_%f')}_{safe_name}"
    upload.save(saved_path)

    try:
        if is_spreadsheet_meal_upload(upload.filename):
            ocr_text = extract_meal_spreadsheet_text(saved_path)
            source_label = "스프레드시트"
        else:
            ocr_text = extract_text(saved_path)
            source_label = "OCR"
    except (OCRServiceError, ValueError) as exc:
        return render_admin_page(
            "AI급식추가",
            "급식표 이미지와 스프레드시트를 AI로 분석하고 저장 전 검토할 수 있어.",
            "meal_ai",
            meal_mode="upload",
            meal_error=safe_str(exc),
            meal_created_by=created_by,
        ), 400

    analysis_input = ocr_text
    if ocr_hint:
        analysis_input = f"{ocr_text}\n\n[관리자 메모]\n{ocr_hint}"
    analysis_input = f"[입력 방식: {source_label}]\n{analysis_input}"

    analysis = analyze_meal_ocr_text(analysis_input, AI_ALLERGY_DICT)
    preview = build_meal_preview_context(analysis)
    return render_admin_page(
        "AI급식추가",
        "AI 분석 결과를 확인하고 food_menu 저장 형식으로 수정할 수 있어.",
        "meal_ai",
        meal_mode="preview",
        meal_preview=preview,
        meal_ocr_text=ocr_text,
        meal_error="",
        meal_created_by=created_by,
    )

@app.post("/admin/meal/save")
def meal_save_page():
    ok, response = require_admin()
    if not ok:
        return response

    created_by = safe_str(request.form.get("created_by")).strip() or session.get("login_name", DEFAULT_ADMIN_NAME)
    preview = build_meal_preview_context_from_form(request.form)
    ocr_text = safe_str(request.form.get("ocr_text")).strip()

    try:
        analysis = parse_preview_form(request.form)
        if not analysis.get("rows"):
            raise ValueError("저장할 급식 행이 없어. 날짜와 메뉴를 1개 이상 입력해줘.")
        save_ai_meal_analysis(menu_ws, analysis, created_by, now_time_str)
    except Exception as exc:
        return render_admin_page(
            "AI급식추가",
            "AI 분석 결과를 확인하고 food_menu 저장 형식으로 수정할 수 있어.",
            "meal_ai",
            meal_mode="preview",
            meal_preview=preview,
            meal_ocr_text=ocr_text,
            meal_error=safe_str(exc) or "AI 급식표 저장에 실패했어.",
            meal_created_by=created_by,
        ), 400

    return redirect(url_for("menu_manage_page"))


@app.post("/api/ai/safety-plan")
def api_ai_safety_plan():
    ok, response = require_admin_api()
    if not ok:
        return response

    data = request.get_json(silent=True) or {}
    date_value = compact_date_value(data.get("date")) or today_sheet_str()
    student_number = safe_str(data.get("student_number")).strip()
    context = safe_str(data.get("context")).strip()

    if not student_number:
        return jsonify({"ok": False, "error": "학생을 먼저 선택해줘."}), 400

    student = get_student_by_student_number(student_number, include_trash=False)
    if not student:
        return jsonify({"ok": False, "error": "학생 정보를 찾지 못했어."}), 404

    student_context = build_student_ai_context(student, date_value)
    if not student_context["menus"]:
        return jsonify({"ok": False, "error": f"{date_value} 날짜 급식 메뉴가 없어."}), 400

    plan = generate_allergy_safety_plan(
        student_name=safe_str(student.get("name")).strip(),
        student_number=student_number,
        student_allergies=student_context["student_allergy_names"],
        unsafe_menus=student_context["unsafe_menus"],
        matched_allergies=student_context["matched_allergies"],
        all_menus=student_context["menu_names"],
        context=context,
    )

    return jsonify({
        "ok": True,
        "date": date_value,
        "student": {
            "name": safe_str(student.get("name")).strip(),
            "student_number": student_number,
            "class_label": safe_str(student.get("class_label")).strip(),
            "student_label": safe_str(student.get("student_label")).strip(),
            "allergy_names": student_context["student_allergy_names"],
        },
        "menus": student_context["menu_names"],
        "unsafe_menus": student_context["unsafe_menus"],
        "matched_allergies": student_context["matched_allergies"],
        "plan": plan,
    })


@app.post("/api/ai/daily-brief")
def api_ai_daily_brief():
    ok, response = require_admin_api()
    if not ok:
        return response

    data = request.get_json(silent=True) or {}
    date_value = compact_date_value(data.get("date")) or today_sheet_str()
    context = build_ai_risk_context(date_value)

    if not context["menus"]:
        return jsonify({"ok": False, "error": f"{date_value} 날짜 급식 메뉴가 없어."}), 400

    brief = generate_daily_ai_brief(
        date_value=context["date"],
        risk_items=context["risk_items"],
        all_menus=context["menu_names"],
    )

    return jsonify({
        "ok": True,
        "date": context["date"],
        "menus": context["menu_names"],
        "risk_count": len(context["risk_items"]),
        "risk_items": context["risk_items"],
        "brief": brief,
    })


@app.post("/api/ai/menu-review")
def api_ai_menu_review():
    ok, response = require_admin_api()
    if not ok:
        return response

    data = request.get_json(silent=True) or {}
    date_value = compact_date_value(data.get("date")) or today_sheet_str()
    menu_rows = []

    for item in data.get("menus") or []:
        menu_name = safe_str((item or {}).get("menu_name")).strip()
        if not menu_name:
            continue
        menu_rows.append({
            "date": date_value,
            "menu_name": menu_name,
            "allergy_codes": ",".join(str(code) for code in sorted(parse_codes((item or {}).get("allergy_codes", [])))),
        })

    if not menu_rows:
        menu_rows = [
            {
                "date": row.get("date"),
                "menu_name": row.get("menu_name"),
                "allergy_codes": row.get("allergy_codes"),
            }
            for row in get_menu_rows_by_date(date_value, include_trash=False)
        ]

    if not menu_rows:
        return jsonify({"ok": False, "error": f"{date_value} 날짜 급식 메뉴가 없어."}), 400

    review = generate_menu_allergy_review(menu_rows, AI_ALLERGY_DICT)
    return jsonify({
        "ok": True,
        "date": date_value,
        "menus": menu_rows,
        "review": review,
    })

# =========================
# 학생 페이지
# =========================
@app.get("/student-home")
def student_home_page():
    ok, response = require_student()
    if not ok:
        if isinstance(response, str):
            return response
        return response

    login_id = safe_str(session.get("login_id")).strip()
    my_student = get_student_by_student_number(login_id, include_trash=False)
    login_settings = get_login_by_id(login_id, include_trash=False) or {}

    today_menus = [x for x in get_all_menu_rows() if x["date"] == today_sheet_str()]
    my_codes = parse_codes(my_student["allergy_codes"]) if my_student else set()

    menu_alerts = []
    for menu in today_menus:
        menu_codes = parse_codes(menu["allergy_codes"])
        matched = sorted(list(my_codes & menu_codes))
        menu_alerts.append({
            "menu_name": menu["menu_name"],
            "allergy_names": menu["allergy_names"],
            "matched_names": ", ".join(codes_to_names(matched)) if matched else "해당 없음",
            "has_conflict": len(matched) > 0
        })

    return render_template_string(
        STUDENT_HTML,
        login_name=session.get("login_name", "학생"),
        login_id=session.get("login_id", ""),
        allergy_map=ALLERGY_MAP,
        my_allergy_codes=sorted(list(my_codes)),
        my_allergy_names=my_student["allergy_names"] if my_student else "없음",
        menu_alerts=menu_alerts,
        today_menus=today_menus,
        notification_email=safe_str(login_settings.get("email")).strip(),
        email_notifications=bool(login_settings.get("email_notifications")),
        notification_time=safe_str(login_settings.get("notification_time")).strip() or "07:30",
        email_delivery_configured=email_delivery_configured(),
        rfid_dashboard_url=RFID_DASHBOARD_URL,
        public_base_url=load_public_base_url()
    )

# =========================
# 내 계정관리 API
# =========================
@app.post("/api/my-account/update")
def api_my_account_update():
    if not is_logged_in():
        return jsonify({"ok": False, "error": "로그인이 필요해"}), 401

    data = request.get_json(force=True) or {}

    current_id = safe_str(session.get("login_id")).strip()
    current_role = safe_str(session.get("role")).strip()
    new_id = safe_str(data.get("new_id")).strip()
    new_rfid_id = safe_str(data.get("new_rfid_id")).strip()
    new_name = safe_str(data.get("new_name")).strip()
    new_pw = safe_str(data.get("new_pw")).strip()

    if not new_id or not new_name:
        return jsonify({"ok": False, "error": "아이디와 이름은 입력해야 해"}), 400

    for row in get_all_login_rows():
        if row["id"] == new_id and row["id"] != current_id:
            return jsonify({"ok": False, "error": "이미 사용 중인 아이디야"}), 400

    login_row = get_login_by_id(current_id, include_trash=True)
    if not login_row:
        return jsonify({"ok": False, "error": "계정을 찾지 못했어"}), 404

    update_cell_by_header(id_ws, login_row["row_index"], "id", new_id)
    update_cell_by_header(id_ws, login_row["row_index"], "name", new_name)
    if new_pw:
        update_cell_by_header(id_ws, login_row["row_index"], "pw", hash_password(new_pw))

    if current_role == "s":
        student_row = get_student_by_student_number(current_id, include_trash=True)
        if student_row:
            update_cell_by_headers(student_ws, student_row["row_index"], ["student_number", "student_no"], new_id)
            update_cell_by_header(student_ws, student_row["row_index"], "name", new_name)

    session["login_id"] = new_id
    session["login_name"] = new_name

    return jsonify({"ok": True})

@app.post("/api/my-account/reset-password")
def api_my_account_reset_password():
    if not is_logged_in():
        return jsonify({"ok": False, "error": "로그인이 필요해"}), 401

    current_id = safe_str(session.get("login_id")).strip()
    login_row = get_login_by_id(current_id, include_trash=True)

    if not login_row:
        return jsonify({"ok": False, "error": "계정을 찾지 못했어"}), 404

    update_cell_by_header(id_ws, login_row["row_index"], "pw", hash_password(DEFAULT_STUDENT_PASSWORD))
    return jsonify({"ok": True})

# =========================
# 학생 본인 계정관리 API
# =========================
@app.post("/api/student/my-account/update")
def api_student_my_account_update():
    ok, response = require_student()
    if not ok:
        if isinstance(response, str):
            return jsonify({"ok": False, "error": response}), 403
        return jsonify({"ok": False, "error": "로그인이 필요해"}), 401

    data = request.get_json(force=True) or {}
    current_id = safe_str(session.get("login_id")).strip()

    new_id = safe_str(data.get("new_id")).strip()
    new_name = safe_str(data.get("new_name")).strip()
    new_pw = safe_str(data.get("new_pw")).strip()
    allergy_codes = data.get("allergy_codes", [])
    email = safe_str(data.get("email")).strip()
    email_notifications = value_is_enabled(data.get("email_notifications"))
    notification_time = safe_str(data.get("notification_time")).strip() or "07:30"

    if email and not is_valid_email(email):
        return jsonify({"ok": False, "error": "이메일 주소 형식을 확인해줘"}), 400
    if email_notifications and not email:
        return jsonify({"ok": False, "error": "이메일 알림을 켜려면 받을 이메일을 입력해야 해"}), 400
    if not re.fullmatch(r"(?:[01]\d|2[0-3]):[0-5]\d", notification_time):
        return jsonify({"ok": False, "error": "알림 시간을 다시 선택해줘"}), 400

    if not new_id or not new_name:
        return jsonify({"ok": False, "error": "아이디와 이름은 입력해야 해"}), 400

    for row in get_all_login_rows():
        if row["id"] == new_id and row["id"] != current_id:
            return jsonify({"ok": False, "error": "이미 사용 중인 아이디야"}), 400

    for row in get_all_student_rows():
        if row["student_number"] == new_id and row["student_number"] != current_id:
            return jsonify({"ok": False, "error": "이미 사용 중인 학번이야"}), 400

    cleaned_codes = []
    for code in allergy_codes:
        try:
            n = int(code)
            if 1 <= n <= 19:
                cleaned_codes.append(n)
        except:
            pass
    cleaned_codes = sorted(list(set(cleaned_codes)))

    login_row = get_login_by_id(current_id, include_trash=True)
    student_row = get_student_by_student_number(current_id, include_trash=True)

    if not login_row or not student_row:
        return jsonify({"ok": False, "error": "계정 정보를 찾지 못했어"}), 404

    update_cell_by_header(id_ws, login_row["row_index"], "id", new_id)
    update_cell_by_header(id_ws, login_row["row_index"], "name", new_name)
    if new_pw:
        update_cell_by_header(id_ws, login_row["row_index"], "pw", hash_password(new_pw))
    update_cell_by_header_create(id_ws, login_row["row_index"], "email", email)
    update_cell_by_header_create(id_ws, login_row["row_index"], "email_notifications", "1" if email_notifications else "0")
    update_cell_by_header_create(id_ws, login_row["row_index"], "notification_time", notification_time)

    update_cell_by_headers(student_ws, student_row["row_index"], ["student_number", "student_no"], new_id)
    update_cell_by_header(student_ws, student_row["row_index"], "name", new_name)
    update_cell_by_header(student_ws, student_row["row_index"], "allergy_codes", ",".join(str(x) for x in cleaned_codes))

    session["login_id"] = new_id
    session["login_name"] = new_name

    return jsonify({"ok": True})

@app.post("/api/student/notification/email-test")
def api_student_notification_email_test():
    ok, response = require_student()
    if not ok:
        if isinstance(response, str):
            return jsonify({"ok": False, "error": response}), 403
        return jsonify({"ok": False, "error": "로그인이 필요해"}), 401

    if not email_delivery_configured():
        return jsonify({"ok": False, "error": "서버의 이메일 API 설정이 아직 완료되지 않았어"}), 503

    data = request.get_json(silent=True) or {}
    current_id = safe_str(session.get("login_id")).strip()
    login_row = get_login_by_id(current_id, include_trash=False) or {}
    target_email = safe_str(data.get("email") or login_row.get("email")).strip()
    if not is_valid_email(target_email):
        return jsonify({"ok": False, "error": "받을 이메일 주소를 확인해줘"}), 400

    student = get_student_by_student_number(current_id, include_trash=False)
    if not student:
        return jsonify({"ok": False, "error": "학생 정보를 찾지 못했어"}), 404

    try:
        message = build_student_email_message(student, today_sheet_str())
        email_id = send_email(target_email, message)
    except Exception as exc:
        logger.exception("급식 테스트 이메일 발송 실패")
        return jsonify({"ok": False, "error": safe_str(exc) or "이메일 발송에 실패했어"}), 502

    return jsonify({
        "ok": True,
        "email_id": email_id,
        "status": message.get("status"),
    })


@app.post("/api/notifications/email-dispatch")
def api_notification_email_dispatch():
    if not notification_dispatch_authorized():
        return jsonify({"ok": False, "error": "알림 발송 토큰이 올바르지 않아"}), 401
    if not email_delivery_configured():
        return jsonify({"ok": False, "error": "이메일 API 설정이 완료되지 않았어"}), 503

    data = request.get_json(silent=True) or {}
    try:
        local_now = datetime.now(ZoneInfo(NOTIFICATION_TIMEZONE))
    except Exception:
        logger.warning("알림 시간대 설정을 읽지 못해 Asia/Seoul을 사용해")
        local_now = datetime.now(ZoneInfo("Asia/Seoul"))

    target_time = safe_str(data.get("time")).strip() or local_now.strftime("%H:%M")
    date_value = compact_date_value(data.get("date")) or local_now.strftime("%Y%m%d")
    if not re.fullmatch(r"(?:[01]\d|2[0-3]):[0-5]\d", target_time):
        return jsonify({"ok": False, "error": "발송 시간 형식은 HH:MM이어야 해"}), 400

    sent = []
    skipped = []
    failed = []
    for account in get_all_login_rows(include_trash=False):
        if account.get("role") != "s":
            continue
        if not account.get("email_notifications"):
            continue
        if safe_str(account.get("notification_time")).strip() != target_time:
            continue

        recipient = safe_str(account.get("email")).strip()
        if not is_valid_email(recipient):
            skipped.append({"id": account.get("id"), "reason": "이메일 주소 없음 또는 오류"})
            continue

        student = get_student_by_student_number(account.get("id"), include_trash=False)
        if not student:
            skipped.append({"id": account.get("id"), "reason": "학생 정보 없음"})
            continue

        try:
            message = build_student_email_message(student, date_value)
            email_id = send_email(recipient, message)
            sent.append({
                "id": account.get("id"),
                "email": recipient,
                "email_id": email_id,
                "status": message.get("status"),
            })
        except Exception as exc:
            logger.exception("예약 급식 이메일 발송 실패: %s", account.get("id"))
            failed.append({"id": account.get("id"), "email": recipient, "error": safe_str(exc)})

    payload = {
        "ok": not failed,
        "date": date_value,
        "time": target_time,
        "sent_count": len(sent),
        "skipped_count": len(skipped),
        "failed_count": len(failed),
        "sent": sent,
        "skipped": skipped,
        "failed": failed,
    }
    return jsonify(payload), (207 if failed else 200)


@app.post("/api/student/my-account/reset-password")
def api_student_my_account_reset_password():
    ok, response = require_student()
    if not ok:
        if isinstance(response, str):
            return jsonify({"ok": False, "error": response}), 403
        return jsonify({"ok": False, "error": "로그인이 필요해"}), 401

    current_id = safe_str(session.get("login_id")).strip()
    login_row = get_login_by_id(current_id, include_trash=True)

    if not login_row:
        return jsonify({"ok": False, "error": "계정을 찾지 못했어"}), 404

    update_cell_by_header(id_ws, login_row["row_index"], "pw", hash_password(DEFAULT_STUDENT_PASSWORD))
    return jsonify({"ok": True})

# =========================
# 메뉴 API
# =========================
@app.post("/api/menu")
def api_add_menu():
    ok, response = require_admin()
    if not ok:
        if isinstance(response, str):
            return jsonify({"ok": False, "error": response}), 403
        return jsonify({"ok": False, "error": "로그인이 필요해"}), 401

    data = request.get_json(force=True) or {}

    date = safe_str(data.get("date")).replace("-", "").strip()
    menu_name = safe_str(data.get("menu_name")).strip()
    created_by = safe_str(data.get("created_by")).strip()
    allergy_codes = data.get("allergy_codes", [])

    cleaned_codes = []
    for code in allergy_codes:
        try:
            n = int(code)
            if 1 <= n <= 19:
                cleaned_codes.append(n)
        except:
            pass
    cleaned_codes = sorted(list(set(cleaned_codes)))

    if not date or not menu_name or not created_by:
        return jsonify({"ok": False, "error": "날짜, 메뉴명, 추가자는 모두 입력해야 해"}), 400

    menu_ws.append_row([
        date,
        menu_name,
        ",".join(str(x) for x in cleaned_codes),
        created_by,
        now_time_str(),
        "0"
    ])

    return jsonify({"ok": True})

@app.post("/api/menu/group-detail")
def api_menu_group_detail():
    ok, response = require_admin()
    if not ok:
        if isinstance(response, str):
            return jsonify({"ok": False, "error": response}), 403
        return jsonify({"ok": False, "error": "로그인이 필요해"}), 401

    data = request.get_json(force=True) or {}
    row_indexes = data.get("row_indexes", [])

    cleaned = []
    for idx in row_indexes:
        try:
            n = int(idx)
            if n >= 2:
                cleaned.append(n)
        except:
            pass

    cleaned = sorted(list(set(cleaned)))
    if not cleaned:
        return jsonify({"ok": False, "error": "행 정보가 없어"}), 400

    all_rows = get_all_menu_rows(include_trash=True)
    target_rows = [x for x in all_rows if x["row_index"] in cleaned]

    if not target_rows:
        return jsonify({"ok": False, "error": "급식 정보를 찾지 못했어"}), 404

    target_rows.sort(key=lambda x: x["menu_name"])
    return jsonify({
        "ok": True,
        "date": target_rows[0]["date"],
        "menus": [x["menu_name"] for x in target_rows],
        "items": [
            {
                "menu_name": x["menu_name"],
                "allergy_codes": sorted(list(parse_codes(x["allergy_codes"])))
            }
            for x in target_rows
        ]
    })

@app.post("/api/menu/group-update")
def api_menu_group_update():
    ok, response = require_admin()
    if not ok:
        if isinstance(response, str):
            return jsonify({"ok": False, "error": response}), 403
        return jsonify({"ok": False, "error": "로그인이 필요해"}), 401

    data = request.get_json(force=True) or {}

    original_date = safe_str(data.get("original_date")).strip()
    new_date = safe_str(data.get("new_date")).strip()
    row_indexes = data.get("row_indexes", [])
    menu_items = data.get("menu_items", [])

    cleaned_row_indexes = []
    for idx in row_indexes:
        try:
            n = int(idx)
            if n >= 2:
                cleaned_row_indexes.append(n)
        except:
            pass
    cleaned_row_indexes = sorted(list(set(cleaned_row_indexes)))

    cleaned_items = []
    for item in menu_items:
        name = safe_str((item or {}).get("menu_name", "")).strip()
        if not name:
            continue

        raw_codes = (item or {}).get("allergy_codes", [])
        codes_cleaned = []
        for code in raw_codes:
            try:
                c = int(code)
                if 1 <= c <= 19:
                    codes_cleaned.append(c)
            except:
                pass
        codes_cleaned = sorted(list(set(codes_cleaned)))

        cleaned_items.append({
            "menu_name": name,
            "allergy_codes": codes_cleaned,
        })

    if not original_date or not new_date:
        return jsonify({"ok": False, "error": "날짜 정보가 올바르지 않아"}), 400

    if not cleaned_row_indexes:
        return jsonify({"ok": False, "error": "수정할 급식 행이 없어"}), 400

    if not cleaned_items:
        return jsonify({"ok": False, "error": "메뉴를 1개 이상 입력해야 해"}), 400

    all_rows = get_all_menu_rows(include_trash=True)
    target_rows = [x for x in all_rows if x["row_index"] in cleaned_row_indexes]

    if not target_rows:
        return jsonify({"ok": False, "error": "수정할 급식 정보를 찾지 못했어"}), 404

    base_created_by = target_rows[0]["created_by"] or session.get("login_name", DEFAULT_ADMIN_NAME)

    target_rows.sort(key=lambda x: x["row_index"])
    reuse_count = min(len(target_rows), len(cleaned_items))

    # 기존 행 재사용
    for i in range(reuse_count):
        update_cell_by_header(menu_ws, target_rows[i]["row_index"], "date", new_date)
        update_cell_by_header(menu_ws, target_rows[i]["row_index"], "menu_name", cleaned_items[i]["menu_name"])
        update_cell_by_header(
            menu_ws,
            target_rows[i]["row_index"],
            "allergy_codes",
            ",".join(str(c) for c in cleaned_items[i]["allergy_codes"])
        )

    # 남는 기존 행 삭제
    extra_old_rows = [x["row_index"] for x in target_rows[reuse_count:]]
    for row_index in sorted(extra_old_rows, reverse=True):
        hard_delete_rows(menu_ws, [row_index])

    # 메뉴가 더 많아지면 새로 추가
    for item in cleaned_items[reuse_count:]:
        menu_ws.append_row([
            new_date,
            item["menu_name"],
            ",".join(str(c) for c in item["allergy_codes"]),
            base_created_by,
            now_time_str(),
            "0"
        ])

    return jsonify({"ok": True})

@app.post("/api/menu/delete-selected")
def api_delete_selected_menus():
    ok, response = require_admin()
    if not ok:
        if isinstance(response, str):
            return jsonify({"ok": False, "error": response}), 403
        return jsonify({"ok": False, "error": "로그인이 필요해"}), 401

    data = request.get_json(force=True) or {}
    row_indexes = data.get("row_indexes", [])

    cleaned = []
    for idx in row_indexes:
        try:
            n = int(idx)
            if n >= 2:
                cleaned.append(n)
        except:
            pass

    cleaned = sorted(list(set(cleaned)))
    if not cleaned:
        return jsonify({"ok": False, "error": "삭제할 급식이 없어"}), 400

    soft_delete_rows(menu_ws, cleaned)
    return jsonify({"ok": True})

# =========================
# 학생 API
# =========================
@app.post("/api/student")
def api_add_student():
    ok, response = require_admin()
    if not ok:
        if isinstance(response, str):
            return jsonify({"ok": False, "error": response}), 403
        return jsonify({"ok": False, "error": "로그인이 필요해"}), 401

    data = request.get_json(force=True) or {}

    rfid_id = safe_str(data.get("rfid_id")).strip()
    name = safe_str(data.get("name")).strip()
    student_number = safe_str(data.get("student_number")).strip()
    grade = safe_str(data.get("grade")).strip()
    class_no = safe_str(data.get("class_no")).strip()
    student_seq = safe_str(data.get("student_seq")).strip()
    password = safe_str(data.get("password")).strip() or DEFAULT_STUDENT_PASSWORD

    if not student_number:
        student_number = build_student_number(grade, class_no, student_seq)

    if not rfid_id or not name or not student_number:
        return jsonify({"ok": False, "error": "RFID ID, 이름, 학년/반/번호는 모두 입력해야 해"}), 400

    parts = student_number_to_parts(student_number)
    if parts["class_key"] == "unknown":
        return jsonify({"ok": False, "error": "학번 형식이 올바르지 않아. 예: 10101"}), 400

    cleaned_codes = sorted(parse_codes(data.get("allergy_codes", [])))

    student_records = get_all_student_rows(include_trash=True)
    for r in student_records:
        if r["trash"] == "1":
            continue
        if r["rfid_id"] == rfid_id:
            return jsonify({"ok": False, "error": "이미 등록된 RFID ID야"}), 400
        if r["student_number"] == student_number:
            return jsonify({"ok": False, "error": "이미 등록된 학번이야"}), 400

    login_records = get_all_login_rows(include_trash=True)
    for r in login_records:
        if r["trash"] == "1":
            continue
        if r["id"] == student_number:
            return jsonify({"ok": False, "error": "이미 로그인 계정으로 등록된 학번이야"}), 400

    append_row_by_headers(student_ws, {
        "rfid_id": rfid_id,
        "name": name,
        "allergy_codes": ",".join(str(x) for x in cleaned_codes),
        "student_number": student_number,
        "student_no": student_number,
        "created_by": session.get("login_name", DEFAULT_ADMIN_NAME),
        "created_at": now_time_str(),
        "trash": "0",
    })

    append_row_by_headers(id_ws, {
        "id": student_number,
        "name": name,
        "pw": hash_password(password),
        "role": "s",
        "trash": "0",
    })

    return jsonify({"ok": True})


def build_student_sheet_import_plan(rows):
    existing_students = get_all_student_rows(include_trash=True)
    existing_logins = get_all_login_rows(include_trash=True)
    used_rfid = {row["rfid_id"] for row in existing_students if row["trash"] == "0" and row["rfid_id"]}
    used_student_numbers = {row["student_number"] for row in existing_students if row["trash"] == "0" and row["student_number"]}
    used_login_ids = {row["id"] for row in existing_logins if row["trash"] == "0" and row["id"]}

    accepted = []
    skipped = []

    for index, row in enumerate(rows or [], start=1):
        rfid_id = safe_str((row or {}).get("rfid_id")).strip()
        name = safe_str((row or {}).get("name")).strip()
        student_number = re.sub(r"\D", "", safe_str((row or {}).get("student_number")).strip())
        allergy_codes = sorted(parse_codes((row or {}).get("allergy_codes", [])))

        parts = student_number_to_parts(student_number)
        reason = ""
        if not name or not student_number:
            reason = "이름 또는 학번이 비어 있어."
        elif parts["class_key"] == "unknown":
            reason = "학번 형식이 올바르지 않아."
        elif rfid_id and rfid_id in used_rfid:
            reason = "이미 등록된 RFID야."
        elif student_number in used_student_numbers or student_number in used_login_ids:
            reason = "이미 등록된 학번이야."

        if reason:
            skipped.append({"index": index, "name": name, "student_number": student_number, "reason": reason})
            continue

        accepted.append({
            "rfid_id": rfid_id,
            "name": name,
            "student_number": student_number,
            "allergy_codes": allergy_codes,
        })
        if rfid_id:
            used_rfid.add(rfid_id)
        used_student_numbers.add(student_number)
        used_login_ids.add(student_number)

    return accepted, skipped


@app.post("/api/student/sheet-add")
def api_student_sheet_add():
    ok, response = require_admin_api()
    if not ok:
        return response

    upload = request.files.get("student_sheet")
    if not upload or not safe_str(upload.filename).strip():
        return jsonify({"ok": False, "error": "학생 명단 시트 파일을 선택해줘."}), 400

    suffix = Path(upload.filename).suffix.lower()
    if suffix not in {".xlsx", ".csv", ".tsv"}:
        return jsonify({"ok": False, "error": "XLSX, CSV, TSV 파일만 업로드할 수 있어."}), 400

    safe_name = secure_filename(upload.filename) or "student_sheet"
    saved_path = MEAL_UPLOAD_DIR / f"{datetime.now().strftime('%Y%m%d_%H%M%S_%f')}_{safe_name}"
    upload.save(saved_path)

    try:
        analysis = parse_student_spreadsheet_file(saved_path)
    except Exception as exc:
        return jsonify({"ok": False, "error": safe_str(exc) or "학생 시트를 읽지 못했어."}), 400

    rows, skipped = build_student_sheet_import_plan(analysis.get("rows", []))
    if not rows and skipped:
        return jsonify({
            "ok": False,
            "error": "미리보기 가능한 학생이 없어. " + "; ".join(f"{item['name'] or item['student_number'] or item['index']}: {item['reason']}" for item in skipped[:5]),
            "skipped": skipped,
            "warnings": analysis.get("warnings", []),
        }), 400

    return jsonify({
        "ok": True,
        "preview_count": len(rows),
        "rows": rows,
        "skipped": skipped,
        "skipped_count": len(skipped),
        "warnings": analysis.get("warnings", []),
    })


@app.post("/api/student/sheet-confirm")
def api_student_sheet_confirm():
    ok, response = require_admin_api()
    if not ok:
        return response

    data = request.get_json(force=True) or {}
    rows, skipped = build_student_sheet_import_plan(data.get("rows", []))
    if not rows:
        return jsonify({
            "ok": False,
            "error": "추가할 수 있는 학생이 없어. " + "; ".join(f"{item['name'] or item['student_number'] or item['index']}: {item['reason']}" for item in skipped[:5]),
            "skipped": skipped,
        }), 400

    added = []
    created_by = session.get("login_name", DEFAULT_ADMIN_NAME)
    created_at = now_time_str()
    student_rows_to_append = []
    login_rows_to_append = []

    for row in rows:
        rfid_id = row["rfid_id"]
        name = row["name"]
        student_number = row["student_number"]
        allergy_codes = row["allergy_codes"]
        student_rows_to_append.append({
            "rfid_id": rfid_id,
            "name": name,
            "allergy_codes": ",".join(str(x) for x in allergy_codes),
            "student_number": student_number,
            "student_no": student_number,
            "created_by": created_by,
            "created_at": created_at,
            "trash": "0",
        })

        login_rows_to_append.append({
            "id": student_number,
            "name": name,
            "pw": hash_password(DEFAULT_STUDENT_PASSWORD),
            "role": "s",
            "trash": "0",
        })

        added.append({"name": name, "student_number": student_number})

    append_rows_by_headers(student_ws, student_rows_to_append)
    append_rows_by_headers(id_ws, login_rows_to_append)

    return jsonify({
        "ok": True,
        "added_count": len(added),
        "skipped_count": len(skipped),
        "added": added,
        "skipped": skipped,
    })

@app.post("/api/student/update")
def api_student_update():
    ok, response = require_admin()
    if not ok:
        if isinstance(response, str):
            return jsonify({"ok": False, "error": response}), 403
        return jsonify({"ok": False, "error": "로그인이 필요해"}), 401

    data = request.get_json(force=True) or {}

    try:
        row_index = int(data.get("row_index", 0))
    except:
        return jsonify({"ok": False, "error": "학생 정보가 올바르지 않아"}), 400

    new_rfid_id = safe_str(data.get("new_rfid_id")).strip()
    new_name = safe_str(data.get("new_name")).strip()
    new_pw = safe_str(data.get("new_pw")).strip()
    grade = safe_str(data.get("grade")).strip()
    class_no = safe_str(data.get("class_no")).strip()
    student_seq = safe_str(data.get("student_seq")).strip()
    allergy_codes = data.get("allergy_codes", [])
    new_id = build_student_number(grade, class_no, student_seq)

    if row_index < 2 or not new_rfid_id or not new_name or not new_id:
        return jsonify({"ok": False, "error": "RFID ID, 이름, 학년/반/번호는 모두 입력해야 해"}), 400

    current_student_rows = get_all_student_rows(include_trash=True)
    target = next((row for row in current_student_rows if row["row_index"] == row_index), None)
    if not target:
        return jsonify({"ok": False, "error": "학생을 찾지 못했어"}), 404

    old_student_id = target["student_number"]

    for row in current_student_rows:
        if row["row_index"] != row_index and row["rfid_id"] == new_rfid_id and row["trash"] == "0":
            return jsonify({"ok": False, "error": "이미 사용 중인 RFID ID야"}), 400
        if row["row_index"] != row_index and row["student_number"] == new_id and row["trash"] == "0":
            return jsonify({"ok": False, "error": "이미 사용 중인 학번이야"}), 400

    for row in get_all_login_rows():
        if row["id"] == new_id and row["id"] != old_student_id:
            return jsonify({"ok": False, "error": "이미 사용 중인 로그인 아이디야"}), 400

    cleaned_codes = sorted(parse_codes(allergy_codes))

    update_cell_by_header(student_ws, row_index, "rfid_id", new_rfid_id)
    update_cell_by_header(student_ws, row_index, "name", new_name)
    update_cell_by_headers(student_ws, row_index, ["student_number", "student_no"], new_id)
    update_cell_by_header(student_ws, row_index, "allergy_codes", ",".join(str(x) for x in cleaned_codes))

    login_rows = get_all_login_rows(include_trash=True)
    for row in login_rows:
        if row["id"] == old_student_id:
            update_cell_by_header(id_ws, row["row_index"], "id", new_id)
            update_cell_by_header(id_ws, row["row_index"], "name", new_name)
            if new_pw:
                update_cell_by_header(id_ws, row["row_index"], "pw", hash_password(new_pw))
            break

    return jsonify({"ok": True})

@app.post("/api/student/reset-password")
def api_student_reset_password():
    ok, response = require_admin()
    if not ok:
        if isinstance(response, str):
            return jsonify({"ok": False, "error": response}), 403
        return jsonify({"ok": False, "error": "로그인이 필요해"}), 401

    data = request.get_json(force=True) or {}

    try:
        row_index = int(data.get("row_index", 0))
    except:
        return jsonify({"ok": False, "error": "학생 정보가 올바르지 않아"}), 400

    if row_index < 2:
        return jsonify({"ok": False, "error": "학생 정보가 올바르지 않아"}), 400

    target = None
    for row in get_all_student_rows(include_trash=True):
        if row["row_index"] == row_index:
            target = row
            break

    if not target:
        return jsonify({"ok": False, "error": "학생을 찾지 못했어"}), 404

    student_id = target["student_number"]

    for row in get_all_login_rows(include_trash=True):
        if row["id"] == student_id:
            update_cell_by_header(id_ws, row["row_index"], "pw", hash_password(DEFAULT_STUDENT_PASSWORD))
            return jsonify({"ok": True})

    return jsonify({"ok": False, "error": "로그인 계정을 찾지 못했어"}), 404

@app.post("/api/student/delete-selected")
def api_delete_selected_students():
    ok, response = require_admin()
    if not ok:
        if isinstance(response, str):
            return jsonify({"ok": False, "error": response}), 403
        return jsonify({"ok": False, "error": "로그인이 필요해"}), 401

    data = request.get_json(force=True) or {}
    row_indexes = data.get("row_indexes", [])

    cleaned = []
    for idx in row_indexes:
        try:
            n = int(idx)
            if n >= 2:
                cleaned.append(n)
        except:
            pass

    cleaned = sorted(list(set(cleaned)))
    if not cleaned:
        return jsonify({"ok": False, "error": "삭제할 학생이 없어"}), 400

    student_rows = get_all_student_rows(include_trash=True)
    login_rows = get_all_login_rows(include_trash=True)

    for row_index in cleaned:
        update_cell_by_header(student_ws, row_index, "trash", "1")

        target_student = next((x for x in student_rows if x["row_index"] == row_index), None)
        if target_student:
            student_id = target_student["student_number"]
            target_login = next((x for x in login_rows if x["id"] == student_id), None)
            if target_login:
                update_cell_by_header(id_ws, target_login["row_index"], "trash", "1")

    return jsonify({"ok": True})

# =========================
# 휴지통 API
# =========================
@app.post("/api/trash/menu/restore")
def api_trash_menu_restore():
    ok, response = require_admin()
    if not ok:
        if isinstance(response, str):
            return jsonify({"ok": False, "error": response}), 403
        return jsonify({"ok": False, "error": "로그인이 필요해"}), 401

    data = request.get_json(force=True) or {}
    row_index = data.get("row_index")

    try:
        row_index = int(row_index)
    except:
        return jsonify({"ok": False, "error": "row_index가 잘못됐어"}), 400

    if row_index < 2:
        return jsonify({"ok": False, "error": "복구할 수 없는 행이야"}), 400

    restore_rows(menu_ws, [row_index])
    return jsonify({"ok": True})

@app.post("/api/trash/student/restore")
def api_trash_student_restore():
    ok, response = require_admin()
    if not ok:
        if isinstance(response, str):
            return jsonify({"ok": False, "error": response}), 403
        return jsonify({"ok": False, "error": "로그인이 필요해"}), 401

    data = request.get_json(force=True) or {}
    row_index = data.get("row_index")

    try:
        row_index = int(row_index)
    except:
        return jsonify({"ok": False, "error": "row_index가 잘못됐어"}), 400

    if row_index < 2:
        return jsonify({"ok": False, "error": "복구할 수 없는 행이야"}), 400

    target_student = None
    for row in get_all_student_rows(include_trash=True):
        if row["row_index"] == row_index:
            target_student = row
            break

    if not target_student:
        return jsonify({"ok": False, "error": "학생을 찾지 못했어"}), 404

    student_id = target_student["student_number"]
    target_login = get_login_by_id(student_id, include_trash=True)

    update_cell_by_header(student_ws, row_index, "trash", "0")
    if target_login:
        update_cell_by_header(id_ws, target_login["row_index"], "trash", "0")

    return jsonify({"ok": True})

@app.post("/api/trash/menu/hard-delete")
def api_trash_menu_hard_delete():
    ok, response = require_admin()
    if not ok:
        if isinstance(response, str):
            return jsonify({"ok": False, "error": response}), 403
        return jsonify({"ok": False, "error": "로그인이 필요해"}), 401

    data = request.get_json(force=True) or {}
    row_index = data.get("row_index")

    try:
        row_index = int(row_index)
    except:
        return jsonify({"ok": False, "error": "row_index가 잘못됐어"}), 400

    if row_index < 2:
        return jsonify({"ok": False, "error": "삭제할 수 없는 행이야"}), 400

    hard_delete_rows(menu_ws, [row_index])
    return jsonify({"ok": True})

@app.post("/api/trash/student/hard-delete")
def api_trash_student_hard_delete():
    ok, response = require_admin()
    if not ok:
        if isinstance(response, str):
            return jsonify({"ok": False, "error": response}), 403
        return jsonify({"ok": False, "error": "로그인이 필요해"}), 401

    data = request.get_json(force=True) or {}
    row_index = data.get("row_index")

    try:
        row_index = int(row_index)
    except:
        return jsonify({"ok": False, "error": "row_index가 잘못됐어"}), 400

    if row_index < 2:
        return jsonify({"ok": False, "error": "삭제할 수 없는 행이야"}), 400

    target_student = None
    for row in get_all_student_rows(include_trash=True):
        if row["row_index"] == row_index:
            target_student = row
            break

    if not target_student:
        return jsonify({"ok": False, "error": "학생을 찾지 못했어"}), 404

    student_id = target_student["student_number"]
    target_login = get_login_by_id(student_id, include_trash=True)

    if target_login:
        hard_delete_rows(id_ws, [target_login["row_index"]])
    hard_delete_rows(student_ws, [row_index])

    return jsonify({"ok": True})

@app.post("/api/kiosk/scan")
def api_kiosk_scan():
    global LATEST_SCANNED_RFID_UID, LATEST_SCANNED_RFID_AT

    data = request.get_json(force=True) or {}
    token = safe_str(data.get("token")).strip()
    uid = safe_str(data.get("uid")).strip()

    logger.info(f"[KIOSK_SCAN] 요청 수신 uid={uid or '-'} token_present={'yes' if token else 'no'} from={request.remote_addr}")

    if KIOSK_SCAN_API_TOKEN and token != KIOSK_SCAN_API_TOKEN:
        logger.warning(f"[KIOSK_SCAN] 토큰 불일치 from={request.remote_addr}")
        return jsonify({"ok": False, "error": "인증 토큰이 올바르지 않아"}), 403

    if uid:
        LATEST_SCANNED_RFID_UID = uid
        LATEST_SCANNED_RFID_AT = now_time_str()
        logger.info(f"[KIOSK_SCAN] 최신 RFID 갱신 uid={uid} at={LATEST_SCANNED_RFID_AT}")
    else:
        logger.warning("[KIOSK_SCAN] uid 없이 요청이 들어왔어")

    result = process_scan(uid)
    update_last_scan_state_from_result(result)

    scan = (result or {}).get("scan", {}) or {}
    logger.info(
        "[KIOSK_SCAN] 처리 완료 uid=%s status=%s name=%s student_number=%s",
        uid or '-',
        safe_str(scan.get("status")) or '-',
        safe_str(scan.get("name")) or '-',
        safe_str(scan.get("student_number")) or '-',
    )

    status_code = 200 if result.get("ok", False) else 400
    return jsonify(result), status_code

@app.get("/api/latest-rfid")
def api_latest_rfid():
    ok, response = require_admin()
    if not ok:
        if isinstance(response, str):
            return jsonify({"ok": False, "error": response}), 403
        return jsonify({"ok": False, "error": "로그인이 필요해"}), 401

    return jsonify({
        "ok": True,
        "uid": LATEST_SCANNED_RFID_UID,
        "scanned_at": LATEST_SCANNED_RFID_AT,
    })

@app.get("/api/last_scan")
def api_last_scan():
    return jsonify(get_display_state())

@app.get("/api/kiosk/health")
def api_kiosk_health():
    return jsonify({"ok": True, "time": now_time_str(), "public_admin_url": load_public_base_url()})

@app.get("/health")
def health():
    return jsonify({"ok": True, "time": now_time_str(), "public_admin_url": load_public_base_url()})

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.getenv("PORT", "5001")), debug=False)














